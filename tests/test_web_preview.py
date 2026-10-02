"""Тесты постраничного рендера для визуального preview."""

import fitz
import pytest

from textalchemy.web import preview


def _make_pdf(path, pages: int = 1) -> None:
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page(width=200, height=300)
        page.insert_text((20, 30), f"Preview {index + 1}")
    doc.save(path)
    doc.close()


def test_pdf_page_count_and_render(tmp_path):
    pdf = tmp_path / "pages.pdf"
    _make_pdf(pdf, pages=3)

    assert preview.pdf_page_count(pdf) == 3

    data = preview.render_pdf_page_png(pdf, 0, dpi=72)
    assert data and data.startswith(b"\x89PNG")

    assert preview.render_pdf_page_png(pdf, 3, dpi=72) is None
    assert preview.render_pdf_page_png(pdf, -1, dpi=72) is None
    assert preview.render_pdf_page_png(pdf, 1, dpi=9999) is not None  # dpi clamped, still renders


def test_cached_page_count_and_png(tmp_path):
    pdf = tmp_path / "pages.pdf"
    _make_pdf(pdf, pages=2)
    preview_dir = tmp_path / "preview"
    preview_dir.mkdir()

    assert preview.cached_page_count(preview_dir, pdf, "target") == 2
    assert (preview_dir / "target-pages").is_file()

    data = preview.cached_page_png(preview_dir, pdf, "target", 0, dpi=72)
    assert data and data.startswith(b"\x89PNG")
    assert (preview_dir / "target-0-72dpi.png").is_file()
    assert preview.cached_page_png(preview_dir, pdf, "target", 0, dpi=72) == data

    high_resolution = preview.cached_page_png(preview_dir, pdf, "target", 0, dpi=144)
    assert high_resolution != data
    assert (preview_dir / "target-0-144dpi.png").is_file()


def test_ensure_pdf_returns_source_directly_for_pdf(tmp_path):
    pdf = tmp_path / "in.pdf"
    _make_pdf(pdf)
    preview_dir = tmp_path / "preview"
    preview_dir.mkdir()

    assert preview.ensure_pdf(preview_dir, pdf, "source") == pdf


def test_ensure_pdf_caches_office_conversion(tmp_path, monkeypatch):
    source = tmp_path / "sample.docx"
    source.write_bytes(b"docx")
    preview_dir = tmp_path / "preview"
    preview_dir.mkdir()
    calls = []

    def fake_convert(src, out):
        calls.append(src)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(b"%PDF")
        return True

    monkeypatch.setattr(preview, "convert_to_pdf", fake_convert)

    first = preview.ensure_pdf(preview_dir, source, "source")
    second = preview.ensure_pdf(preview_dir, source, "source")
    assert first == second == preview_dir / "source.pdf"
    assert len(calls) == 1


def test_failed_office_conversion_does_not_cache_unavailability(tmp_path, monkeypatch):
    source = tmp_path / "example.docx"
    source.write_bytes(b"office document")
    cache = tmp_path / "preview"
    cache.mkdir()
    monkeypatch.setattr(preview, "convert_to_pdf", lambda *_args: False)
    assert preview.cached_page_count(cache, source, "source") == 0
    assert not (cache / "source-pages").exists()
    (cache / "source-pages").write_text("0")  # Прежний нулевой кэш тоже не блокирует восстановление.

    def recovered(_source, destination):
        _make_pdf(destination, pages=2)
        return True

    monkeypatch.setattr(preview, "convert_to_pdf", recovered)
    assert preview.cached_page_count(cache, source, "source") == 2


def test_office_timeout_removes_profile_and_partial_files(tmp_path, monkeypatch):
    import subprocess

    source = tmp_path / "example.docx"
    source.write_bytes(b"office document")
    profiles = []
    monkeypatch.setattr(preview, "libreoffice_path", lambda: "test-office")

    def timeout(command, **kwargs):
        profile = next(arg.removeprefix("-env:UserInstallation=file:///") for arg in command if arg.startswith("-env:"))
        profiles.append(preview.Path(profile))
        raise subprocess.TimeoutExpired(command, kwargs["timeout"])

    monkeypatch.setattr(preview.subprocess, "run", timeout)
    assert not preview.convert_to_pdf(source, tmp_path / "result.pdf")
    assert profiles and all(not profile.parent.exists() for profile in profiles)
    assert source.read_bytes() == b"office document" and not (tmp_path / "result.pdf").exists()


def test_convert_to_pdf_via_libreoffice(tmp_path):
    executable = preview.libreoffice_path()
    if executable is None:
        pytest.skip("LibreOffice not available")
    try:
        import docx
    except ImportError:
        pytest.skip("python-docx not available")

    source = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_heading("Preview heading", level=1)
    document.add_paragraph("Some content for conversion.")
    document.save(source)

    output_pdf = tmp_path / "converted.pdf"
    assert preview.convert_to_pdf(source, output_pdf) is True
    assert output_pdf.is_file()
    assert preview.pdf_page_count(output_pdf) >= 1
