from __future__ import annotations

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from textalchemy.extract.latex import _format_run_text, _get_paragraph_style, clean_text, docx_to_latex


def test_clean_text_escapes_special_chars():
    assert clean_text("a & b") == "a \\& b"
    assert clean_text("100%") == "100\\%"
    assert clean_text("_foo_") == "\\_{}foo\\_{}"
    assert clean_text("hello") == "hello"


def test_clean_text_handles_backslash():
    assert clean_text("a\\b") == "a\\textbackslash{}b"


def test_clean_text_handles_ellipsis():
    assert clean_text("…") == "\\dots{}"


def test_format_run_text_bold(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("bold text")
    run.bold = True
    result = _format_run_text(run)
    assert result == "\\textbf{bold text}"


def test_format_run_text_italic(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("italic text")
    run.italic = True
    result = _format_run_text(run)
    assert result == "\\textit{italic text}"


def test_format_run_text_bold_italic(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("bold italic")
    run.bold = True
    run.italic = True
    result = _format_run_text(run)
    assert result == "\\textit{\\textbf{bold italic}}"


def test_format_run_text_plain(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("plain")
    result = _format_run_text(run)
    assert result == "plain"


def test_get_paragraph_style_heading1():
    doc = Document()
    p = doc.add_heading("Title", level=1)
    assert _get_paragraph_style(p) == "section"


def test_get_paragraph_style_heading2():
    doc = Document()
    p = doc.add_heading("Subtitle", level=2)
    assert _get_paragraph_style(p) == "subsection"


def test_get_paragraph_style_heading3():
    doc = Document()
    p = doc.add_heading("Subsubtitle", level=3)
    assert _get_paragraph_style(p) == "subsubsection"


def test_get_paragraph_style_normal():
    doc = Document()
    p = doc.add_paragraph("Normal text")
    assert _get_paragraph_style(p) is None


def test_docx_to_latex_bold_text(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Hello ")
    run2 = p.add_run("bold")
    run2.bold = True
    p.add_run(" world")
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "Hello " in tex
    assert "\\textbf{bold}" in tex
    assert " world" in tex


def test_docx_to_latex_italic_text(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("emphasized")
    run.italic = True
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\textit{emphasized}" in tex


def test_docx_to_latex_heading(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_heading("Chapter 1", level=1)
    doc.add_heading("Section 1", level=2)
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\section{Chapter 1}" in tex
    assert "\\subsection{Section 1}" in tex


def test_docx_to_latex_table(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\begin{tabular}" in tex
    assert "A & B" in tex
    assert "1 & 2" in tex
    assert "\\end{tabular}" in tex


def test_docx_to_latex_center_aligned(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    p = doc.add_paragraph("centered text")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\begin{center}" in tex
    assert "centered text" in tex
    assert "\\end{center}" in tex


def test_docx_to_latex_right_aligned(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    p = doc.add_paragraph("right text")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\begin{flushright}" in tex
    assert "right text" in tex
    assert "\\end{flushright}" in tex


def test_docx_to_latex_writes_to_file(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("Save test")
    doc.save(str(docx_path))

    tex_path = tmp_path / "output.tex"
    result = docx_to_latex(str(docx_path), str(tex_path))
    assert tex_path.is_file()
    assert tex_path.read_text(encoding="utf-8") == result


def test_docx_to_latex_preamble_present(tmp_path):
    docx_path = tmp_path / "input.docx"
    doc = Document()
    doc.add_paragraph("test")
    doc.save(str(docx_path))

    tex = docx_to_latex(str(docx_path))
    assert "\\documentclass[12pt,a4paper]{article}" in tex
    assert "\\begin{document}" in tex
    assert "\\end{document}" in tex


def test_docx_to_latex_missing_file():
    with pytest.raises(Exception):
        docx_to_latex("/nonexistent/path.docx")
