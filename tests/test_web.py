import importlib
import json
from html.parser import HTMLParser
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from textalchemy import __version__
from textalchemy.core.artifacts import ArtifactWorkspace
from textalchemy.core.database import Database
from textalchemy.core.diagnostics import ConversionReport
from textalchemy.core.document_model import ConversionMode
from textalchemy.core.inspection import DocumentInspection
from textalchemy.core.types import DocFormat
from textalchemy.web.main import app

client = TestClient(app)


@pytest.fixture(autouse=True)
def _isolate_task_store(monkeypatch, tmp_path_factory):
    """Фоновые задачи пишутся в отдельное временное хранилище, а не в user-data."""
    import importlib

    from textalchemy.web.tasks import TaskStore

    web_app = importlib.import_module("textalchemy.web.app")
    data = tmp_path_factory.mktemp("web-data")
    monkeypatch.setattr(web_app, "data_dir", data)
    monkeypatch.setattr(web_app, "db", Database(db_path=data / "library.db"))
    store = TaskStore(tmp_path_factory.mktemp("web-tasks"))
    monkeypatch.setattr(web_app, "tasks_store", store)
    from textalchemy.web.routes import convert as convert_route

    monkeypatch.setattr(convert_route, "tasks_store", store)
    yield
    _wait_for_convert_tasks()


def _client_task_store():
    from textalchemy.web.routes import convert as convert_route

    return convert_route.tasks_store


def _wait_for_convert_tasks(timeout=15.0):
    """Дождаться завершения фоновых задач очереди (воркер вне event loop)."""
    from textalchemy.web.queue import task_queue

    assert task_queue.wait_idle(timeout=timeout), "background convert tasks did not finish in time"


def test_web_version_uses_package_metadata():
    assert app.version == __version__


def test_locale_catalog_defaults_to_russian_and_is_extensible():
    manifest = client.get("/api/locales")
    catalog = client.get("/api/locales/ru")

    assert manifest.status_code == 200
    assert manifest.json() == {"default": "ru", "locales": [{"code": "ru", "name": "Русский"}]}
    assert catalog.status_code == 200
    assert catalog.json()["code"] == "ru"
    assert catalog.json()["messages"]["nav.convert"] == "Преобразовать"
    assert client.get("/api/locales/en").status_code == 404


def test_dashboard():
    resp = client.get("/")
    assert resp.status_code == 200
    assert "text/html" in resp.headers["content-type"]
    assert f"/static/css/style.css?v={__version__}" in resp.text
    assert f"/static/js/app.js?v={__version__}" in resp.text
    assert f"/static/favicon.svg?v={__version__}" in resp.text
    assert 'aria-label="Основная навигация"' in resp.text
    assert 'href="/generate"' in resp.text
    assert 'href="/reports"' in resp.text
    assert "Библиотека источников" in resp.text
    assert "Автоматизировать обработку" in resp.text
    assert 'data-task-center-open' in resp.text
    assert '/static/js/components/task-center.js' in resp.text
    assert '/static/js/i18n.js' in resp.text
    assert 'data-i18n="nav.convert"' in resp.text
    assert "tracked changes" not in resp.text
    assert "payload" not in resp.text


def test_global_task_center_reports_local_storage_policy():
    store = _client_task_store()
    store.set(
        "task-center-test",
        {
            "status": "done",
            "source_format": "docx",
            "target_format": "pdf",
            "filename": "result.pdf",
            "artifact": "result.pdf",
        },
    )

    response = client.get("/api/tasks")

    assert response.status_code == 200
    payload = response.json()
    assert payload["active"] == 0
    assert payload["retention_seconds"] == 3600
    assert payload["storage"]["scope"] == "local-device"
    assert payload["storage"]["root_label"] == "каталог данных TextAlchemy"
    assert isinstance(payload["storage"]["bytes"], int)
    assert payload["storage"]["automatic_cleanup"] is True
    assert payload["storage"]["external_uploads"] is False
    assert payload["tasks"][0]["result_url"] == "/api/convert/result/task-center-test"


def test_global_task_center_clears_only_finished_tasks():
    store = _client_task_store()
    store.set("finished-task", {"status": "done"})
    store.set("active-task", {"status": "running"})

    response = client.delete("/api/tasks/finished")

    assert response.json() == {"success": True, "removed": 1}
    assert store.get("finished-task") is None
    assert store.get("active-task") is not None


def test_global_task_center_can_rerun_individual_conversion(monkeypatch):
    web_app = importlib.import_module("textalchemy.web.app")

    store = _client_task_store()
    task_id = "rerun-single"
    source = Path(store.root) / task_id / "source"
    source.mkdir(parents=True)
    (source / "input.pdf").write_bytes(b"%PDF")
    store.set(
        task_id,
        {"status": "error", "queue_kind": "convert", "source_format": "pdf", "target_format": "docx", "mode": "balanced"},
    )
    launched = []
    monkeypatch.setattr(web_app.task_queue, "submit_named", lambda name, fn, *args: launched.append((name, fn, args)))

    snapshot = client.get("/api/tasks").json()["tasks"][0]
    response = client.post(f"/api/tasks/{task_id}/rerun")

    assert snapshot["can_rerun"] is True
    assert response.status_code == 200
    assert response.json()["status"] == "queued"
    assert launched and launched[0][0] == task_id


def test_global_task_center_can_cancel_running_conversion(monkeypatch):
    web_app = importlib.import_module("textalchemy.web.app")

    store = _client_task_store()
    task_id = "cancel-single"
    store.set(task_id, {"status": "running", "queue_kind": "convert"})
    monkeypatch.setattr(web_app.task_queue, "cancel", lambda name: name == task_id)

    response = client.post(f"/api/tasks/{task_id}/cancel")

    assert response.status_code == 200
    assert response.json()["status"] == "cancelled"
    assert store.get(task_id)["status"] == "cancelled"


def test_extract_page():
    resp = client.get("/extract")
    assert resp.status_code == 200
    assert f'/static/js/pages/extract.js?v={__version__}' in resp.text
    assert "Исходник LaTeX" in resp.text


def test_convert_page():
    resp = client.get("/convert")
    assert resp.status_code == 200
    assert "Максимально похожий вид" in resp.text
    assert "Можно удобно редактировать" in resp.text
    assert "Визуальный preview" not in resp.text
    assert "Preview, diff" not in resp.text
    assert "Скачать результат" in resp.text
    assert 'id="sourceInspection"' in resp.text
    assert 'id="comparisonSection"' in resp.text
    assert 'id="objectDiff"' in resp.text
    assert f'/static/js/pages/convert.js?v={__version__}' in resp.text
    assert 'type="module"' in resp.text
    assert 'convert-history.js' not in resp.text
    assert "async function pollTask" not in resp.text


def test_organize_page():
    resp = client.get("/pipeline")
    assert resp.status_code == 200
    assert "Соберите сценарий обработки" in resp.text
    assert 'id="opsPalette"' in resp.text
    assert 'id="expertMode"' in resp.text
    assert f'/static/js/pages/pipeline.js?v={__version__}' in resp.text
    assert 'type="module"' in resp.text
    assert "function renderPalette" not in resp.text


def test_api_operations_includes_param_schema():
    resp = client.get("/api/operations")
    assert resp.status_code == 200
    data = resp.json()
    assert isinstance(data, list)
    ops = {op["id"]: op for op in data}
    assert "ingest.file" in ops
    op = ops["ingest.file"]
    assert op["input_param"] == "path"
    assert "params" in op
    assert op["params"]["path"]["required"] is True
    assert ops["extract.html_model"]["input_type"] == "Document"
    assert ops["extract.html_model"]["output_type"] == "DocumentModel"
    for renderer in ("bibtex", "gost", "markdown", "json"):
        operation = ops[f"render.{renderer}"]
        assert operation["input_param"] == "items"
        assert operation["params"]["items"]["required"] is True


@pytest.mark.parametrize("renderer", ["bibtex", "gost", "markdown", "json"])
def test_api_pipeline_bibliography_linked_input(renderer):
    spec = {
        "bib": [{"index": 1, "authors": ["Иванов"], "title": "Исследование", "year": 2020}],
        "steps": [{"op": f"render.{renderer}", "input": "bib", "output": "rendered"}],
        "output": "rendered",
    }
    response = client.post("/api/pipeline/run", data={"spec": json.dumps(spec)})
    assert response.status_code == 200
    payload = response.json()
    assert payload["success"] is True
    assert payload["result"]["ok"] is True, payload["result"]["error"]
    assert "Исследование" in payload["result"]["final"]
    assert "Иванов" in payload["result"]["final"]


def test_api_pipeline_parse_normalizes_spec():
    spec = json.dumps({
        "path": "input.txt",
        "steps": [
            {"op": "ingest.file", "output": "doc", "params": {"path": "input.txt"}},
            {"op": "extract.text", "input": "doc", "output": "text"},
        ],
        "output": "text",
    })
    resp = client.post("/api/pipeline/parse", data={"spec": spec})
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert [s["op"] for s in data["spec"]["steps"]] == ["ingest.file", "extract.text"]
    assert data["spec"]["steps"][1]["input"] == "doc"
    assert data["spec"]["ctx"] == {"path": "input.txt"}
    assert data["spec"]["output"] == "text"


def test_api_pipeline_parse_rejects_bad_spec():
    resp = client.post("/api/pipeline/parse", data={"spec": "steps: [некорректно"})
    assert resp.status_code == 200
    assert resp.json()["success"] is False


def test_api_pipeline_yaml_serializes():
    spec = json.dumps({"steps": [{"op": "ingest.file", "output": "doc"}]})
    resp = client.post("/api/pipeline/yaml", data={"spec": spec})
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert "ingest.file" in data["yaml"]


def test_api_pipeline_validate_reports_broken_links():
    spec = json.dumps({
        "steps": [
            {"op": "ingest.file", "output": "doc", "params": {"path": "input.txt"}},
            {"op": "extract.text", "input": "missing", "output": "text"},
        ],
    })
    resp = client.post("/api/pipeline/validate", data={"spec": spec})
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is False
    assert any("missing" in e["message"] for e in data["errors"])


def test_api_pipeline_validate_ok():
    spec = json.dumps({
        "steps": [
            {"op": "ingest.file", "output": "doc", "params": {"path": "input.txt"}},
            {"op": "extract.text", "input": "doc", "output": "text"},
        ],
        "output": "text",
    })
    resp = client.post("/api/pipeline/validate", data={"spec": spec})
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["errors"] == []


def test_api_pipeline_validate_rejects_self_reference():
    spec = {'steps': [{'op': 'extract.text', 'input': 'text', 'output': 'text'}]}
    data = client.post('/api/pipeline/validate', data={'spec': json.dumps(spec)}).json()
    assert data['success'] is False
    assert data['errors'][0]['step'] == 0
    # Replacing a value already supplied in the initial context is valid.
    spec['text'] = 'initial'
    data = client.post('/api/pipeline/validate', data={'spec': json.dumps(spec)}).json()
    assert data['success'] is True


def test_api_pipeline_run_reports_step_failure():
    spec = json.dumps({
        "path": "no-such-file-12345.txt",
        "steps": [{"op": "ingest.file", "output": "doc", "params": {"path": "no-such-file-12345.txt"}}],
    })
    resp = client.post("/api/pipeline/run", data={"spec": spec})
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["result"]["ok"] is False


def test_pipeline_builder_end_to_end_flow(tmp_path):
    """Полный пользовательский поток конструктора через те же API, что и фронтенд."""
    src = tmp_path / "input.txt"
    src.write_text("e2e", encoding="utf-8")

    operations = {op["id"]: op for op in client.get("/api/operations").json()}
    assert "ingest.file" in operations
    assert "extract.text" in operations

    spec = json.dumps({
        "path": str(src),
        "steps": [
            {"op": "ingest.file", "output": "doc", "params": {"path": str(src)}},
            {"op": "extract.text", "input": "doc", "output": "text"},
        ],
        "output": "text",
    })

    validated = client.post("/api/pipeline/validate", data={"spec": spec}).json()
    assert validated["success"] is True

    yaml_resp = client.post("/api/pipeline/yaml", data={"spec": spec}).json()
    assert yaml_resp["success"] is True
    assert "ingest.file" in yaml_resp["yaml"]

    parsed = client.post("/api/pipeline/parse", data={"spec": yaml_resp["yaml"]}).json()
    assert parsed["success"] is True
    assert [s["op"] for s in parsed["spec"]["steps"]] == ["ingest.file", "extract.text"]

    ran = client.post("/api/pipeline/run", data={"spec": yaml_resp["yaml"]}).json()
    assert ran["success"] is True
    assert ran["result"]["ok"] is True
    assert ran["result"]["final"]["plain"] == "e2e"


_PAGES = ["/", "/extract", "/convert", "/pipeline", "/bibliography", "/matching",
          "/reports", "/recognize", "/generate", "/export"]

_LABELABLE = {"input", "select", "textarea"}
_EXEMPT_CONTROL_TYPES = {"hidden", "submit", "reset", "button", "image", "file"}


class _A11yScanner(HTMLParser):
    def __init__(self):
        super().__init__()
        self.viewport = False
        self.h1 = False
        self.label_for = []
        self.controls = []  # (tag, attrs, in_wrapping_label)
        self.images = []
        self.dropzones = []
        self.scripts = []
        self._label_for = None
        self._in_wrapping_label = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "meta" and attrs.get("name") == "viewport":
            self.viewport = True
        elif tag == "h1":
            self.h1 = True
        elif tag == "label":
            self._label_for = attrs.get("for")
            self._in_wrapping_label = not self._label_for
            if self._label_for:
                self.label_for.append(self._label_for)
        elif tag == "img":
            self.images.append(attrs)
        elif tag == "script" and attrs.get("src"):
            self.scripts.append(attrs["src"])
        if tag == "div" and "file-drop" in (attrs.get("class") or "").split():
            self.dropzones.append(attrs)
        if tag in _LABELABLE:
            self.controls.append((tag, attrs, self._in_wrapping_label))

    def handle_endtag(self, tag):
        if tag == "label":
            self._label_for = None
            self._in_wrapping_label = False


def _scan_page(path):
    resp = client.get(path)
    assert resp.status_code == 200
    scanner = _A11yScanner()
    scanner.feed(resp.text)
    scanner.html = resp.text
    return scanner


def test_pages_are_mobile_ready():
    for path in _PAGES:
        assert _scan_page(path).viewport, f"{path} без <meta name='viewport'>"


def test_pages_have_main_heading():
    for path in _PAGES:
        assert _scan_page(path).h1, f"{path} без <h1>"


def test_pages_form_controls_have_labels():
    for path in _PAGES:
        scanner = _scan_page(path)
        labelled = set(scanner.label_for)
        for tag, attrs, in_wrapping_label in scanner.controls:
            if attrs.get("type") in _EXEMPT_CONTROL_TYPES:
                continue
            if attrs.get("aria-label") or attrs.get("aria-labelledby"):
                continue
            if in_wrapping_label:
                continue
            control_id = attrs.get("id")
            assert control_id and control_id in labelled, (
                f"{path}: <{tag} id={control_id!r}> не имеет связанной <label for> или aria-label"
            )


def test_pages_images_have_alt():
    for path in _PAGES:
        scanner = _scan_page(path)
        for attrs in scanner.images:
            assert "alt" in attrs, f"{path}: <img> без атрибута alt"


def test_dropzones_are_keyboard_operable():
    for path in _PAGES:
        scanner = _scan_page(path)
        for attrs in scanner.dropzones:
            assert attrs.get("role") == "button", f"{path}: file-drop без role='button'"
            assert attrs.get("tabindex") == "0", f"{path}: file-drop без tabindex='0'"
        if scanner.dropzones:
            scripts = "\n".join(client.get(source).text for source in scanner.scripts)
            if "components/ingest-input.js" in scripts:
                scripts += client.get("/static/js/components/ingest-input.js").text
            assert "keydown" in scanner.html + scripts, f"{path}: file-drop без обработчика keydown (Enter/Space)"


def test_bibliography_page():
    resp = client.get("/bibliography")
    assert resp.status_code == 200
    assert 'aria-label="Разделы библиотеки"' in resp.text
    assert f'/static/js/pages/bibliography.js?v={__version__}' in resp.text


def test_rename_page():
    resp = client.get("/matching")
    assert resp.status_code == 200


def test_matching_page():
    resp = client.get("/matching")
    assert resp.status_code == 200
    assert f'/static/js/pages/matching.js?v={__version__}' in resp.text


def test_reports_page():
    resp = client.get("/reports")
    assert resp.status_code == 200
    assert f'/static/js/pages/reports.js?v={__version__}' in resp.text


def test_export_page():
    resp = client.get("/export")
    assert resp.status_code == 200
    assert f'/static/js/pages/export.js?v={__version__}' in resp.text


def test_recognize_page():
    resp = client.get("/recognize")
    assert resp.status_code == 200


def test_generate_page():
    resp = client.get("/generate")
    assert resp.status_code == 200


def test_api_info():
    resp = client.get("/api/info")
    assert resp.status_code == 200
    data = resp.json()
    assert "name" in data
    assert "version" in data
    assert "modules" in data


def test_legacy_bibliography_migration_reports_invalid_json(monkeypatch, tmp_path, caplog):
    web_app = importlib.import_module("textalchemy.web.app")
    legacy = tmp_path / "bibliography.json"
    legacy.write_text("{broken", encoding="utf-8")
    database = Database(tmp_path / "migration.db")
    monkeypatch.setattr(web_app, "data_dir", tmp_path)
    monkeypatch.setattr(web_app, "db", database)

    with caplog.at_level("ERROR", logger="textalchemy.web.app"):
        web_app._migrate_json_to_db()

    assert legacy.exists()
    assert database.all_items() == []
    assert "Failed to migrate legacy bibliography" in caplog.text
    database.engine.dispose()


def test_legacy_bibliography_migration_is_successful(monkeypatch, tmp_path):
    web_app = importlib.import_module("textalchemy.web.app")
    legacy = tmp_path / "bibliography.json"
    legacy.write_text('[{"title": "Migrated"}]', encoding="utf-8")
    database = Database(tmp_path / "migration.db")
    monkeypatch.setattr(web_app, "data_dir", tmp_path)
    monkeypatch.setattr(web_app, "db", database)

    web_app._migrate_json_to_db()

    assert not legacy.exists()
    assert (tmp_path / "bibliography.json.imported").exists()
    assert [item.title for item in database.all_items()] == ["Migrated"]
    database.engine.dispose()


def test_api_bibliography():
    resp = client.get("/api/bibliography")
    assert resp.status_code == 200
    assert isinstance(resp.json(), list)


def test_api_add_bibliography():
    resp = client.post("/api/bibliography", data={
        "authors": "Ivanov I.I.",
        "title": "Test Article",
        "doc_type": "article",
        "year": "2024",
    })
    assert resp.status_code == 200
    data = resp.json()
    assert data["success"] is True
    assert data["item"]["title"] == "Test Article"


def test_api_config():
    resp = client.get("/api/config")
    assert resp.status_code == 200
    cfg = resp.json()
    assert "source_dir" in cfg


def test_api_stats():
    resp = client.get("/api/stats")
    assert resp.status_code == 200
    data = resp.json()
    assert "total_bib" in data
    assert "total_files" in data


def test_api_export_json():
    resp = client.get("/api/export/json")
    assert resp.status_code == 200
    assert "application/json" in resp.headers["content-type"]


def test_api_export_markdown():
    resp = client.get("/api/export/markdown")
    assert resp.status_code == 200
    assert "text/markdown" in resp.headers["content-type"]


def test_api_export_gost():
    resp = client.get("/api/export/gost")
    assert resp.status_code == 200
    assert "text/plain" in resp.headers["content-type"]


def test_api_export_ris():
    resp = client.get("/api/export/ris")
    assert resp.status_code == 200


def test_api_export_csv():
    resp = client.get("/api/export/csv")
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]


def test_api_export_bibtex():
    resp = client.get("/api/export/bibtex")
    assert resp.status_code == 200


def test_api_export_unknown():
    resp = client.get("/api/export/unknown")
    assert resp.status_code == 400


def test_api_templates():
    resp = client.get("/api/generate/templates")
    assert resp.status_code == 200
    assert isinstance(resp.json(), list)


# ── Bibliography CRUD ──────────────────────────────

def test_api_bibliography_update():
    resp = client.put("/api/bibliography/1", data={
        "authors": "Updated Author",
        "title": "Updated Title",
        "doc_type": "article",
    })
    assert resp.status_code in (200, 404)
    if resp.status_code == 200:
        assert resp.json()["success"] is True


def test_api_bibliography_update_not_found():
    resp = client.put("/api/bibliography/9999", data={
        "authors": "None",
        "title": "None",
        "doc_type": "article",
    })
    assert resp.status_code == 404


def test_api_bibliography_delete():
    resp = client.delete("/api/bibliography/1")
    assert resp.status_code == 200


# ── Matching ───────────────────────────────────────

def test_api_match_run():
    resp = client.post("/api/match/run", data={
        "source_dir": str(Path.cwd() / "tests"),
        "threshold": "0.30",
    })
    assert resp.status_code == 200
    data = resp.json()
    assert "matched" in data
    assert "unmatched" in data


def test_api_match_report():
    resp = client.get("/api/match/report")
    assert resp.status_code == 200
    data = resp.json()
    assert "matched" in data


# ── Config ─────────────────────────────────────────

def test_api_update_config():
    resp = client.post("/api/config", data={"cfg": '{"source_dir": "./docs"}'})
    assert resp.status_code == 200
    assert resp.json()["success"] is True


# ── Convert ────────────────────────────────────────

def test_api_convert_no_file():
    resp = client.post("/api/convert")
    assert resp.status_code in (400, 422)


def test_api_convert_inspect_no_file():
    response = client.post("/api/convert/inspect")
    assert response.status_code in (400, 422)


def test_api_convert_status_not_found():
    resp = client.get("/api/convert/status/nonexistent")
    assert resp.status_code == 404


def test_api_convert_inspect_returns_public_structure_and_cleans_workspace(monkeypatch, tmp_path):
    def fake_inspect(path):
        return DocumentInspection(
            path,
            "pptx",
            metrics={"pages": 2, "paragraphs": 5, "tables": 1, "images": 2},
        )

    monkeypatch.setattr("textalchemy.web.routes.convert.inspect_path", fake_inspect)
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    response = client.post(
        "/api/convert/inspect",
        files={"file": (r"..\deck.pptx", b"pptx", "application/octet-stream")},
    )

    assert response.status_code == 200
    inspection = response.json()["inspection"]
    assert inspection["source_path"] == "deck.pptx"
    assert inspection["source_format"] == "pptx"
    assert inspection["metrics"]["tables"] == 1
    assert list(tmp_path.iterdir()) == []


def test_api_convert_capabilities_are_runtime_plans(monkeypatch):
    monkeypatch.setattr("textalchemy.convert.executor.requirement_available", lambda _requirement: True)
    response = client.get("/api/convert/capabilities")
    assert response.status_code == 200
    sources = {source["format"]: source for source in response.json()["sources"]}
    assert {"pdf", "docx", "pptx", "model"} <= sources.keys()

    pdf_targets = {target["format"]: target for target in sources["pdf"]["targets"]}
    assert "docx" in pdf_targets
    assert pdf_targets["html"]["plans"]["balanced"]["steps"] == ["pdf.model", "model.html"]
    # Every intermediate is an in-memory model, never a path-only file conversion.
    for target in pdf_targets.values():
        for plan in target["plans"].values():
            assert all(step.endswith(".model") for step in plan["steps"][:-1])
    assert set(pdf_targets["docx"]["modes"]) == {"balanced", "faithful", "editable"}

    docx_targets = {target["format"]: target for target in sources["docx"]["targets"]}
    assert {"pdf", "html", "latex", "model"} <= docx_targets.keys()
    assert docx_targets["pdf"]["plans"]["faithful"]["steps"] == ["docx.model", "model.pdf"]
    faithful = docx_targets["pdf"]["plans"]["faithful"]
    assert 0 < faithful["visual_score"] < 1
    assert faithful["editability_score"] < faithful["visual_score"]
    assert set(faithful["preservation"]["scores"]) == {
        "content", "semantics", "geometry", "style", "relationships", "editability",
    }


def test_api_convert_blocks_route_below_selected_loss_budget(monkeypatch):
    monkeypatch.setattr("textalchemy.convert.executor.requirement_available", lambda _requirement: True)

    response = client.post(
        "/api/convert",
        files={"file": ("budget.docx", b"placeholder", "application/octet-stream")},
        data={"target_format": "pdf", "mode": "faithful", "min_retention": "0.9"},
    )

    assert response.status_code == 422
    assert "ниже выбранного порога 90%" in response.json()["detail"]


def test_api_convert_txt_to_native_pptx():
    from io import BytesIO

    from pptx import Presentation

    started = client.post(
        "/api/convert", files={"file": ("slide.txt", "Редактируемый слайд".encode(), "text/plain")},
        data={"source_format": "txt", "target_format": "pptx", "mode": "editable"},
    )
    assert started.status_code == 200
    _wait_for_convert_tasks()
    status = client.get(started.json()["status"]).json()
    assert status["status"] == "done"
    result = client.get(started.json()["result"])
    assert result.status_code == 200
    assert "presentationml.presentation" in result.headers["content-type"]
    assert "slide.pptx" in result.headers["content-disposition"]
    assert Presentation(BytesIO(result.content)).slides[0].shapes[0].text == "Редактируемый слайд"


def test_api_convert_returns_report_and_separate_artifact(monkeypatch):
    captured = {}

    def fake_execute(_executor, request):
        captured["request"] = request
        request.output_path.write_bytes(b"converted-document")
        report = ConversionReport(request.output_path)
        report.metrics["executed_steps"] = ["pdf.docx.test"]
        return report

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    started = client.post(
        "/api/convert",
        files={"file": ("article.pdf", b"%PDF-test", "application/pdf")},
        data={"source_format": "pdf", "target_format": "docx", "mode": "editable"},
    )
    assert started.status_code == 200
    endpoints = started.json()
    _wait_for_convert_tasks()

    status = client.get(endpoints["status"])
    assert status.status_code == 200
    payload = status.json()
    assert payload["status"] == "done"
    assert payload["report"]["success"] is True
    assert payload["report"]["metrics"]["executed_steps"] == ["pdf.docx.test"]
    assert "content" not in payload

    request = captured["request"]
    assert request.source is DocFormat.PDF
    assert request.target is DocFormat.DOCX
    assert request.mode is ConversionMode.EDITABLE

    result = client.get(endpoints["result"])
    assert result.status_code == 200
    assert result.content == b"converted-document"
    assert "article.docx" in result.headers["content-disposition"]


def test_api_convert_compares_source_and_result_structure(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted-pdf")
        return ConversionReport(request.output_path)

    def fake_inspect(path):
        is_source = path.suffix == ".docx"
        return DocumentInspection(
            path,
            path.suffix.removeprefix("."),
            metrics={
                "pages": 1,
                "paragraphs": 2,
                "characters": 100 if is_source else 80,
                "tables": 1 if is_source else 0,
            },
            pages=[
                {
                    "index": 0,
                    "width_pt": 595.28,
                    "height_pt": 841.89,
                    "margin_top_pt": 72,
                    "margin_right_pt": 72,
                    "margin_bottom_pt": 72,
                    "margin_left_pt": 72,
                }
            ],
        )

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    monkeypatch.setattr("textalchemy.web.routes.convert.inspect_path", fake_inspect)
    monkeypatch.setattr("textalchemy.convert.executor.requirement_available", lambda _requirement: True)
    response = client.post(
        "/api/convert",
        files={"file": ("source.docx", b"docx", "application/octet-stream")},
        data={"source_format": "docx", "target_format": "pdf"},
    )

    assert response.status_code == 200, response.text
    _wait_for_convert_tasks()
    status = client.get(response.json()["status"]).json()
    assert status["source_inspection"]["source_path"] == "source.docx"
    assert status["target_inspection"]["source_path"] == "source.pdf"
    assert status["comparison"]["retention"]["characters"]["ratio"] == 0.8
    assert status["comparison"]["retention"]["tables"]["ratio"] == 0
    assert status["comparison"]["geometry_summary"]["max_dimension_error_pt"] is None
    assert status["comparison"]["geometry_summary"]["available"] is False


def test_api_convert_rejects_unknown_mode():
    response = client.post(
        "/api/convert",
        files={"file": ("article.pdf", b"%PDF-test", "application/pdf")},
        data={"source_format": "pdf", "target_format": "docx", "mode": "impossible"},
    )
    assert response.status_code == 400


def test_api_convert_rejects_oversized_upload_and_cleans_workspace(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path, max_bytes=5),
    )

    response = client.post(
        "/api/convert",
        files={"file": ("article.pdf", b"123456", "application/pdf")},
        data={"source_format": "pdf", "target_format": "docx"},
    )

    assert response.status_code == 413
    assert list(tmp_path.iterdir()) == []


def test_api_convert_sanitizes_uploaded_path(monkeypatch, tmp_path):
    captured = {}

    def fake_execute(_executor, request):
        captured["source"] = request.input_path
        request.output_path.write_bytes(b"converted")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    response = client.post(
        "/api/convert",
        files={"file": (r"..\..\article.pdf", b"%PDF-test", "application/pdf")},
        data={"source_format": "pdf", "target_format": "docx"},
    )

    assert response.status_code == 200
    _wait_for_convert_tasks()
    assert captured["source"].name == "article.pdf"
    assert list(tmp_path.iterdir()) == []


def test_api_convert_legacy_format_remains_supported(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"legacy-result")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    response = client.post(
        "/api/convert",
        files={"file": ("article.pdf", b"%PDF-test", "application/pdf")},
        data={"fmt": "pdf"},
    )
    assert response.status_code == 200
    _wait_for_convert_tasks()
    assert client.get(response.json()["status"]).json()["status"] == "done"


def test_api_convert_serves_single_file_html(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_text("<html>converted</html>", encoding="utf-8")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    response = client.post(
        "/api/convert",
        files={"file": ("article.docx", b"docx-test", "application/octet-stream")},
        data={"source_format": "docx", "target_format": "html", "mode": "balanced"},
    )
    endpoints = response.json()
    _wait_for_convert_tasks()
    result = client.get(endpoints["result"])
    assert result.status_code == 200
    assert result.headers["content-type"].startswith("text/html")
    assert result.content == b"<html>converted</html>"


# ── Convert: visual preview ──────────────────────────

def _seed_preview_task(store, task_id="preview-task"):
    task_dir = store.root / task_id
    source_dir = task_dir / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    (source_dir / "in.pdf").write_bytes(b"%PDF-1.4 preview source")
    (task_dir / "out.pdf").write_bytes(b"%PDF-1.4 preview result")
    store.set(task_id, {"status": "done", "artifact": "out.pdf", "filename": "out.pdf"})
    return task_id


def test_api_convert_single_preserves_source_for_preview(monkeypatch):
    from textalchemy.web.routes import convert as convert_route

    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted-document")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    response = client.post(
        "/api/convert",
        files={"file": ("article.pdf", b"%PDF-test", "application/pdf")},
        data={"source_format": "pdf", "target_format": "docx"},
    )
    assert response.status_code == 200
    _wait_for_convert_tasks()
    assert convert_route.tasks_store.source_path(response.json()["task_id"]) is not None


def test_api_convert_preview_meta_and_page(monkeypatch):
    store = _client_task_store()
    task_id = _seed_preview_task(store)

    def fake_count(_preview_dir, _file, _side):
        return 3

    def fake_png(_preview_dir, _file, _side, page_index, dpi=110):
        return b"\x89PNG\r\n\x1a\npreview-page"

    monkeypatch.setattr("textalchemy.web.routes.convert.cached_page_count", fake_count)
    monkeypatch.setattr("textalchemy.web.routes.convert.cached_page_png", fake_png)

    meta = client.get(f"/api/convert/preview/{task_id}/meta")
    assert meta.status_code == 200
    body = meta.json()
    assert body["source"]["available"] is True
    assert body["source"]["pages"] == 3
    assert body["target"]["pages"] == 3
    assert body["target"]["error"] is None

    page = client.get(f"/api/convert/preview/{task_id}?side=target&page=2")
    assert page.status_code == 200
    assert page.headers["content-type"] == "image/png"
    assert page.content == b"\x89PNG\r\n\x1a\npreview-page"
    assert "max-age" in page.headers.get("cache-control", "")

    assert client.get(f"/api/convert/preview/{task_id}?side=target&page=9").status_code == 404
    assert client.get(f"/api/convert/preview/{task_id}?side=other").status_code == 400
    assert client.get(f"/api/convert/preview/{task_id}?side=target&page=0").status_code == 400
    assert client.get("/api/convert/preview/nonexistent/meta").status_code == 404


def test_api_convert_preview_diff_returns_heatmap_and_metrics(monkeypatch):
    from io import BytesIO

    from PIL import Image, ImageDraw

    store = _client_task_store()
    task_id = _seed_preview_task(store, "diff-preview")

    def png(offset):
        image = Image.new("RGB", (80, 100), "white")
        ImageDraw.Draw(image).rectangle((10 + offset, 10, 50 + offset, 40), fill="black")
        output = BytesIO()
        image.save(output, "PNG")
        return output.getvalue()

    def fake_png(_preview_dir, _file, side, _page_index, dpi=110):
        return png(0 if side == "source" else 5)

    monkeypatch.setattr("textalchemy.web.routes.convert.cached_page_png", fake_png)

    response = client.get(f"/api/convert/preview/{task_id}/diff?page=1")

    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert 0 < float(response.headers["x-visual-similarity"]) < 1
    assert float(response.headers["x-visual-rmse"]) > 0
    assert response.content.startswith(b"\x89PNG")


def test_api_convert_preview_requires_done_task(monkeypatch):
    store = _client_task_store()
    task_id = "pending-preview"
    store.set(task_id, {"status": "running"})

    meta = client.get(f"/api/convert/preview/{task_id}/meta")
    assert meta.status_code == 409
    page = client.get(f"/api/convert/preview/{task_id}?side=source")
    assert page.status_code == 409


def test_api_convert_preview_target_unavailable_for_directory_artifacts(monkeypatch):
    store = _client_task_store()
    task_id = "dir-artifact"
    task_dir = store.root / task_id
    task_dir.mkdir(parents=True, exist_ok=True)
    source_dir = task_dir / "source"
    source_dir.mkdir(exist_ok=True)
    (source_dir / "in.pdf").write_bytes(b"%PDF")
    (task_dir / "deck-html.zip").write_bytes(b"zip")
    store.set(task_id, {"status": "done", "artifact": "deck-html.zip", "filename": "deck-html.zip"})

    def fake_count(_preview_dir, file, _side):
        return 3 if file.suffix == ".pdf" else 0

    monkeypatch.setattr("textalchemy.web.routes.convert.cached_page_count", fake_count)

    meta = client.get(f"/api/convert/preview/{task_id}/meta").json()
    assert meta["source"]["available"] is True
    assert meta["source"]["pages"] == 3
    assert meta["target"]["available"] is False
    assert meta["target"]["pages"] == 0


# ── Convert: batch queue + history ──────────────────

def _batch_files():
    return [
        ("files", ("a.pdf", b"%PDF-a", "application/pdf")),
        ("files", ("b.pdf", b"%PDF-b", "application/pdf")),
    ]


def test_api_convert_batch_runs_each_file_and_cleans(monkeypatch, tmp_path):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted-" + request.input_path.name.encode())
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    response = client.post("/api/convert/batch", files=_batch_files(), data={"target_format": "docx"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["success"] is True
    assert len(body["tasks"]) == 2
    assert {task["name"] for task in body["tasks"]} == {"a.pdf", "b.pdf"}
    assert {task["source_format"] for task in body["tasks"]} == {"pdf"}
    _wait_for_convert_tasks()

    job = client.get(f"/api/convert/jobs/{body['job_id']}").json()
    assert [task["status"] for task in job["tasks"]] == ["done", "done"]
    for task in job["tasks"]:
        result = client.get(task["result_url"])
        assert result.status_code == 200
        assert result.content == b"converted-" + task["name"].encode()
    assert list(tmp_path.iterdir()) == []


def test_api_convert_batch_requires_files():
    response = client.post("/api/convert/batch", data={"target_format": "docx"})
    assert response.status_code in (400, 422)


def test_api_convert_batch_limits_file_count():
    files = [("files", (f"file{i}.pdf", b"%PDF", "application/pdf")) for i in range(21)]
    response = client.post("/api/convert/batch", files=files)
    assert response.status_code == 400
    assert "максимум" in response.json()["detail"]


def test_api_convert_batch_rejects_unavailable_target_and_cleans(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    response = client.post("/api/convert/batch", files=_batch_files(), data={"target_format": "epub"})
    assert response.status_code == 400
    assert "недоступен" in response.json()["detail"]
    assert list(tmp_path.iterdir()) == []


def test_api_convert_batch_cleans_all_workspaces_on_unexpected_preparation_error(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    calls = 0

    def fail_on_second(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("preparation failed")
        return DocFormat.PDF, DocFormat.DOCX

    monkeypatch.setattr("textalchemy.web.routes.convert._resolve_conversion", fail_on_second)

    with pytest.raises(RuntimeError, match="preparation failed"):
        client.post("/api/convert/batch", files=_batch_files(), data={"target_format": "docx"})
    assert list(tmp_path.iterdir()) == []


def test_api_convert_batch_cleans_workspaces_when_job_persistence_fails(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    monkeypatch.setattr(
        "textalchemy.web.routes.convert.tasks_store.set_job",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("job store failed")),
    )

    with pytest.raises(OSError, match="job store failed"):
        client.post("/api/convert/batch", files=_batch_files(), data={"target_format": "docx"})
    assert list(tmp_path.iterdir()) == []


def test_api_convert_batch_rejects_unknown_mode():
    response = client.post("/api/convert/batch", files=_batch_files(), data={"mode": "impossible"})
    assert response.status_code == 400


def test_api_convert_jobs_lists_history(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    created = client.post("/api/convert/batch", files=_batch_files()).json()
    _wait_for_convert_tasks()

    history = client.get("/api/convert/jobs").json()
    assert history["jobs"]
    latest = history["jobs"][0]
    assert latest["job_id"] == created["job_id"]
    assert latest["counts"]["done"] == 2
    assert latest["files"] == ["a.pdf", "b.pdf"]


def test_api_convert_job_not_found():
    assert client.get("/api/convert/jobs/nope").status_code == 404
    assert client.delete("/api/convert/jobs/nope").status_code == 404
    assert client.post("/api/convert/jobs/nope/rerun").status_code == 404


def test_api_convert_job_delete_removes_tasks(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    created = client.post("/api/convert/batch", files=_batch_files()).json()
    _wait_for_convert_tasks()

    response = client.delete(f"/api/convert/jobs/{created['job_id']}")
    assert response.status_code == 200
    assert client.get(f"/api/convert/jobs/{created['job_id']}").status_code == 404
    for task in created["tasks"]:
        assert client.get(task["status"]).status_code == 404


def test_api_convert_job_rerun_uses_stored_sources(monkeypatch):
    calls = []

    def fake_execute(_executor, request):
        assert request.quality_policy.max_loss_issues == 2
        assert request.object_loss_policy.max_lost_objects == 3
        assert request.text_preservation_policy is not None
        calls.append(request.input_path.name)
        request.output_path.write_bytes(b"converted")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    created = client.post("/api/convert/batch", files=_batch_files(), data={
        "max_loss_issues": 2, "max_lost_objects": 3, "require_unchanged_text": True,
    }).json()
    _wait_for_convert_tasks()
    assert len(calls) == 2

    rerun = client.post(f"/api/convert/jobs/{created['job_id']}/rerun")
    assert rerun.status_code == 200
    assert len(rerun.json()["launched"]) == 2
    _wait_for_convert_tasks()
    assert len(calls) == 4

    job = client.get(f"/api/convert/jobs/{created['job_id']}").json()
    assert [task["status"] for task in job["tasks"]] == ["done", "done"]


def test_completed_conversion_can_rerun_with_persisted_quality_policy():
    response = client.post(
        "/api/convert", files={"file": ("source.txt", "Первая\nВторая".encode(), "text/plain")},
        data={"target_format": "model", "max_loss_issues": 0, "max_lost_objects": 0, "require_unchanged_text": True},
    )
    assert response.status_code == 200
    created = response.json()
    _wait_for_convert_tasks()
    first = client.get(created["status"]).json()
    assert first["status"] == "done", first
    assert first["max_loss_issues"] == 0
    assert first["report"]["metrics"]["quality_gate"]["accepted"] is True
    first_result = client.get(created["result"]).json()

    rerun = client.post(f"/api/tasks/{created['task_id']}/rerun")
    assert rerun.status_code == 200, rerun.text
    _wait_for_convert_tasks()
    second = client.get(created["status"]).json()
    assert second["status"] == "done", second
    assert second["source_format"] == "txt"
    assert second["report"]["metrics"]["quality_gate"]["max_loss_issues"] == 0
    assert second["report"]["metrics"]["object_quality_gate"]["max_lost_objects"] == 0
    assert second["max_lost_objects"] == 0
    assert second["require_unchanged_text"] is True
    assert second["report"]["metrics"]["text_quality_gate"]["accepted"] is True
    assert client.get(created["result"]).json() == first_result


@pytest.mark.parametrize("batch", [False, True])
def test_text_flow_mode_survives_rerun(batch):
    endpoint = "/api/convert/batch" if batch else "/api/convert"
    field = "files" if batch else "file"
    response = client.post(endpoint, files={field: ("source.txt", b"First Second", "text/plain")},
                           data={"target_format": "model", "text_preservation": "flow", "max_text_edits": 2})
    assert response.status_code == 200
    created = response.json()
    tasks = created["tasks"] if batch else [created]
    rerun_url = f"/api/convert/jobs/{created['job_id']}/rerun" if batch else f"/api/tasks/{created['task_id']}/rerun"
    for attempt in range(2):
        if attempt:
            assert client.post(rerun_url).status_code == 200
        _wait_for_convert_tasks()
        for item in tasks:
            task = client.get(item["status"]).json()
            assert task["status"] == "done", task
            assert task["text_preservation"] == "flow"
            assert task["max_text_edits"] == 2
            assert task["report"]["metrics"]["text_quality_gate"]["max_text_edits"] == 2
            assert task["report"]["metrics"]["text_quality_gate"]["mode"] == "flow"


@pytest.mark.parametrize("endpoint, field", [("/api/convert", "file"), ("/api/convert/batch", "files")])
@pytest.mark.parametrize("params", [
    {"text_preservation": "invalid"}, {"text_preservation": "flow", "require_unchanged_text": True},
])
def test_web_rejects_conflicting_or_unknown_text_mode(endpoint, field, params):
    response = client.post(endpoint, files={field: ("source.txt", b"Text", "text/plain")},
                           data={"target_format": "model", **params})
    assert response.status_code == 400


def test_unverifiable_object_budget_survives_web_rerun():
    response = client.post(
        "/api/convert", files={"file": ("source.txt", b"Unique text", "text/plain")},
        data={"target_format": "html", "max_lost_objects": 0},
    )
    assert response.status_code == 200
    created = response.json()
    for attempt in range(2):
        if attempt:
            assert client.post(f"/api/tasks/{created['task_id']}/rerun").status_code == 200
        _wait_for_convert_tasks()
        task = client.get(created["status"]).json()
        assert task["status"] == "error"
        assert task["max_lost_objects"] == 0
        assert task["report"]["metrics"]["object_quality_gate"]["reason"] == "unavailable"
        assert client.get(created["result"]).status_code != 200


@pytest.mark.parametrize("endpoint, field", [("/api/convert", "file"), ("/api/convert/batch", "files")])
@pytest.mark.parametrize("limit_field", ["max_loss_issues", "max_lost_objects", "max_text_edits"])
def test_web_conversion_rejects_negative_loss_budget(endpoint, field, limit_field):
    response = client.post(
        endpoint, files={field: ("source.txt", b"text", "text/plain")},
        data={"target_format": "model", limit_field: -1},
    )
    assert response.status_code == 422


def test_rejected_web_conversion_keeps_policy_for_rerun(monkeypatch):
    from textalchemy.core.diagnostics import IssueSeverity

    def lossy_export(model, output):
        output.write_text("<p>Incomplete</p>", encoding="utf-8")
        report = ConversionReport(output)
        report.add(IssueSeverity.LOSS, "text", "Text lost during export")
        return report

    monkeypatch.setattr("textalchemy.convert.executor._write_html", lossy_export)
    response = client.post(
        "/api/convert", files={"file": ("source.txt", b"text", "text/plain")},
        data={"target_format": "html", "max_loss_issues": 0},
    )
    assert response.status_code == 200
    created = response.json()
    for attempt in range(2):
        if attempt:
            assert client.post(f"/api/tasks/{created['task_id']}/rerun").status_code == 200
        _wait_for_convert_tasks()
        task = client.get(created["status"]).json()
        assert task["status"] == "error", task
        assert task["report"]["metrics"]["quality_gate"]["accepted"] is False
        assert task["max_loss_issues"] == 0
        assert client.get(created["result"]).status_code != 200
        stored = _client_task_store().get(created["task_id"])
        assert not stored.get("artifact")


def test_api_convert_jobs_report_interrupted_tasks(monkeypatch):
    def fake_execute(_executor, request):
        request.output_path.write_bytes(b"converted")
        return ConversionReport(request.output_path)

    monkeypatch.setattr("textalchemy.web.routes.convert.ConversionExecutor.execute", fake_execute)
    created = client.post("/api/convert/batch", files=_batch_files()).json()
    _wait_for_convert_tasks()

    store = _client_task_store()
    for task in created["tasks"]:
        task_id = task["task_id"]
        meta = store.get(task_id)
        store.set(task_id, {**meta, "status": "interrupted", "error": "Задача прервана перезапуском сервера"})

    job = client.get(f"/api/convert/jobs/{created['job_id']}").json()
    assert [entry["status"] for entry in job["tasks"]] == ["interrupted", "interrupted"]
    assert all("прервана" in entry["error"] for entry in job["tasks"])

    history = client.get("/api/convert/jobs").json()
    latest = history["jobs"][0]
    assert latest["counts"]["interrupted"] == 2


# ── Import ─────────────────────────────────────────

def test_api_import_no_file():
    resp = client.post("/api/bibliography/import")
    assert resp.status_code in (400, 422)


# ── Preview rename ─────────────────────────────────

def test_api_preview_rename():
    resp = client.post("/api/preview/rename", data={
        "source_dir": str(Path.cwd() / "tests"),
    })
    assert resp.status_code == 200


# ── Extract ────────────────────────────────────────

def test_api_extract_text_no_file():
    resp = client.post("/api/extract/text")
    assert resp.status_code in (400, 422)


def test_api_extract_latex_no_file():
    resp = client.post("/api/extract/latex")
    assert resp.status_code in (400, 422)


# ── Recognize ──────────────────────────────────────

def test_api_recognize_no_file():
    resp = client.post("/api/recognize")
    assert resp.status_code in (400, 422)


# ── Generate ───────────────────────────────────────

def test_api_generate_no_data():
    resp = client.post("/api/generate")
    assert resp.status_code in (400, 422)


# ── 404 handler ────────────────────────────────────

def test_404():
    resp = client.get("/nonexistent-route")
    assert resp.status_code == 404


# ── Extract: текст и LaTeX ─────────────────────────

def test_api_extract_text_returns_plain(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.extract.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    resp = client.post(
        "/api/extract/text",
        files={"file": ("notes.txt", "Привет из файла".encode("utf-8"), "text/plain")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert "Привет из файла" in body["text"]
    assert list(tmp_path.iterdir()) == []


def test_api_extract_text_error(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.extract.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    def boom(**kwargs):
        raise RuntimeError("reader exploded")

    monkeypatch.setattr("textalchemy.web.routes.extract.extract_text", boom)
    resp = client.post(
        "/api/extract/text",
        files={"file": ("notes.txt", b"x", "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json()["success"] is False
    assert "reader exploded" in resp.json()["error"]


def test_api_extract_latex_returns_tex(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.extract.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    def fake_latex(src, out, doc_type="manuscript"):
        Path(out).write_text("\\section{Раздел}", encoding="utf-8")

    monkeypatch.setattr("textalchemy.web.routes.extract.docx_to_latex", fake_latex)
    resp = client.post(
        "/api/extract/latex",
        files={"file": ("doc.docx", b"docx-bytes", "application/octet-stream")},
    )
    assert resp.status_code == 200
    assert "text/plain" in resp.headers["content-type"]
    assert "\\section{Раздел}" in resp.text
    assert list(tmp_path.iterdir()) == []


def test_api_extract_latex_error(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.extract.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    def boom(src, out, doc_type="manuscript"):
        raise RuntimeError("latex failed")

    monkeypatch.setattr("textalchemy.web.routes.extract.docx_to_latex", boom)
    resp = client.post(
        "/api/extract/latex",
        files={"file": ("doc.docx", b"docx-bytes", "application/octet-stream")},
    )
    assert resp.json()["success"] is False
    assert "latex failed" in resp.json()["error"]


# ── Recognize: OCR ─────────────────────────────────

def test_api_recognize_success(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    class FakeEngine:
        is_available = True

        def __init__(self, languages, use_gpu=False):
            self.languages = languages

        @property
        def backend_name(self):
            return "tesseract"

        def recognize(self, path, handwriting=False):
            return type("R", (), {"text": "распознано", "confidence": 0.9})()

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    resp = client.post(
        "/api/recognize",
        files={"file": ("page.png", b"png-bytes", "image/png")},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is True
    assert body["text"] == "распознано"
    assert body["backend"] == "tesseract"
    assert list(tmp_path.iterdir()) == []


def test_api_recognize_error_when_backend_unavailable(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    class FakeEngine:
        is_available = False

        def __init__(self, languages, use_gpu=False):
            pass

        @property
        def backend_name(self):
            return None

        def recognize(self, path, handwriting=False):
            raise AssertionError("не должен вызываться")

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    resp = client.post(
        "/api/recognize",
        files={"file": ("page.png", b"png-bytes", "image/png")},
    )
    body = resp.json()
    assert body["success"] is False
    assert "OCR-движок недоступен" in body["error"]
    assert "draft_id" not in body


def test_api_recognize_error(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    class FakeEngine:
        is_available = True

        def __init__(self, languages, use_gpu=False):
            pass

        @property
        def backend_name(self):
            return "easyocr"

        def recognize(self, path, handwriting=False):
            raise RuntimeError("ocr backend crashed")

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    resp = client.post(
        "/api/recognize",
        files={"file": ("page.png", b"png-bytes", "image/png")},
    )
    assert resp.json()["success"] is False
    assert "ocr backend crashed" in resp.json()["error"]


def test_api_recognize_pdf_scenario_fast(tmp_path, monkeypatch):
    """PDF with a native text layer is readable in fast mode without OCR."""
    import fitz

    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    pdf = tmp_path / "text_layer.pdf"
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.insert_text((36, 48), "Fast mode text layer", fontname="helv", fontsize=12)
    doc.save(str(pdf))
    doc.close()

    class FakeEngine:
        is_available = False

        def __init__(self, languages, use_gpu=False):
            pass

        @property
        def backend_name(self):
            return None

        def recognize_pdf_geometry(self, *args, **kwargs):
            raise AssertionError("OCR must not run in fast mode")

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    with open(pdf, "rb") as fh:
        resp = client.post(
            "/api/recognize",
            files={"file": ("scan.pdf", fh.read(), "application/pdf")},
            data={"scenario": "fast"},
        )
    body = resp.json()
    assert body["success"] is True
    assert "Fast mode text layer" in body["text"]
    assert body["backend"] == "text_layer"
    assert body["scenario"] == "fast"


def test_api_recognize_pdf_scan_error_when_no_backend(tmp_path, monkeypatch):
    """Missing OCR cannot produce a successful editable document."""
    import fitz

    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    pdf = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page(width=300, height=400)
    doc.save(str(pdf))
    doc.close()

    class FakeEngine:
        is_available = False

        def __init__(self, languages, use_gpu=False):
            pass

        @property
        def backend_name(self):
            return None

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    with open(pdf, "rb") as fh:
        resp = client.post(
            "/api/recognize",
            files={"file": ("scan.pdf", fh.read(), "application/pdf")},
            data={"scenario": "scan"},
        )
    body = resp.json()
    assert body["success"] is False
    assert "OCR-движок недоступен" in body["error"]
    assert "draft_id" not in body


def test_api_recognize_pdf_unknown_scenario(tmp_path, monkeypatch):
    """Unknown scenario returns an error payload."""
    import fitz

    from textalchemy.core.artifacts import ArtifactWorkspace

    monkeypatch.setattr(
        "textalchemy.web.routes.recognize.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )

    pdf = tmp_path / "x.pdf"
    doc = fitz.open()
    doc.new_page(width=300, height=400)
    doc.save(str(pdf))
    doc.close()

    class FakeEngine:
        is_available = False

        def __init__(self, languages, use_gpu=False):
            pass

    monkeypatch.setattr("textalchemy.web.routes.recognize.OcrEngine", FakeEngine)
    with open(pdf, "rb") as fh:
        resp = client.post(
            "/api/recognize",
            files={"file": ("x.pdf", fh.read(), "application/pdf")},
            data={"scenario": "bogus"},
        )
    body = resp.json()
    assert body["success"] is False
    assert "unknown scenario" in body["error"]


# ── Generate: шаблоны ──────────────────────────────

def test_api_generate_lists_templates(monkeypatch):
    from textalchemy.generate.template import DocumentTemplate

    monkeypatch.setattr('textalchemy.web.routes.generate.custom_templates', lambda: [])
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.list_templates",
        lambda: [DocumentTemplate(name="report", description="Отчёт")],
    )
    resp = client.get("/api/generate/templates")
    assert resp.status_code == 200
    assert resp.json() == [{"name": "report", "description": "Отчёт", "template_type": "docx"}]


def test_api_generate_template_schema(monkeypatch):
    from textalchemy.generate.template_schema import TemplateSchema

    monkeypatch.setattr(
        "textalchemy.web.routes.generate.template_schema_for",
        lambda name, templates_dir=None: (TemplateSchema(allow_extra=True), "derived"),
    )
    resp = client.get("/api/generate/templates/report/schema")
    assert resp.status_code == 200
    body = resp.json()
    assert body["name"] == "report"
    assert body["source"] == "derived"
    assert body["schema"]["allow_extra"] is True


def test_api_generate_template_schema_missing():
    resp = client.get("/api/generate/templates/does-not-exist/schema")
    assert resp.status_code == 404


def test_api_generate_template_preview_meta(monkeypatch):
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.cached_page_count",
        lambda preview_dir, source, side: 3,
    )
    resp = client.get("/api/generate/templates/report/preview/meta")
    assert resp.status_code == 200
    assert resp.json() == {"available": True, "pages": 3, "error": None}


def test_api_generate_template_preview_png(monkeypatch):
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.cached_page_png",
        lambda preview_dir, source, side, page, **kwargs: b"\x89PNG-preview",
    )
    resp = client.get("/api/generate/templates/report/preview?page=1")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "image/png"
    assert resp.content == b"\x89PNG-preview"


def test_api_generate_rejects_bad_json():
    resp = client.post(
        "/api/generate",
        data={"template": "report", "params": "{not json"},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is False
    assert "JSON" in body["error"]


def test_api_generate_validation_errors(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace
    from textalchemy.generate.template_schema import TemplateField, TemplateSchema, TemplateValueType

    monkeypatch.setattr(
        "textalchemy.web.routes.generate.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.template_schema_for",
        lambda name, templates_dir=None: (
            TemplateSchema(
                fields=[
                    TemplateField("title", TemplateValueType.STRING),
                    TemplateField("count", TemplateValueType.INTEGER),
                ],
                allow_extra=False,
            ),
            "sidecar",
        ),
    )
    resp = client.post(
        "/api/generate",
        data={"template": "report", "params": '{"count": "many"}'},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["success"] is False
    assert body["errors"]["title"] == "required value is missing"
    assert body["errors"]["count"] == "expected integer, got str"


def test_api_generate_returns_file(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace
    from textalchemy.core.diagnostics import ConversionReport
    from textalchemy.generate.template_schema import TemplateSchema

    monkeypatch.setattr(
        "textalchemy.web.routes.generate.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.template_schema_for",
        lambda name, templates_dir=None: (TemplateSchema(allow_extra=True), "derived"),
    )

    def fake_generate(template_path, output_path, data, *, strict=True, schema=None):
        Path(output_path).write_bytes(b"docx-content")
        return ConversionReport(Path(output_path))

    monkeypatch.setattr("textalchemy.web.routes.generate.generate_docx_template", fake_generate)
    resp = client.post(
        "/api/generate",
        data={"template": "report", "output": "out.docx", "params": '{"a": 1}'},
    )
    assert resp.status_code == 200
    assert "out.docx" in resp.headers["content-disposition"]
    assert resp.content == b"docx-content"


def test_api_generate_error(tmp_path, monkeypatch):
    from textalchemy.core.artifacts import ArtifactWorkspace
    from textalchemy.generate.template_schema import TemplateSchema

    monkeypatch.setattr(
        "textalchemy.web.routes.generate.create_web_workspace",
        lambda: ArtifactWorkspace(parent=tmp_path),
    )
    monkeypatch.setattr(
        "textalchemy.web.routes.generate.template_schema_for",
        lambda name, templates_dir=None: (TemplateSchema(allow_extra=True), "derived"),
    )

    def boom(template_path, output_path, data, *, strict=True, schema=None):
        raise RuntimeError("generation failed")

    monkeypatch.setattr("textalchemy.web.routes.generate.generate_docx_template", boom)
    resp = client.post(
        "/api/generate",
        data={"template": "report", "params": "{}"},
    )
    assert resp.json()["success"] is False
    assert "generation failed" in resp.json()["error"]
