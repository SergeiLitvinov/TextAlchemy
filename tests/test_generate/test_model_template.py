"""Тесты безопасного шаблонизатора DocumentModel."""

import pytest
from docx import Document

from textalchemy.core.document_model import (
    DocumentModel,
    Formula,
    FormulaFormat,
    Image,
    Paragraph,
    Section,
    Table,
    TableCell,
    TableRow,
    TextRun,
)
from textalchemy.core.exceptions import GenerateError
from textalchemy.generate import TemplateFormula, TemplateImage, generate_docx_template, render_document_template


def _paragraph(text: str) -> Paragraph:
    return Paragraph(content=[TextRun(text)])


def test_render_variables_conditions_and_loops_without_mutating_template():
    template = DocumentModel(
        metadata={"title": "Report: {{ title }}"},
        sections=[
            Section(
                blocks=[
                    _paragraph("Hello, {{ user.name }}!"),
                    _paragraph("{% if show_details %}"),
                    _paragraph("Details: {{ details }}"),
                    _paragraph("{% else %}"),
                    _paragraph("No details"),
                    _paragraph("{% endif %}"),
                    _paragraph("{% for item in items %}"),
                    _paragraph("- {{ item }}"),
                    _paragraph("{% endfor %}"),
                ]
            )
        ],
    )

    result = render_document_template(
        template,
        {
            "title": "July",
            "user": {"name": "Alice"},
            "show_details": True,
            "details": "Ready",
            "items": ["one", "two"],
        },
    )

    assert result.metadata["title"] == "Report: July"
    assert [block.plain_text for block in result.sections[0].blocks] == [
        "Hello, Alice!",
        "Details: Ready",
        "- one",
        "- two",
    ]
    assert template.sections[0].blocks[0].plain_text == "Hello, {{ user.name }}!"


def test_render_repeating_table_rows():
    marker_start = TableRow(cells=[TableCell(blocks=[_paragraph("{% for row in rows %}")])])
    body = TableRow(
        cells=[
            TableCell(blocks=[_paragraph("{{ row.name }}")]),
            TableCell(blocks=[_paragraph("{{ row.value }}")]),
        ]
    )
    marker_end = TableRow(cells=[TableCell(blocks=[_paragraph("{% endfor %}")])])
    template = DocumentModel(sections=[Section(blocks=[Table(rows=[marker_start, body, marker_end])])])

    result = render_document_template(template, {"rows": [{"name": "A", "value": 1}, {"name": "B", "value": 2}]})

    table = result.sections[0].blocks[0]
    assert isinstance(table, Table)
    assert len(table.rows) == 2
    assert [[cell.blocks[0].plain_text for cell in row.cells] for row in table.rows] == [["A", "1"], ["B", "2"]]


def test_render_typed_image_and_formula():
    template = DocumentModel(
        sections=[Section(blocks=[Paragraph(content=[TextRun("{{ logo }}"), TextRun("{{ equation }}")])])]
    )

    result = render_document_template(
        template,
        {
            "logo": TemplateImage(data=b"svg", media_type="image/svg+xml", filename="logo.svg", alt_text="Logo"),
            "equation": TemplateFormula("E=mc^2", FormulaFormat.LATEX, fallback_text="E = mc²"),
        },
    )

    content = result.sections[0].blocks[0].content
    assert isinstance(content[0], Image)
    assert isinstance(content[1], Formula)
    assert content[1].fallback_text == "E = mc²"
    assert result.resources[content[0].resource_id].media_type == "image/svg+xml"
    assert result.resources[content[0].resource_id].data == b"svg"


def test_strict_mode_reports_missing_variable():
    template = DocumentModel(sections=[Section(blocks=[_paragraph("{{ missing }}")])])

    with pytest.raises(GenerateError, match="missing"):
        render_document_template(template, {})


def test_unclosed_control_block_is_rejected():
    template = DocumentModel(sections=[Section(blocks=[_paragraph("{% if enabled %}"), _paragraph("text")])])

    with pytest.raises(GenerateError, match="endif"):
        render_document_template(template, {"enabled": True})


def test_generate_docx_template_end_to_end(tmp_path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "result.docx"
    source = Document()
    source.add_heading("{{ title }}", level=1)
    source.add_paragraph("{% for item in items %}")
    source.add_paragraph("Item: {{ item }}")
    source.add_paragraph("{% endfor %}")
    source.save(template_path)

    report = generate_docx_template(template_path, output_path, {"title": "Report", "items": ["A", "B"]})
    generated = Document(output_path)

    assert report.success is True
    assert [paragraph.text for paragraph in generated.paragraphs] == ["Report", "Item: A", "Item: B"]


def test_render_bibliography_placeholder_builds_numbered_list():
    template = DocumentModel(sections=[Section(blocks=[_paragraph("{{ bibliography() }}")])])

    result = render_document_template(
        template,
        {
            "bibliography": [
                {"authors": ["Иванов, И. И."], "title": "Анализ данных", "year": 2024},
                {"authors": ["Петров, П. П."], "title": "Методы", "year": 2023},
            ]
        },
    )

    blocks = result.sections[0].blocks
    assert len(blocks) == 2
    assert blocks[0].plain_text.startswith("[1]")
    assert "Анализ данных" in blocks[0].plain_text
    assert blocks[1].plain_text.startswith("[2]")
    assert "Методы" in blocks[1].plain_text
    assert blocks[0].properties["bibliography_item"]["index"] == 1


def test_render_toc_placeholder_lists_headings_with_levels():
    template = DocumentModel(
        sections=[
            Section(blocks=[_paragraph("{{ toc() }}")]),
            Section(blocks=[Paragraph(content=[TextRun("Введение")], properties={"heading_level": 1}), _paragraph("text")]),
            Section(blocks=[Paragraph(content=[TextRun("Раздел")], properties={"heading_level": 2}), _paragraph("text")]),
        ]
    )

    result = render_document_template(template, {})

    toc = [block for block in result.sections[0].blocks if "Введение" in block.plain_text or "Раздел" in block.plain_text]
    assert len(toc) == 2
    assert toc[0].plain_text.startswith("Введение")
    assert toc[0].plain_text.endswith("2")
    assert toc[1].plain_text.startswith("    Раздел")
    assert toc[1].plain_text.endswith("3")
    assert toc[0].properties["toc_entry"]["level"] == 1


def test_render_cross_references_number_figures_tables_and_captions():
    template = DocumentModel(
        sections=[
            Section(
                blocks=[
                    _paragraph("см. {{ ref('fig:logo') }}"),
                    _paragraph("Рис. {{ id('fig:logo') }} Логотип"),
                    _paragraph("{{ ref('tbl:data') }}"),
                    _paragraph("Таблица {{ id('tbl:data') }} данные"),
                    _paragraph("в таб. {{ ref('fig:logo') }} и {{ ref('tbl:data') }}"),
                ]
            )
        ]
    )

    result = render_document_template(template, {})

    assert [block.plain_text for block in result.sections[0].blocks] == [
        "см. Рис. 1",
        "Рис. 1 Логотип",
        "Табл. 1",
        "Таблица 1 данные",
        "в таб. Рис. 1 и Табл. 1",
    ]


def test_render_cross_reference_inside_table_cell():
    template = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Table(
                        rows=[
                            TableRow(
                                cells=[
                                    TableCell(
                                        blocks=[
                                            _paragraph("Иллюстрация {{ id('fig:map') }}"),
                                            _paragraph("см. {{ ref('fig:map') }}"),
                                        ]
                                    )
                                ]
                            )
                        ]
                    )
                ]
            )
        ]
    )

    result = render_document_template(template, {})

    cell = result.sections[0].blocks[0].rows[0].cells[0]
    assert [block.plain_text for block in cell.blocks] == ["Иллюстрация 1", "см. Рис. 1"]


def test_render_undefined_cross_reference_is_rejected():
    template = DocumentModel(sections=[Section(blocks=[_paragraph("см. {{ ref('fig:missing') }}")])])

    with pytest.raises(GenerateError, match="fig:missing"):
        render_document_template(template, {})
