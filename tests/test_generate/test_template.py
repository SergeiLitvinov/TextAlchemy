import tempfile
from pathlib import Path

from textalchemy.generate import DocumentTemplate, TemplateEngine, generate_document, list_templates


def test_list_templates_empty():
    with tempfile.TemporaryDirectory() as tmp:
        templates = list_templates(tmp)
        assert templates == []


def test_template_engine_list_empty():
    with tempfile.TemporaryDirectory() as tmp:
        engine = TemplateEngine(tmp)
        assert engine.list_templates() == []


def test_generate_document_no_template():
    with tempfile.TemporaryDirectory() as tmp:
        import pytest
        with pytest.raises(Exception):
            generate_document("nonexistent", Path(tmp) / "out.docx", templates_dir=tmp)


def test_document_template():
    t = DocumentTemplate(name="test", description="Test template")
    assert t.name == "test"
    assert t.template_type == "docx"


def test_generate_document_replaces_placeholders():
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        tmpl_dir = Path(tmp) / "templates"
        tmpl_dir.mkdir()
        tmpl_path = tmpl_dir / "letter.docx"
        doc = Document()
        doc.add_paragraph("Dear {{name}}, your code is {{code}}.")
        doc.save(str(tmpl_path))

        out = Path(tmp) / "out.docx"
        generate_document("letter", out, {"name": "Ivan", "code": "ABC"}, templates_dir=tmpl_dir)

        result = Document(str(out))
        text = "\n".join(p.text for p in result.paragraphs)
        assert "Dear Ivan, your code is ABC." in text
        assert "{{name}}" not in text
        assert "{{code}}" not in text


def test_generate_document_preserves_run_formatting():
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        tmpl_dir = Path(tmp) / "templates"
        tmpl_dir.mkdir()
        tmpl_path = tmpl_dir / "bold.docx"
        doc = Document()
        para = doc.add_paragraph()
        run = para.add_run("Hello {{who}}")
        run.bold = True
        doc.save(str(tmpl_path))

        out = Path(tmp) / "out.docx"
        generate_document("bold", out, {"who": "World"}, templates_dir=tmpl_dir)

        result = Document(str(out))
        run = result.paragraphs[0].runs[0]
        assert run.text == "Hello World"
        assert run.bold is True
