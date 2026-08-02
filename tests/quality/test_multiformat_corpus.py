"""Golden-проверки воспроизводимого PDF/PPTX/EMF corpus."""

from __future__ import annotations

import json
import struct
import zipfile
from pathlib import Path

import fitz
import pytest
from lxml import etree
from pptx import Presentation

from tests.corpus.multiformat import EMF_SIGNATURE, build_multiformat_corpus, sha256
from textalchemy.core.document_model import Paragraph, Table
from textalchemy.formats.pdf import read_pdf_geometry
from textalchemy.formats.pptx import read_pptx_model

CORPUS_DIR = Path(__file__).parents[1] / "corpus"
GOLDEN = json.loads((CORPUS_DIR / "multiformat.golden.json").read_text(encoding="utf-8"))


def test_multiformat_corpus_is_byte_reproducible(tmp_path):
    first = build_multiformat_corpus(tmp_path / "first")
    second = build_multiformat_corpus(tmp_path / "second")

    assert set(first) == {"emf", "pdf", "pptx"}
    assert {name: sha256(path) for name, path in first.items()} == {
        name: sha256(path) for name, path in second.items()
    }


def test_emf_corpus_matches_binary_golden(tmp_path):
    emf = build_multiformat_corpus(tmp_path)["emf"]
    data = emf.read_bytes()

    record_type, header_size = struct.unpack_from("<II", data)
    signature = struct.unpack_from("<I", data, 40)[0]
    total_bytes, records = struct.unpack_from("<II", data, 48)
    eof_type, eof_size = struct.unpack_from("<II", data, len(data) - 20)

    assert len(data) == GOLDEN["emf"]["bytes"]
    assert record_type == 1
    assert header_size == 108
    assert signature == EMF_SIGNATURE == GOLDEN["emf"]["signature"]
    assert total_bytes == len(data)
    assert records == GOLDEN["emf"]["records"]
    assert (eof_type, eof_size) == (14, 20)


def test_pdf_corpus_matches_geometry_and_content_golden(tmp_path):
    pdf = build_multiformat_corpus(tmp_path)["pdf"]
    geometry = read_pdf_geometry(pdf)

    assert len(geometry.pages) == GOLDEN["pdf"]["pages"]
    assert geometry.metadata["title"] == GOLDEN["pdf"]["title"]
    assert [[page.width, page.height] for page in geometry.pages] == GOLDEN["pdf"]["page_sizes_pt"]
    assert sum(len(page.text_blocks) for page in geometry.pages) >= GOLDEN["pdf"]["minimum_text_blocks"]
    assert sum(len(page.image_blocks) for page in geometry.pages) >= GOLDEN["pdf"]["minimum_images"]
    assert sum(len(page.tables) for page in geometry.pages) == 1
    with fitz.open(pdf) as document:
        drawings = sum(len(page.get_drawings()) for page in document)
        text = "\n".join(page.get_text() for page in document)
    assert drawings >= GOLDEN["pdf"]["minimum_drawings"]
    assert "Left column" in text
    assert "Landscape formulas" in text


def test_pptx_corpus_matches_ooxml_and_semantic_golden(tmp_path):
    artifacts = build_multiformat_corpus(tmp_path)
    presentation = Presentation(artifacts["pptx"])
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    notes = [slide.notes_slide.notes_text_frame.text for slide in presentation.slides]

    assert len(presentation.slides) == GOLDEN["pptx"]["slides"]
    assert sum(int(getattr(shape, "has_table", False)) for shape in shapes) >= GOLDEN["pptx"]["minimum_tables"]
    assert sum(int(getattr(shape, "has_chart", False)) for shape in shapes) >= GOLDEN["pptx"]["minimum_charts"]
    assert sum(bool(note.strip()) for note in notes) >= GOLDEN["pptx"]["minimum_notes"]
    assert "EMF must remain vector" in notes[1]

    with zipfile.ZipFile(artifacts["pptx"]) as archive:
        names = set(archive.namelist())
        for name in names:
            if name.endswith((".xml", ".rels")):
                etree.fromstring(archive.read(name))
        relationships = archive.read("ppt/slides/_rels/slide2.xml.rels").decode("utf-8")
        content_types = archive.read("[Content_Types].xml").decode("utf-8")
        emf = archive.read(GOLDEN["pptx"]["emf_part"])
    assert GOLDEN["pptx"]["emf_part"] in names
    assert "scientific-diagram.emf" in relationships
    assert 'ContentType="image/x-emf"' in content_types
    assert emf == artifacts["emf"].read_bytes()


def test_pptx_model_matches_visual_positioning_golden(tmp_path):
    pptx = build_multiformat_corpus(tmp_path)["pptx"]
    model = read_pptx_model(pptx)

    page_sizes = [[section.page.width.pt, section.page.height.pt] for section in model.sections]
    for actual, expected in zip(page_sizes, GOLDEN["pptx"]["model_page_sizes_pt"], strict=True):
        assert actual == pytest.approx(expected, abs=0.01)
    assert all(
        margin.pt == 0
        for section in model.sections
        for margin in (
            section.page.margin_top,
            section.page.margin_right,
            section.page.margin_bottom,
            section.page.margin_left,
        )
    )

    first_slide, second_slide = model.sections
    title = next(
        block
        for block in first_slide.blocks
        if isinstance(block, Paragraph) and block.plain_text.startswith("Scientific")
    )
    table = next(block for block in first_slide.blocks if isinstance(block, Table))
    formula_text = next(
        block for block in second_slide.blocks if isinstance(block, Paragraph) and block.plain_text.startswith("E = mc")
    )
    expected = GOLDEN["pptx"]["model_key_boxes_pt"]
    assert _box_values(title) == pytest.approx(expected["title"], abs=0.01)
    assert _box_values(table) == pytest.approx(expected["table"], abs=0.01)
    assert _box_values(formula_text) == pytest.approx(expected["formula_text"], abs=0.01)


def _box_values(block):
    assert block.box is not None
    return [block.box.x, block.box.y, block.box.width, block.box.height]


def test_libreoffice_profile_matches_provenance_manifest():
    office_dir = CORPUS_DIR / "office"
    manifest = json.loads((office_dir / "manifest.json").read_text(encoding="utf-8"))
    fixture_name, expected = next(iter(manifest.items()))
    fixture = office_dir / fixture_name

    assert fixture.stat().st_size == expected["bytes"]
    assert sha256(fixture) == expected["sha256"]
    with zipfile.ZipFile(fixture) as archive:
        names = set(archive.namelist())
        application = archive.read("docProps/app.xml").decode("utf-8")
        presentation = Presentation(fixture)
    assert expected["producer"] in application
    assert expected["producer_version"].split()[0] in application
    assert len(presentation.slides) == expected["slides"]
    assert expected["expected_vector_profile"] in names
