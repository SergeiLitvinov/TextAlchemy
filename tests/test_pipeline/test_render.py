"""Тесты pipeline.render."""
from __future__ import annotations

from textalchemy.core.types import BibItem, Block, BlockType, DocFormat, Table, Text
from textalchemy.pipeline.render import (
    render_bibtex,
    render_docx,
    render_gost,
    render_json,
    render_latex,
    render_latex_pandoc,  # noqa: F401 - tested below
    render_markdown,
)


def _text() -> Text:
    return Text(
        blocks=[
            Block(type=BlockType.PARAGRAPH, text="Hello world"),
            Block(type=BlockType.HEADING, text="Section 1", level=1),
            Block(type=BlockType.PARAGRAPH, text="Body text"),
        ],
        tables=[Table(rows=[["A1", "A2"], ["B1", "B2"]])],
        plain="Hello world\nSection 1\nBody text",
        source_format=DocFormat.DOCX,
        engine="python-docx",
        pages=1,
    )


def _items() -> list[BibItem]:
    return [
        BibItem(index=1, raw_text="x", authors=["Иванов И.И.", "Петров П.П."],
                title="Ferroresonance in Power Grids", year=2020,
                doc_type="article", source="Электричество", pages="10-20",
                doi="10.1109/abc.2020"),
        BibItem(index=2, raw_text="y", authors=["Smith J."],
                title="Quantum Computing", year=2019, doc_type="book",
                isbn="978-5-12345-678-9"),
    ]


class TestRenderLatex:
    def test_preamble_present(self):
        out = render_latex(text=_text(), title="T", author="A")
        assert r"\documentclass" in out
        assert r"\begin{document}" in out
        assert r"\end{document}" in out

    def test_escapes_specials(self):
        t = Text(blocks=[Block(type=BlockType.PARAGRAPH, text="price: $5 & more")])
        out = render_latex(text=t)
        assert r"\$5" in out
        assert r"\&" in out

    def test_heading_levels(self):
        t = Text(blocks=[Block(type=BlockType.HEADING, text="H", level=2)])
        out = render_latex(text=t)
        assert r"\subsection{H}" in out


class TestRenderLatexPandoc:
    def test_fallback_no_pandoc(self):
        """Без pandoc — fallback на render.latex."""
        out = render_latex_pandoc(text=_text())
        assert r"\documentclass" in out
        assert r"\begin{document}" in out

    def test_fallback_no_input_path(self):
        """Без input_path — fallback на render.latex."""
        out = render_latex_pandoc(text=_text(), input_path=None)
        assert r"\documentclass" in out

    def test_fallback_nonexistent_pandoc(self):
        """С input_path но без pandoc — fallback."""
        out = render_latex_pandoc(text=_text(), input_path="nonexistent.docx")
        assert r"\documentclass" in out


class TestRenderDocx:
    def test_writes_file(self, tmp_path):
        out = tmp_path / "out.docx"
        path = render_docx(text=_text(), output_path=out)
        assert path == out
        assert out.is_file()
        # Проверим, что контент попал в DOCX
        from docx import Document
        d = Document(str(out))
        all_text = "\n".join(p.text for p in d.paragraphs)
        assert "Hello world" in all_text
        assert "Body text" in all_text
        # Таблица
        assert len(d.tables) == 1
        assert d.tables[0].cell(0, 0).text == "A1"
        assert d.tables[0].cell(1, 1).text == "B2"


class TestRenderBibtex:
    def test_basic(self):
        out = render_bibtex(items=_items())
        assert "@misc{item001," in out
        assert "@misc{item002," in out
        assert "Иванов И.И. and Петров П.П." in out
        assert "Ferroresonance in Power Grids" in out
        assert "year   = {2020}" in out
        assert "doi    = {10.1109/abc.2020}" in out


class TestRenderGost:
    def test_basic(self):
        out = render_gost(items=_items())
        # 1. авторы, название, // источник, год, страницы, DOI
        assert "1. Иванов И.И., Петров П.П. Ferroresonance in Power Grids // Электричество" in out
        assert "2020" in out
        assert "10-20" in out
        assert "DOI: 10.1109/abc.2020" in out
        # 2. книга с ISBN
        assert "ISBN: 978-5-12345-678-9" in out


class TestRenderMarkdown:
    def test_basic(self):
        out = render_markdown(items=_items())
        assert "1. Иванов" in out
        assert "**Ferroresonance in Power Grids**" in out


class TestRenderJson:
    def test_basic(self):
        out = render_json(items=_items())
        import json
        data = json.loads(out)
        assert len(data) == 2
        assert data[0]["title"] == "Ferroresonance in Power Grids"
        assert data[0]["doi"] == "10.1109/abc.2020"


class TestRenderRegistry:
    def test_all_registered(self):
        from textalchemy.core.registry import all_operations

        ids = {s.id for s in all_operations()}
        assert "render.latex" in ids
        assert "render.latex.pandoc" in ids
        assert "render.docx" in ids
        assert "render.bibtex" in ids
        assert "render.gost" in ids
        assert "render.markdown" in ids
        assert "render.json" in ids
