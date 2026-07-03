"""Тесты pipeline/render_html.py."""
from __future__ import annotations

from pathlib import Path

from textalchemy.core.types import DocFormat, Document
from textalchemy.pipeline.ingest import ingest_file  # noqa: F401
from textalchemy.pipeline.render_html import render_html_pptx  # noqa: F401


def _make_pptx(path: Path) -> Path:
    """Создаёт минимальный .pptx через python-pptx."""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Test"
    prs.save(str(path))
    return path


class TestRenderHtmlRegistry:
    def test_registered(self):
        from textalchemy.core.registry import all_operations
        ids = {s.id for s in all_operations()}
        assert "render.html.pptx" in ids


class TestRenderHtmlPptx:
    def test_creates_html(self, tmp_path):
        src = _make_pptx(tmp_path / "test.pptx")
        out_dir = tmp_path / "out"
        doc = ingest_file(path=src)
        assert doc.format == DocFormat.PPTX

        result = render_html_pptx(doc=doc, output_dir=out_dir)
        assert result.success, result.error
        # Должен быть index.html
        assert (out_dir / "index.html").is_file()
        # Ассеты скопированы
        assert (out_dir / "assets" / "css" / "main.css").is_file()
        assert (out_dir / "assets" / "js" / "main.js").is_file()

    def test_no_assets(self, tmp_path):
        src = _make_pptx(tmp_path / "t.pptx")
        out_dir = tmp_path / "out2"
        doc = ingest_file(path=src)
        result = render_html_pptx(
            doc=doc, output_dir=out_dir, copy_assets=False,
        )
        assert result.success
        assert (out_dir / "index.html").is_file()

    def test_wrong_format_fails_gracefully(self, tmp_path):
        f = tmp_path / "test.txt"
        f.write_text("not pptx")
        doc = ingest_file(path=f)
        out_dir = tmp_path / "out3"
        result = render_html_pptx(doc=doc, output_dir=out_dir)
        assert not result.success
        assert "not a pptx" in (result.error or "")

    def test_missing_file(self, tmp_path):
        doc = Document(
            path=tmp_path / "nope.pptx",
            format=DocFormat.PPTX, size=0, sha256="",
        )
        out_dir = tmp_path / "out4"
        result = render_html_pptx(doc=doc, output_dir=out_dir)
        assert not result.success
