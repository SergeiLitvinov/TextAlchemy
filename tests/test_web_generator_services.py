"""Сценарии генератора выполняются без HTTP; ошибки сохраняют наборы и освобождают файлы."""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from textalchemy.core.artifacts import ArtifactWorkspace
from textalchemy.generate.template_schema import TemplateField, TemplateSchema, TemplateValueType
from textalchemy.web.services.generator_datasets import DatasetStore
from textalchemy.web.services.generator_execution import GeneratedFile, GeneratorService
from textalchemy.web.services.generator_sessions import GeneratorDatasetService
from textalchemy.web.services.ocr_drafts import DraftConflictError


@pytest.fixture
def generator(tmp_path: Path) -> GeneratorService:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("{{ title }}")
    document.save(source)
    schema = TemplateSchema(fields=[TemplateField("title", TemplateValueType.STRING, required=True)])
    return GeneratorService(
        store_generated_preview=lambda path, fmt: {"success": True, "filename": path.name},
        schema_provider=lambda name: (schema, "sidecar"),
        resolve_template=lambda name: source,
        workspace_factory=lambda: ArtifactWorkspace(parent=tmp_path),
    )


@pytest.mark.parametrize("format", ["docx", "html"])
def test_generate_real_file_and_release_after_consumer_reads(
    generator: GeneratorService,
    tmp_path: Path,
    format: str,
) -> None:
    original = (tmp_path / "source.docx").read_bytes()
    result = generator.generate(template="sample", format=format, params='{"title": "Проверка"}')
    assert isinstance(result, GeneratedFile)
    if format == "docx":
        assert Document(result.path).paragraphs[0].text == "Проверка"
    else:
        assert "Проверка" in result.path.read_text(encoding="utf-8")
    assert (tmp_path / "source.docx").read_bytes() == original
    result.cleanup()
    assert not result.path.parent.exists()


def test_preview_retains_copy_and_marks_missing_fields(generator: GeneratorService, tmp_path: Path) -> None:
    retained = tmp_path / "retained.docx"

    def retain(path: Path, format: str) -> dict[str, Any]:
        retained.write_bytes(path.read_bytes())
        return {"success": True, "filename": path.name}

    service = replace(generator, store_generated_preview=retain)
    result = service.generate(template="sample", preview=True)
    assert isinstance(result, dict) and result["draft"] is True
    assert result["missing_fields"][0]["name"] == "title"
    assert "Черновик" in result["filename"]
    assert retained.exists() and not list(tmp_path.glob("textalchemy_*"))


@pytest.mark.parametrize("params", ["[]", "null", "{", "{}"])
def test_invalid_data_allocates_no_workspace(generator: GeneratorService, tmp_path: Path, params: str) -> None:
    result = generator.generate(template="sample", params=params)
    assert isinstance(result, dict) and result["success"] is False
    assert not list(tmp_path.glob("textalchemy_*"))


def test_failed_export_removes_partial_file(generator: GeneratorService, tmp_path: Path) -> None:
    def fail(source: Path, output: Path, data: dict[str, Any], **options: Any) -> None:
        output.write_bytes(b"partial")
        raise RuntimeError("export failed")

    service = replace(generator, generate_docx_template=fail, fill_text_package=lambda *args: False)
    result = service.generate(template="sample", params='{"title": "Valid"}')
    assert isinstance(result, dict) and result["error"] == "export failed"
    assert not list(tmp_path.glob("textalchemy_*"))


def test_failed_preview_storage_cleans_workspace(generator: GeneratorService, tmp_path: Path) -> None:
    def fail(path: Path, format: str) -> dict[str, Any]:
        raise OSError("storage failed")

    result = replace(generator, store_generated_preview=fail).generate(template="sample", preview=True)
    assert isinstance(result, dict) and result["error"] == "storage failed"
    assert not list(tmp_path.glob("textalchemy_*"))


def test_failed_path_allocation_cleans_workspace(generator: GeneratorService, tmp_path: Path) -> None:
    class BrokenWorkspace(ArtifactWorkspace):
        def artifact_path(self, name: str, *, fallback: str = "artifact") -> Path:
            raise OSError("path failed")

    service = replace(generator, workspace_factory=lambda: BrokenWorkspace(parent=tmp_path))
    result = service.generate(template="sample", params='{"title": "Valid"}')
    assert isinstance(result, dict) and result["error"] == "path failed"
    assert not list(tmp_path.glob("textalchemy_*"))


def test_dataset_service_preserves_saved_data_on_schema_change(tmp_path: Path) -> None:
    schema = TemplateSchema(fields=[TemplateField("title", TemplateValueType.STRING)])
    service = GeneratorDatasetService(DatasetStore(tmp_path), lambda name: (schema, "sidecar"))
    snapshot = {
        "version": 1,
        "format": "docx",
        "output": "result.docx",
        "schema": json.dumps(schema.to_dict()),
        "values": [{"name": "title", "value": "Сохранённое"}],
    }
    item = service.save(template="sample", name="Набор", snapshot=snapshot)
    assert service.get(item["id"])["snapshot"] == snapshot
    service.schema_provider = lambda name: (TemplateSchema(fields=[]), "derived")
    with pytest.raises(DraftConflictError):
        service.get(item["id"])
    assert service.store.get(item["id"])["snapshot"] == snapshot


def test_dataset_update_requires_revision_and_does_not_overwrite(tmp_path: Path) -> None:
    schema = TemplateSchema(fields=[])
    service = GeneratorDatasetService(DatasetStore(tmp_path), lambda name: (schema, "derived"))
    snapshot = {"version": 1, "format": "docx", "output": "a.docx", "schema": json.dumps(schema.to_dict()), "values": []}
    item = service.save(template="sample", name="Original", snapshot=snapshot)
    with pytest.raises(ValueError, match="ревизия"):
        service.save(template="sample", name="Changed", snapshot=snapshot, dataset_id=item["id"])
    assert service.get(item["id"])["name"] == "Original"
