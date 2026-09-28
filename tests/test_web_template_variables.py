"""Edited template copies retain their schema and produce real documents."""
import importlib
import io
import json
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from textalchemy.web.services.template_variables import custom_path


@pytest.fixture
def variable_client(tmp_path, monkeypatch):
    templates = tmp_path / 'templates'
    templates.mkdir()
    source = templates / 'example.docx'
    doc = Document()
    para = doc.add_paragraph()
    para.add_run('Before ').bold = True
    para.add_run('{{ti').italic = True
    para.add_run('tle}} / {{ title }} after')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '{{title}}'
    doc.sections[0].footer.paragraphs[0].text = '{{title}}'
    doc.save(source)
    schema = {'fields': [{'name': 'title', 'type': 'string', 'required': True}], 'allow_extra': False}
    source.with_suffix('.schema.json').write_text(json.dumps(schema), encoding='utf-8')
    monkeypatch.setattr('textalchemy.generate.template.TEMPLATES_DIR', templates)
    web = importlib.import_module('textalchemy.web.app')
    monkeypatch.setattr(web, 'data_dir', tmp_path / 'data')
    with TestClient(web.app) as client:
        yield client, source


def test_rename_copy_schema_reopen_and_generate(variable_client):
    client, source = variable_client
    original = source.read_bytes()
    inspection = client.get('/api/generate/templates/example/variables').json()
    assert len(inspection['variables'][0]['contexts']) == 4
    result = client.post('/api/generate/templates/example/variables', json={
        'old': 'title', 'new': 'subject', 'revision': inspection['revision']})
    assert result.status_code == 200, result.text
    name = result.json()['name']
    assert source.read_bytes() == original
    assert name in {item['name'] for item in client.get('/api/generate/templates').json()}
    schema = client.get(f'/api/generate/templates/{name}/schema').json()['schema']
    assert schema['fields'][0]['name'] == 'subject'
    assert schema['fields'][0]['required'] is True
    edited = Document(custom_path(name))
    assert edited.paragraphs[0].text == 'Before {{subject}} / {{ subject }} after'
    assert edited.paragraphs[0].runs[0].bold
    assert edited.paragraphs[0].runs[1].italic
    assert edited.tables[0].cell(0, 0).text == '{{subject}}'
    assert edited.sections[0].footer.paragraphs[0].text == '{{subject}}'
    with ZipFile(io.BytesIO(original)) as before, ZipFile(custom_path(name)) as after:
        assert before.read('word/styles.xml') == after.read('word/styles.xml')
    response = client.post('/api/generate', data={'template': name, 'params': json.dumps({'subject': 'Edited value'})})
    assert response.headers['content-type'].startswith('application/vnd.openxmlformats'), response.text
    generated = Document(io.BytesIO(response.content))
    assert generated.paragraphs[0].text == 'Before Edited value / Edited value after'
    assert generated.tables[0].cell(0, 0).text == 'Edited value'


@pytest.mark.parametrize('new', ['title', 'a.b', '../escape', '2name', ''])
def test_invalid_rename_creates_no_copy(variable_client, new):
    client, source = variable_client
    original = source.read_bytes()
    revision = client.get('/api/generate/templates/example/variables').json()['revision']
    response = client.post('/api/generate/templates/example/variables', json={'old': 'title', 'new': new, 'revision': revision})
    assert response.status_code == 422
    assert source.read_bytes() == original
    assert len(client.get('/api/generate/templates').json()) == 1


def test_changed_template_and_complex_syntax_rejected(variable_client):
    client, source = variable_client
    revision = client.get('/api/generate/templates/example/variables').json()['revision']
    doc = Document(source)
    doc.add_paragraph('Changed')
    doc.save(source)
    assert client.post('/api/generate/templates/example/variables', json={
        'old': 'title', 'new': 'subject', 'revision': revision}).status_code == 422
    doc.add_paragraph('{% if title %}Conditional{% endif %}')
    doc.save(source)
    assert client.get('/api/generate/templates/example/variables').status_code == 422
    assert len(client.get('/api/generate/templates').json()) == 1
