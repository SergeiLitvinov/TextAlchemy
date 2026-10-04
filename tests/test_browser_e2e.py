"""Реальные browser E2E и автоматическая accessibility-проверка (TODO:128).

Запускается настоящий uvicorn-сервер в отдельном потоке, тесты управляют
headless Chromium через Playwright и прогоняют axe-core для WCAG-проверок.

Требует dev-зависимостей ``playwright`` и ``axe-core-python``, а также
установленного браузера::

    uv run playwright install chromium

Если браузер не установлен, тесты пропускаются, чтобы не ломать обычный
``pytest`` на машинах без playwright.
"""

from __future__ import annotations

import socket
import threading
import time
from pathlib import Path

import pytest
import uvicorn

pytest.importorskip("playwright")
pytest.importorskip("axe_core_python")

from playwright.sync_api import Browser, Page, sync_playwright  # noqa: E402

from textalchemy.web.app import app  # noqa: E402
from textalchemy.web.queue import task_queue  # noqa: E402
from textalchemy.web.routes import convert as convert_route  # noqa: E402
from textalchemy.web.tasks import TaskStore  # noqa: E402

PAGES = [
    "/",
    "/extract",
    "/convert",
    "/pipeline",
    "/bibliography",
    "/matching",
    "/reports",
    "/recognize",
    "/generate",
    "/pdf-order",
    "/export",
]

_VIEWPORT_WIDTHS = (375, 768, 1280)


def _make_pdf(path: Path, text: str = "E2E Convert Report") -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=400, height=500)
    page.insert_text((50, 72), text, fontsize=14)
    doc.save(str(path))
    doc.close()
    return path


def _free_port() -> int:
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


class _ServerThread(threading.Thread):
    def __init__(self, port: int) -> None:
        super().__init__(daemon=True, name="e2e-uvicorn")
        self.port = port
        self.config = uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning")
        self.server = uvicorn.Server(self.config)

    def run(self) -> None:
        self.server.run()


@pytest.fixture(scope="session")
def e2e_server(tmp_path_factory):
    """Реальный HTTP-сервер (uvicorn) поверх того же app, что и TestClient."""
    import importlib

    web_app = importlib.import_module("textalchemy.web.app")

    # Session-фикстура не может использовать function-scoped monkeypatch,
    # поэтому патчим вручную и откатываем в teardown.
    session_patch = pytest.MonkeyPatch()
    data = tmp_path_factory.mktemp("e2e-data")
    session_patch.setattr(web_app, "data_dir", data)
    from textalchemy.core.database import Database

    session_patch.setattr(web_app, "db", Database(db_path=data / "library.db"))
    session_store = TaskStore(tmp_path_factory.mktemp("e2e-session-tasks"))
    session_patch.setattr(web_app, "tasks_store", session_store)
    session_patch.setattr(convert_route, "tasks_store", session_store)

    port = _free_port()
    thread = _ServerThread(port)
    thread.start()
    try:
        for _ in range(200):
            if thread.server.started:
                break
            time.sleep(0.05)
        else:
            raise RuntimeError("uvicorn server did not start")
        yield f"http://127.0.0.1:{port}"
    finally:
        thread.server.should_exit = True
        thread.join(timeout=15)
        session_patch.undo()


@pytest.fixture()
def task_store(tmp_path_factory, monkeypatch):
    """Изолированное хранилище задач на время одного E2E-теста."""
    import importlib

    web_app = importlib.import_module("textalchemy.web.app")
    store = TaskStore(tmp_path_factory.mktemp("e2e-tasks"))
    monkeypatch.setattr(web_app, "tasks_store", store)
    monkeypatch.setattr(convert_route, "tasks_store", store)
    yield store
    task_queue.wait_idle(timeout=30)


@pytest.fixture()
def browser() -> Browser:
    with sync_playwright() as pw:
        try:
            launched = pw.chromium.launch(headless=True)
        except Exception as error:  # noqa: BLE001 - skip gracefully when browsers missing
            pytest.skip(f"playwright browser unavailable: {error}")
        yield launched
        launched.close()


@pytest.fixture()
def page(browser: Browser) -> Page:
    context = browser.new_context()
    current = context.new_page()
    errors: list[str] = []
    current.on("console", lambda msg: errors.append(msg.text) if msg.type == "error" else None)
    current.on("pageerror", lambda exc: errors.append(str(exc)))
    current.e2e_errors = errors  # type: ignore[attr-defined]
    yield current
    context.close()


def _wait_convert_ready(current: Page) -> None:
    current.wait_for_function(
        "document.getElementById('drop-hint').textContent.includes('PDF')",
        timeout=30000,
    )


def _open_batch_history(current: Page, job_id: str | None = None) -> None:
    history = current.locator("#conversionHistory")
    if not history.evaluate("element => element.open"):
        history.locator(":scope > summary").click()
    if job_id is not None:
        entry = current.locator("#historyList .job-entry").filter(has=current.locator(f'[data-open-job="{job_id}"]'))
        if not entry.evaluate("element => element.open"):
            entry.locator(":scope > summary").click()


def test_e2e_dashboard_loads_without_console_errors(e2e_server, page):
    page.goto(f"{e2e_server}/")
    page.wait_for_load_state("networkidle")
    assert page.locator("h1").first.text_content() == "TextAlchemy"
    assert page.locator("nav a").count() == 11
    assert page.get_by_role("link", name="Руководство", exact=True).get_attribute("href") == "/help/"
    assert page.e2e_errors == []


def test_e2e_conversion_catalog_guidance_and_batch_forecast(e2e_server, page):
    page.set_viewport_size({"width": 375, "height": 812})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    summary = page.locator("#formatCatalog > summary")
    summary.focus()
    page.keyboard.press("Enter")
    assert page.locator("#formatCatalog").get_attribute("open") is not None
    assert "PowerPoint" in page.locator("#formatCatalogContent").text_content()
    assert "EPUB" in page.locator("#formatCatalogContent").text_content()
    page.locator("#fileInput").set_input_files({"name": "note.txt", "mimeType": "text/plain", "buffer": b"Test text"})
    assert page.locator("#target").input_value() == "docx"
    page.locator("#target").select_option("pptx")
    assert "редактируемые объекты" in page.locator("#routeGuidance").text_content()
    assert page.locator('#mode option[value="faithful"]').is_disabled()
    page.locator("#target").select_option("docx")
    assert page.locator("#routeGuidance").is_hidden()
    page.locator("#fileInput").set_input_files(
        [
            {"name": "first.txt", "mimeType": "text/plain", "buffer": b"First"},
            {"name": "second.txt", "mimeType": "text/plain", "buffer": b"Second"},
        ]
    )
    page.locator("#target").select_option("pptx")
    assert "Для каждого файла" in page.locator("#route-help").text_content()
    assert "%" not in page.locator("#route-help").text_content()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    assert page.e2e_errors == []


def test_e2e_conversion_catalog_failure_is_explained(e2e_server, page):
    page.route("**/api/convert/capabilities", lambda route: route.fulfill(status=503, body="Unavailable"))
    page.goto(f"{e2e_server}/convert")
    page.wait_for_function("document.getElementById('formatCatalogContent').textContent.includes('Обновите страницу')")
    assert "Не удалось" in page.locator("#drop-hint").text_content()


def test_e2e_catalog_explains_missing_components(e2e_server, page):
    from textalchemy.convert.executor import ConversionExecutor
    from textalchemy.web.services.conversion_catalog import available_conversions

    catalog = available_conversions(ConversionExecutor(requirement_checker=lambda name: name != "python-pptx"))
    page.route("**/api/convert/capabilities", lambda route: route.fulfill(json=catalog))
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#formatCatalog > summary").click()
    content = page.locator("#formatCatalogContent")
    assert "python-pptx" in content.text_content()
    assert "Преобразование в этом режиме пока не поддерживается" in content.text_content()
    assert "нет доступных направлений" in content.text_content()
    assert ".pptx" not in page.locator("#fileInput").get_attribute("accept")
    page.locator("#fileInput").set_input_files({"name": "note.txt", "mimeType": "text/plain", "buffer": b"Text"})
    assert page.locator('#target option[value="pptx"]').count() == 0
    assert page.e2e_errors == []


def test_e2e_global_task_center_explains_storage(e2e_server, page):
    page.goto(f"{e2e_server}/")
    page.wait_for_load_state("networkidle")

    page.get_by_role("button", name="Задачи").click()

    center = page.get_by_role("complementary", name="Задачи и результаты")
    assert center.get_attribute("inert") is None
    assert "Файлы не отправляются" in center.text_content()
    assert "автоматически очищаются через 1 час" in center.text_content()
    page.keyboard.press("Escape")
    assert page.locator("#taskCenter").get_attribute("inert") is not None


def test_e2e_mobile_navigation_opens_and_closes(e2e_server, page):
    page.set_viewport_size({"width": 390, "height": 844})
    page.goto(f"{e2e_server}/")
    page.wait_for_load_state("networkidle")

    toggle = page.get_by_role("button", name="Открыть меню")
    toggle.click()
    assert page.locator("body").evaluate("element => element.classList.contains('nav-open')") is True
    assert toggle.get_attribute("aria-expanded") == "true"

    page.keyboard.press("Escape")
    assert page.locator("body").evaluate("element => element.classList.contains('nav-open')") is False
    assert toggle.get_attribute("aria-expanded") == "false"


def test_e2e_generate_document_from_template(e2e_server, page):
    page.goto(f"{e2e_server}/generate")
    page.wait_for_selector("#field-title", timeout=30000)

    page.fill("#field-title", "Проверка шаблона")
    page.fill("#field-author", "TextAlchemy")
    page.fill("#field-body", "Документ создан через пользовательский сценарий.")

    with page.expect_download(timeout=60000) as download_info:
        page.click("#genBtn")
    data = download_info.value.path().read_bytes()
    assert data[:4] == b"PK\x03\x04"
    assert page.locator("#status").text_content() == "Документ сгенерирован и загружен."


def test_e2e_recognize_pdf_text_layer(e2e_server, page, tmp_path):
    pdf = _make_pdf(tmp_path / "recognize.pdf", "Recognized E2E text")
    page.goto(f"{e2e_server}/recognize")

    page.locator("#scenario").select_option("fast")
    page.set_input_files("#fileInput", str(pdf))
    page.locator("#processBtn").click()
    page.wait_for_function("document.getElementById('result').value.includes('Recognized E2E text')", timeout=30000)

    assert page.locator("#scenario").input_value() == "fast"
    assert page.locator("#copyBtn").is_enabled()
    assert "Текст готов к проверке" in page.locator("#status").text_content()


def test_e2e_extract_text_from_pdf(e2e_server, page, tmp_path):
    pdf = _make_pdf(tmp_path / "extract.pdf", "Extracted E2E content")
    page.goto(f"{e2e_server}/extract")
    page.set_input_files("#fileInput", str(pdf))
    page.locator("#processBtn").click()
    page.wait_for_function("document.getElementById('result').value.includes('Extracted E2E content')", timeout=30000)

    assert page.locator("#copyBtn").is_enabled()
    assert "Готово:" in page.locator("#status").text_content()


def test_e2e_library_navigation_and_editor(e2e_server, page):
    page.goto(f"{e2e_server}/bibliography")
    tabs = page.get_by_role("navigation", name="Разделы библиотеки")
    assert tabs.get_by_role("link").count() == 4

    page.get_by_role("button", name="Добавить источник").click()
    assert page.locator("#bibEditor").get_attribute("open") is not None
    assert page.locator("#title").evaluate("element => element === document.activeElement") is True

    tabs.get_by_role("link", name="Связать файлы").click()
    page.wait_for_url(f"{e2e_server}/matching")
    library_tabs = page.get_by_role("navigation", name="Разделы библиотеки")
    assert library_tabs.get_by_role("link", name="Связать файлы").get_attribute("aria-current") == "page"
    assert page.locator("#dryRun").is_checked()


def test_e2e_pipeline_visual_expert_roundtrip(e2e_server, page):
    page.goto(f"{e2e_server}/pipeline")
    page.wait_for_selector("#stepsList .step-card", timeout=30000)

    page.click("#validateBtn")
    page.wait_for_function("document.getElementById('status').textContent.includes('Связи в порядке')")

    page.click("#modeExpertBtn")
    page.wait_for_function("document.getElementById('specText').value.includes('render.latex')")
    assert "title: Demo" in page.locator("#specText").input_value()

    page.click("#modeVisualBtn")
    page.wait_for_function("!document.getElementById('visualMode').hidden")
    assert page.locator("#p-2-title").input_value() == "Demo"
    assert page.e2e_errors == []


@pytest.mark.parametrize("width", [375, 1280])
def test_e2e_html_warning_navigation(e2e_server, page, tmp_path, task_store, width):
    pytest.importorskip("bs4")
    pytest.importorskip("tinycss2")
    pytest.importorskip("docx")
    from playwright.sync_api import expect

    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    source = Path(__file__).parent / "corpus/html-warnings.html"
    page.set_input_files("#fileInput", str(source))
    page.select_option("#target", "docx")
    page.click("#convertBtn")
    expect(page.locator("#resultTitle")).to_have_text("Документ готов", timeout=60000)
    css = page.locator("#issueList li").filter(has_text="background-color").get_by_role("button")
    css.click()
    inspector = page.locator("#issueInspector")
    expect(inspector).to_be_visible()
    expect(inspector).to_be_focused()
    expect(page.locator("#issueFragment")).to_contain_text("Абзац с неподдержанным фоном.")
    css_location = inspector.get_attribute("data-location")
    image = page.locator("#issueList li").filter(has_text="Изображение не загружено").get_by_role("button")
    image.click()
    expect(page.locator("#issueFragment")).to_contain_text("Результат измерения:")
    expect(page.locator("#issueFragment")).to_contain_text("missing-chart.png")
    assert inspector.get_attribute("data-location") != css_location
    expect(css).to_have_attribute("aria-pressed", "false")
    expect(image).to_have_attribute("aria-pressed", "true")
    page.screenshot(path=str(tmp_path / f"html-inspector-{width}.png"), full_page=True)
    general = page.locator("#issueList").get_by_role("button", name="Весь документ", exact=True)
    general.focus()
    page.keyboard.press("Enter")
    expect(inspector).to_have_attribute("data-location", "html:document")
    expect(page.locator("#issueLocationTitle")).to_have_text("Весь документ")
    expect(inspector).to_be_focused()
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    # Another conversion replaces both the warning list and its selected fragment.
    page.set_input_files("#fileInput", str(Path(__file__).parent / "corpus/scientific-html.html"))
    page.select_option("#target", "model")
    page.click("#convertBtn")
    expect(page.locator("#resultCard")).to_be_visible(timeout=60000)
    expect(page.locator("#issuesSection")).to_be_hidden()
    expect(inspector).to_be_hidden()
    assert page.e2e_errors == []


def test_e2e_convert_single_pdf_to_docx(e2e_server, page, tmp_path, task_store):
    pdf = _make_pdf(tmp_path / "report.pdf")
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)

    page.set_input_files("#fileInput", str(pdf))
    page.wait_for_function("document.getElementById('sourceFilename').textContent === 'report.pdf'")
    page.wait_for_function(
        "!['Проверяем…'].includes(document.getElementById('inspectionState').textContent)",
        timeout=30000,
    )

    options = page.locator("#target option").all_text_contents()
    assert any("Word (DOCX)" in option for option in options)
    route_help = page.locator("#route-help").text_content()
    assert "сходство" in route_help and "редактируемость" in route_help

    page.click("#convertBtn")
    page.wait_for_selector("#resultCard:not([hidden])", timeout=120000)
    assert page.locator("#resultTitle").text_content() == "Документ готов"
    assert page.locator("#resultFilename").text_content() == "report.docx"

    with page.expect_download(timeout=60000) as download_info:
        page.click("#downloadBtn")
    downloaded = download_info.value
    data = downloaded.path().read_bytes()
    assert data[:4] == b"PK\x03\x04"  # DOCX = ZIP-архив
    assert page.e2e_errors == []


@pytest.mark.parametrize("budget, rejected", [("", False), ("0", True), ("1", False)])
@pytest.mark.parametrize("theme", ["light", "dark"])
def test_e2e_quality_budget(e2e_server, page, task_store, monkeypatch, budget, rejected, theme):
    from textalchemy.core.diagnostics import ConversionReport, IssueSeverity

    def export_with_loss(model, output):
        output.write_text("<p>Result</p>", encoding="utf-8")
        report = ConversionReport(output)
        report.add(IssueSeverity.LOSS, "text", "Текст сокращён")
        return report

    monkeypatch.setattr("textalchemy.convert.executor._write_html", export_with_loss)
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.evaluate("theme => document.documentElement.dataset.theme = theme", theme)
    page.set_viewport_size({"width": 375, "height": 900})
    page.set_input_files("#fileInput", {"name": "source.txt", "mimeType": "text/plain", "buffer": b"Original"})
    page.select_option("#target", "html")
    page.locator("#expertConversionBtn").click()
    page.locator(".conversion-loss-budget summary").click()
    page.select_option("#maxLossIssues", budget)
    with page.expect_response(
        lambda response: response.url.endswith("/api/convert") and response.request.method == "POST"
    ) as response:
        page.click("#convertBtn")
    created = response.value.json()
    page.wait_for_selector("#resultCard:not([hidden])", timeout=30000)
    stored = task_store.get(created["task_id"])
    assert stored["max_loss_issues"] == (int(budget) if budget else None)
    assert page.locator("#downloadBtn").is_visible() is not rejected
    summary = page.locator("#qualityGateSummary").inner_text()
    if rejected:
        assert "Результат не выдан" in summary
        assert page.locator("#qualityBadge").inner_text() == "Превышен бюджет потерь"
        assert not stored.get("artifact")
    elif budget:
        assert "Бюджет соблюдён" in summary
    else:
        assert "не проверялся" in summary
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


def test_e2e_download_editable_pptx(e2e_server, page, tmp_path, task_store):
    from pptx import Presentation

    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.set_input_files(
        "#fileInput",
        {
            "name": "slide.txt",
            "mimeType": "text/plain",
            "buffer": "Редактируемый слайд".encode(),
        },
    )
    page.select_option("#target", "pptx")
    page.click("#convertBtn")
    page.wait_for_selector("#resultCard:not([hidden])", timeout=30000)
    assert page.locator("#downloadBtn").is_visible()
    with page.expect_download() as pending:
        page.click("#downloadBtn")
    path = tmp_path / pending.value.suggested_filename
    pending.value.save_as(path)
    assert path.suffix == ".pptx"
    assert Presentation(path).slides[0].shapes[0].text == "Редактируемый слайд"
    assert page.e2e_errors == []


def test_e2e_batch_passes_quality_budget_to_each_file(e2e_server, page, task_store):
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.set_input_files(
        "#fileInput",
        [
            {"name": "first.txt", "mimeType": "text/plain", "buffer": b"First"},
            {"name": "second.txt", "mimeType": "text/plain", "buffer": b"Second"},
        ],
    )
    page.select_option("#target", "model")
    page.locator("#expertConversionBtn").click()
    page.locator(".conversion-loss-budget summary").click()
    page.select_option("#maxLossIssues", "0")
    page.select_option("#maxLostObjects", "0")
    page.select_option("#textPreservation", "paragraphs")
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.click("#batchConvertBtn")
    created = response.value.json()
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'", timeout=30000)
    for item in created["tasks"]:
        task = task_store.get(item["task_id"])
        assert task["status"] == "done"
        assert task["max_loss_issues"] == 0
        assert task["max_lost_objects"] == 0
        assert task["report"]["metrics"]["object_quality_gate"]["accepted"] is True
        assert task["text_preservation"] == "paragraphs"
        assert task["report"]["metrics"]["text_quality_gate"]["accepted"] is True
        assert task["report"]["metrics"]["quality_gate"]["accepted"] is True
    assert page.e2e_errors == []


@pytest.mark.parametrize("target, accepted", [("docx", True), ("html", False)])
def test_e2e_object_budget_reports_verification(e2e_server, page, task_store, target, accepted):
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.set_input_files("#fileInput", {"name": "source.txt", "mimeType": "text/plain", "buffer": b"Unique text"})
    page.select_option("#target", target)
    page.locator("#expertConversionBtn").click()
    page.locator(".conversion-loss-budget summary").click()
    page.select_option("#maxLostObjects", "0")
    page.click("#convertBtn")
    page.wait_for_selector("#resultCard:not([hidden])", timeout=30000)
    summary = page.locator("#objectGateSummary").inner_text()
    assert ("Бюджет объектов соблюдён" if accepted else "не удалось проверить") in summary
    assert page.locator("#downloadBtn").is_visible() is accepted
    assert page.e2e_errors == []


@pytest.mark.parametrize("changed", [False, True])
def test_e2e_exact_text_check(e2e_server, page, task_store, monkeypatch, changed):
    from textalchemy.convert.docx_writer import write_docx_model
    from textalchemy.core.document_model import TextRun

    def export(model, output):
        if changed:
            model.sections[0].blocks[0].content = [TextRun("Replaced")]
        return write_docx_model(model, output)

    monkeypatch.setattr("textalchemy.convert.executor._write_docx", export)
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.set_viewport_size({"width": 375, "height": 900})
    page.set_input_files("#fileInput", {"name": "source.txt", "mimeType": "text/plain", "buffer": b"Original"})
    page.select_option("#target", "docx")
    page.locator("#expertConversionBtn").click()
    page.locator(".conversion-loss-budget summary").click()
    page.select_option("#textPreservation", "paragraphs")
    page.click("#convertBtn")
    page.wait_for_selector("#resultCard:not([hidden])", timeout=30000)
    summary = page.locator("#textGateSummary").inner_text()
    assert ("Не совпал текст" if changed else "сохранён дословно") in summary
    assert page.locator("#downloadBtn").is_visible() is not changed
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


@pytest.mark.parametrize("mode, accepted", [("paragraphs", False), ("flow", True)])
def test_e2e_text_reflow_mode(e2e_server, page, task_store, monkeypatch, mode, accepted):
    from textalchemy.convert.docx_writer import write_docx_model
    from textalchemy.core.document_model import Paragraph, TextRun

    def export(model, output):
        model.sections[0].blocks = [Paragraph(content=[TextRun(word)]) for word in ["First", "Second"]]
        return write_docx_model(model, output)

    monkeypatch.setattr("textalchemy.convert.executor._write_docx", export)
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.set_viewport_size({"width": 375, "height": 900})
    page.set_input_files("#fileInput", {"name": "source.txt", "mimeType": "text/plain", "buffer": b"First Second"})
    page.select_option("#target", "docx")
    page.locator("#expertConversionBtn").click()
    page.locator(".conversion-loss-budget summary").click()
    page.select_option("#textPreservation", mode)
    page.click("#convertBtn")
    page.wait_for_selector("#resultCard:not([hidden])", timeout=30000)
    summary = page.locator("#textGateSummary").inner_text()
    assert ("Последовательность текста сохранена" if accepted else "Не совпал текст") in summary
    assert page.locator("#downloadBtn").is_visible() is accepted
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


def test_e2e_convert_batch_flow(e2e_server, page, tmp_path, task_store):
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)

    pdfs = [_make_pdf(tmp_path / "a.pdf", "Batch A"), _make_pdf(tmp_path / "b.pdf", "Batch B")]
    page.set_input_files("#fileInput", [str(path) for path in pdfs])
    page.wait_for_function("document.getElementById('batchFileBlock') && !document.getElementById('batchFileBlock').hidden")
    assert "2 файлов" in page.locator("#batchFileSummary").text_content()

    page.click("#batchConvertBtn")
    page.wait_for_selector("#batchProgressCard:not([hidden])", timeout=60000)
    page.wait_for_function(
        "document.getElementById('batchProgressState').textContent === 'Готово'",
        timeout=120000,
    )
    done_items = page.locator("#batchProgressList .batch-progress-item.done")
    assert done_items.count() == 2

    with page.expect_download() as download:
        page.locator("#batchArchiveBtn").click()
    archive_path = tmp_path / "batch.zip"
    download.value.save_as(archive_path)
    import json
    from zipfile import ZipFile

    with ZipFile(archive_path) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["files"]) == 2
        assert all(item["status"] == "done" and item["file"] in archive.namelist() for item in manifest["files"])

    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.wait_for_selector("#historyList details.job-entry", timeout=30000)
    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.locator("#historyList details.job-entry summary").first.click()
    page.wait_for_selector("[data-rerun-job]", timeout=30000)
    assert page.locator('#historyList [data-download-url$="/archive"]').is_visible()
    page.click("[data-rerun-job]")
    page.locator("#toast-container div").first.wait_for(timeout=30000)
    assert "перезапущена" in page.locator("#toast-container").inner_text().lower()
    page.wait_for_function(
        "document.getElementById('batchProgressState').textContent === 'Готово'",
        timeout=120000,
    )

    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.wait_for_selector("#historyList details.job-entry", timeout=30000)
    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.locator("#historyList details.job-entry summary").first.click()
    page.wait_for_selector("[data-delete-job]", timeout=30000)
    page.click("[data-delete-job]")
    page.wait_for_function(
        "document.querySelectorAll('#historyList details.job-entry').length === 0",
        timeout=30000,
    )
    assert page.e2e_errors == []


def test_e2e_dropzone_is_keyboard_operable(e2e_server, page, tmp_path):
    pdf = _make_pdf(tmp_path / "keyboard.pdf")
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)

    dropzone = page.locator("#dropZone")
    dropzone.focus()
    assert page.evaluate("document.activeElement.id") == "dropZone"
    with page.expect_file_chooser() as chooser_info:
        dropzone.press("Enter")
    chooser_info.value.set_files(str(pdf))
    page.wait_for_function("document.getElementById('sourceFilename').textContent === 'keyboard.pdf'")


def test_e2e_dropzone_reachable_via_tab(e2e_server, page):
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    focusable_count = page.locator("a[href], button, input, select, textarea, [tabindex]").count()
    for _ in range(focusable_count + 1):
        page.keyboard.press("Tab")
        if page.evaluate("document.activeElement.id") == "dropZone":
            break
    assert page.evaluate("document.activeElement.id") == "dropZone"


@pytest.mark.parametrize("width", _VIEWPORT_WIDTHS)
def test_e2e_pages_have_no_horizontal_overflow(e2e_server, page, width):
    page.set_viewport_size({"width": width, "height": 800})
    for path in PAGES:
        page.goto(f"{e2e_server}{path}")
        page.wait_for_load_state("networkidle")
        overflow = page.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
        assert overflow <= 0, f"{path} @{width}px: horizontal overflow {overflow}px"


def test_e2e_axe_accessibility_no_serious_violations(e2e_server, page):
    from axe_core_python.sync_playwright import Axe

    axe = Axe()
    for path in ("/", "/convert", "/extract", "/pipeline"):
        page.goto(f"{e2e_server}{path}")
        page.wait_for_load_state("networkidle")
        results = axe.run(page)
        serious = [v for v in results["violations"] if v.get("impact") in ("serious", "critical")]
        assert not serious, f"{path}: axe serious/critical: " + repr([(v["id"], v["impact"], v["help"]) for v in serious])


def test_e2e_convert_controls_have_accessible_names(e2e_server, page, tmp_path, task_store):
    pdf = _make_pdf(tmp_path / "names.pdf")
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)

    assert page.get_by_role("heading", name="Преобразовать документ", exact=True).count() == 1
    dropzone = page.get_by_role("button", name="Загрузите документ")
    assert dropzone.count() == 1

    # Контролы конвертации появляются только после выбора файла.
    page.set_input_files("#fileInput", str(pdf))
    page.wait_for_function("!document.getElementById('conversionSetup').hidden")
    assert page.get_by_role("combobox", name="Формат результата").count() == 1
    assert page.get_by_role("combobox", name="Приоритет").count() == 0
    page.locator("#expertConversionBtn").click()
    assert page.get_by_role("combobox", name="Приоритет").count() == 1
    assert page.get_by_role("button", name="Начать конвертацию").count() == 1
    assert page.get_by_role("button", name="Выбрать другой").count() == 1


def test_e2e_batch_filters_persist_and_switch_jobs(e2e_server, page, task_store):
    files = [
        {"task_id": "filter-running", "name": "В работе.txt"},
        {"task_id": "filter-done", "name": "Отчёт.txt"},
        {"task_id": "filter-error", "name": "Ошибка.txt"},
        {"task_id": "filter-expired", "name": "Старый.txt"},
    ]
    for item, status in zip(files, ["running", "done", "error"]):
        task_store.set(item["task_id"], {"status": status})
    task_store.set_job("filter-batch", {"job_id": "filter-batch", "files": files})
    task_store.set_job("filter-other", {"job_id": "filter-other", "files": [files[1]]})
    page.set_viewport_size({"width": 375, "height": 812})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    _open_batch_history(page, "filter-batch")
    page.locator('[data-open-job="filter-batch"]').click()
    page.wait_for_function("document.getElementById('batchFilterSummary').textContent.includes('4 из 4')")
    if not page.locator("#batchFilters").evaluate("el => el.open"):
        page.locator("#batchFilters > summary").click()
    page.select_option("#batchStateFilter", "attention")
    assert page.locator("#batchProgressList > li").count() == 2
    task_store.set("filter-running", {"status": "error"})
    page.wait_for_function("document.getElementById('batchFilterSummary').textContent.includes('3 из 4')")
    assert page.locator("#batchStateFilter").input_value() == "attention"
    if not page.locator("#batchFilters").evaluate("el => el.open"):
        page.locator("#batchFilters > summary").click()
    page.select_option("#batchStateFilter", "done")
    page.locator("#batchNameFilter").fill("ОТЧЁТ")
    assert page.locator("#batchProgressList > li").count() == 1
    page.locator("#batchNameFilter").fill("Несуществующий")
    assert page.locator("#batchFilterEmpty").is_visible()
    assert page.locator("#batchArchiveBtn").is_visible()
    page.locator("#batchFilterReset").click()
    assert page.locator("#batchProgressList > li").count() == 4
    page.locator("#batchNameFilter").fill("Ошибка")
    page.locator('[data-job="filter-other"] > summary').click()
    _open_batch_history(page, "filter-other")
    page.locator('[data-open-job="filter-other"]').click()
    page.wait_for_function("document.getElementById('batchFilterSummary').textContent.includes('1 из 1')")
    assert page.locator("#batchNameFilter").input_value() == ""
    assert page.locator("#batchStateFilter").input_value() == "all"
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


def test_e2e_individual_batch_formats(e2e_server, page, task_store):
    page.set_viewport_size({"width": 375, "height": 812})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files(
        [
            {"name": "same.txt", "mimeType": "text/plain", "buffer": b"First"},
            {"name": "same.txt", "mimeType": "text/plain", "buffer": b"Second"},
        ]
    )
    page.locator("#expertConversionBtn").click()
    page.locator("#batchIndividualOptions > summary").click()
    page.select_option("#batchTarget0", "model")
    page.select_option("#batchMode0", "editable")
    page.select_option("#target", "html")
    assert page.locator("#batchTarget0").input_value() == "model"
    assert page.locator("#batchTarget1").input_value() == "html"
    page.locator("#batchApplyDefaults").click()
    assert page.locator("#batchTarget0").input_value() == "html"
    page.select_option("#batchTarget0", "model")
    page.select_option("#batchMode0", "editable")
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.locator("#batchConvertBtn").click()
    created = response.value.json()
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    first, second = [task_store.get(item["task_id"]) for item in created["tasks"]]
    assert (first["target_format"], first["mode"], first["status"]) == ("model", "editable", "done")
    assert (second["target_format"], second["status"]) == ("html", "done")
    assert page.e2e_errors == []


@pytest.mark.parametrize("width", [375, 1280])
def test_e2e_batch_storage_failure_keeps_inputs_and_options(e2e_server, page, task_store, monkeypatch, width):
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files([
        {"name": name, "mimeType": "text/plain", "buffer": name.encode()} for name in ("first.txt", "second.txt")
    ])
    page.locator("#expertConversionBtn").click()
    page.select_option("#target", "html")
    page.locator("#batchIndividualOptions > summary").click()
    page.select_option("#batchTarget0", "model")
    page.select_option("#batchMode0", "editable")
    writer = task_store.set_job

    def fail_job(job_id, payload):
        writer(job_id, payload)
        raise OSError("temporary batch storage failure")

    with monkeypatch.context() as failure:
        failure.setattr(task_store, "set_job", fail_job)
        with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
            page.locator("#batchConvertBtn").click()
        assert response.value.status == 500
        page.wait_for_function("!document.getElementById('batchConvertBtn').disabled")
        assert task_store.list_jobs() == [] and task_store.list_tasks() == []
        assert page.locator("#fileInput").evaluate("el => Array.from(el.files).map(file => file.name)") == [
            "first.txt", "second.txt",
        ]
        assert page.locator("#batchTarget0").input_value() == "model"
        assert page.locator("#batchMode0").input_value() == "editable"
        assert page.locator("#batchTarget1").input_value() == "html"
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.locator("#batchConvertBtn").click()
    assert response.value.status == 200
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    assert all(task_store.get(item["task_id"])["status"] == "done" for item in response.value.json()["tasks"])
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), page.evaluate(
        "Array.from(document.querySelectorAll('body *')).filter(e => e.getBoundingClientRect().right > innerWidth + 1)"
        ".slice(0, 15).map(e => [e.tagName, e.className, e.id, e.getBoundingClientRect().width])"
    )
    assert all("500 (Internal Server Error)" in error for error in page.e2e_errors)


@pytest.mark.parametrize("width", [375, 1280])
@pytest.mark.parametrize("failure", ["storage", "queue"])
def test_e2e_retry_storage_failure_shows_reason_and_keeps_results(e2e_server, page, task_store, monkeypatch, width, failure):
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files([
        {"name": name, "mimeType": "text/plain", "buffer": name.encode()} for name in ("good.txt", "retry.txt")
    ])
    page.select_option("#target", "model")
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.locator("#batchConvertBtn").click()
    created = response.value.json()
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    good_id, retry_id = [item["task_id"] for item in created["tasks"]]
    task_store.set(retry_id, {**task_store.get(retry_id), "status": "error", "error": "Временная ошибка"})
    originals = {task_id: task_store.get(task_id) for task_id in (good_id, retry_id)}
    results = {task_id: task_store.result_path(task_id, task["artifact"]).read_bytes() for task_id, task in originals.items()}
    writer = task_store._write_meta_locked

    def fail_retry(task_id, payload):
        if task_id == retry_id and payload.get("status") == "queued":
            raise OSError("temporary storage failure")
        writer(task_id, payload)

    if failure == "storage":
        monkeypatch.setattr(task_store, "_write_meta_locked", fail_retry)
    else:
        def fail_queue(_task_id):
            raise RuntimeError("queue unavailable")
        monkeypatch.setattr(convert_route, "resume_conversion_task", fail_queue)
    _open_batch_history(page)
    page.locator("#historyRefresh").click()
    page.wait_for_function("document.querySelector('#historyList .job-state').textContent.includes('1 ошибок')")
    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.locator("#historyList details > summary").first.click()
    page.get_by_role("checkbox", name="Выбрать для повтора: retry.txt", exact=True).check()
    with page.expect_response(lambda response: response.url.endswith("/rerun")) as response:
        page.locator("[data-retry-selected]").click()
    assert response.value.json()["launched"] == []
    if failure == "storage":
        page.get_by_text("Не удалось сохранить повтор; прежний результат сохранён.", exact=False).wait_for(state="visible")
    else:
        page.locator("#batchProgressList").get_by_text("Ожидает перезапуска очереди", exact=False).wait_for(state="visible")
        assert response.value.json()["queued"] == [retry_id]
        assert task_store.get(retry_id)["status"] == "queued"
        assert task_store.source_path(retry_id).read_bytes() == b"retry.txt"
    for task_id, task in originals.items():
        if failure == "queue" and task_id == retry_id:
            continue
        assert task_store.get(task_id) == task
        assert task_store.result_path(task_id, task["artifact"]).read_bytes() == results[task_id]
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


def test_e2e_batch_without_common_target(e2e_server, page):
    from textalchemy.convert.executor import ConversionExecutor
    from textalchemy.web.services.conversion_catalog import available_conversions

    catalog = available_conversions(ConversionExecutor(requirement_checker=lambda _: True))
    for source in catalog["sources"]:
        if source["format"] in {"txt", "model"}:
            target = "model" if source["format"] == "txt" else "docx"
            source["targets"] = [item for item in source["targets"] if item["format"] == target]
            source["default_target"] = target
    page.route("**/api/convert/capabilities", lambda route: route.fulfill(json=catalog))
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files(
        [
            {"name": "text.txt", "mimeType": "text/plain", "buffer": b"Text"},
            {"name": "model.json", "mimeType": "application/json", "buffer": b"{}"},
        ]
    )
    assert page.locator("#batchIndividualOptions").get_attribute("open") is not None
    assert page.locator("#batchTarget0").input_value() == "model"
    assert page.locator("#batchTarget1").input_value() == "docx"
    assert page.locator("#batchConvertBtn").is_visible()
    assert page.e2e_errors == []


def test_e2e_retry_only_failed_files(e2e_server, page, task_store):
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files(
        [
            {"name": "good.txt", "mimeType": "text/plain", "buffer": b"Good"},
            {"name": "retry.txt", "mimeType": "text/plain", "buffer": b"Retry"},
        ]
    )
    page.select_option("#target", "model")
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.locator("#batchConvertBtn").click()
    created = response.value.json()
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    good_id, retry_id = [item["task_id"] for item in created["tasks"]]
    good = task_store.get(good_id)
    task_store.set(retry_id, {**task_store.get(retry_id), "status": "error", "error": "Временная ошибка"})
    _open_batch_history(page)
    page.locator("#historyRefresh").click()
    page.wait_for_function("document.querySelector('#historyList .job-state').textContent.includes('ошибок')")
    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.locator("#historyList details > summary").first.click()
    with page.expect_response(lambda response: response.url.endswith("/rerun")) as response:
        page.get_by_role("button", name="Повторить неудачные", exact=True).click()
    assert response.value.json()["launched"] == [retry_id]
    assert response.value.json()["failed_only"] is True
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    assert task_store.get(good_id) == good
    assert task_store.get(retry_id)["status"] == "done"
    assert page.e2e_errors == []


@pytest.mark.parametrize("width", [375, 1280])
@pytest.mark.parametrize("status_changed", [False, True])
def test_e2e_retry_selected_file(e2e_server, page, task_store, tmp_path, width, status_changed):
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    page.locator("#fileInput").set_input_files([
        {"name": name, "mimeType": "text/plain", "buffer": name.encode()}
        for name in ("good.txt", "retry.txt", "leave.txt")
    ])
    page.select_option("#target", "model")
    with page.expect_response(lambda response: response.url.endswith("/api/convert/batch")) as response:
        page.locator("#batchConvertBtn").click()
    created = response.value.json()
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    good_id, retry_id, leave_id = [item["task_id"] for item in created["tasks"]]
    good = task_store.get(good_id)
    good_bytes = task_store.result_path(good_id, good["artifact"]).read_bytes()
    for task_id in (retry_id, leave_id):
        task_store.set(task_id, {**task_store.get(task_id), "status": "error", "error": "Временная ошибка"})
    untouched = task_store.get(leave_id)
    untouched_bytes = task_store.result_path(leave_id, untouched["artifact"]).read_bytes()
    _open_batch_history(page)
    page.locator("#historyRefresh").click()
    page.wait_for_function("document.querySelector('#historyList .job-state').textContent.includes('2 ошибок')")
    if not page.locator("#conversionHistory").evaluate("el => el.open"):
        page.locator("#conversionHistory > summary").click()
    page.locator("#historyList details > summary").first.click()
    button = page.locator('[data-retry-selected]')
    button.wait_for(state="visible")
    assert button.is_disabled()
    assert page.locator('[data-retry-task]').count() == 2
    checkbox = page.get_by_role("checkbox", name="Выбрать для повтора: retry.txt", exact=True)
    checkbox.focus()
    checkbox.press("Space")
    assert checkbox.is_checked()
    assert not button.is_disabled()
    assert button.inner_text() == "Повторить выбранные (1)"
    page.screenshot(path=str(tmp_path / f"selected-retry-{width}.png"), full_page=True)
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth"), page.evaluate(
        "Array.from(document.querySelectorAll('body *')).filter(e => e.getBoundingClientRect().right > innerWidth + 1)"
        ".slice(0, 12).map(e => [e.tagName, e.className, e.id, e.getBoundingClientRect().width])"
    )


    if status_changed:
        task_store.set(retry_id, {**task_store.get(retry_id), "status": "done"})
    with page.expect_response(lambda response: response.url.endswith("/rerun")) as response:
        button.click()
    result = response.value.json()
    assert result["launched"] == ([] if status_changed else [retry_id])
    if status_changed:
        assert result["skipped"][0]["task_id"] == retry_id
        page.get_by_text("Пропущены: retry.txt:", exact=False).wait_for(state="visible")
    else:
        page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
        assert task_store.get(retry_id)["status"] == "done"
    assert task_store.get(good_id) == good
    assert task_store.result_path(good_id, good["artifact"]).read_bytes() == good_bytes
    assert task_store.get(leave_id) == untouched
    assert task_store.result_path(leave_id, untouched["artifact"]).read_bytes() == untouched_bytes
    assert page.e2e_errors == []


@pytest.mark.parametrize("surface", ["panel", "history"])
def test_e2e_cancel_remaining_batch_files(e2e_server, page, task_store, surface):
    task_store.set("cancel-pending", {"status": "queued"})
    task_store.set("cancel-ready", {"status": "done"})
    ready = task_store.get("cancel-ready")
    task_store.set_job(
        "cancel-batch",
        {
            "job_id": "cancel-batch",
            "files": [
                {"task_id": "cancel-pending", "name": "pending.txt"},
                {"task_id": "cancel-ready", "name": "ready.txt"},
            ],
        },
    )
    page.set_viewport_size({"width": 375, "height": 812})
    page.goto(f"{e2e_server}/convert")
    _wait_convert_ready(page)
    if surface == "panel":
        _open_batch_history(page, "cancel-batch")
        page.locator('[data-open-job="cancel-batch"]').click()
        button = page.locator("#batchCancelBtn")
    else:
        _open_batch_history(page, "cancel-batch")
        button = page.locator('[data-cancel-job="cancel-batch"]')
    with page.expect_response(lambda response: response.url.endswith("/cancel")) as response:
        button.click()
    assert response.value.json()["requested"] == [{"task_id": "cancel-pending", "status": "cancelled"}]
    page.wait_for_function("document.getElementById('batchProgressState').textContent === 'Готово'")
    assert page.locator("#batchCancelBtn").is_hidden()
    assert page.locator("#batchArchiveBtn").is_visible()
    assert task_store.get("cancel-ready") == ready
    assert page.locator("#batchProgressList .cancelled").count() > 0
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


@pytest.mark.parametrize("width", [375, 1280])
def test_e2e_ocr_edit_reload_and_export(e2e_server, page, task_store, tmp_path, width):
    from docx import Document
    from playwright.sync_api import expect

    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/recognize")
    page.locator("#scenario").select_option("fast")
    page.set_input_files("#fileInput", _make_pdf(tmp_path / "ocr.pdf", "0CR err0r"))
    page.locator("#processBtn").click()
    expect(page.locator("#result")).to_have_value("0CR err0r", timeout=30000)
    page.locator("#result").fill("Исправленный текст\n\nСтрока 2")
    page.locator("#saveTextBtn").click()
    expect(page.locator("#editStatus")).to_have_text("Текст сохранён")
    assert "draft=" in page.url
    page.reload()
    expect(page.locator("#result")).to_have_value("Исправленный текст\n\nСтрока 2")
    page.locator("#result").fill("Последняя правка\n\nСтрока 2")
    page.locator("#ocrExportFormat").select_option("docx")
    with page.expect_download() as event:
        page.locator("#exportTextBtn").click()
    target = tmp_path / "corrected.docx"
    event.value.save_as(target)
    assert "\n".join(p.text for p in Document(target).paragraphs) == "Последняя правка\n\nСтрока 2"
    page.reload()
    expect(page.locator("#result")).to_have_value("Последняя правка\n\nСтрока 2")
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth")
    assert page.e2e_errors == []


def test_e2e_ocr_conflict_keeps_local_edit(e2e_server, page, task_store, tmp_path):
    from playwright.sync_api import expect

    page.goto(f"{e2e_server}/recognize")
    page.locator("#scenario").select_option("fast")
    page.set_input_files("#fileInput", _make_pdf(tmp_path / "ocr.pdf", "Original"))
    page.locator("#processBtn").click()
    expect(page.locator("#result")).to_have_value("Original", timeout=30000)
    draft_id = page.url.split("draft=")[1]
    response = page.request.put(f"{e2e_server}/api/recognize/drafts/{draft_id}",
                                data={"revision": 1, "text": "Other tab"})
    assert response.ok
    page.locator("#result").fill("Local unsaved edit")
    page.locator("#exportTextBtn").click()
    expect(page.locator("#status")).to_contain_text("другой вкладке")
    expect(page.locator("#result")).to_have_value("Local unsaved edit")
    expect(page.locator("#saveTextBtn")).to_be_enabled()
    assert page.request.get(f"{e2e_server}/api/recognize/drafts/{draft_id}").json()["text"] == "Other tab"
    assert all("409" in error for error in page.e2e_errors)


@pytest.mark.parametrize("width", [375, 1280])
def test_e2e_ocr_paragraph_order(e2e_server, page, task_store, tmp_path, width):
    from docx import Document
    from playwright.sync_api import expect

    page.set_viewport_size({"width": width, "height": 900})
    page.goto(f"{e2e_server}/recognize")
    page.locator("#scenario").select_option("fast")
    page.set_input_files("#fileInput", _make_pdf(tmp_path / "order.pdf", "Second\nFirst\nThird"))
    page.locator("#processBtn").click()
    result = page.locator("#result")
    expect(result).to_have_value("Second\nFirst\nThird", timeout=30000)
    page.locator("#paragraphTools summary").click()
    result.focus()
    page.keyboard.press("Control+Home")
    expect(page.locator("#paragraphUp")).to_be_disabled()
    page.keyboard.press("ArrowDown")
    page.locator("#paragraphUp").focus()
    page.keyboard.press("Enter")
    expect(result).to_have_value("First\nSecond\nThird")
    expect(page.locator("#editStatus")).to_have_text("Есть несохранённые правки")
    page.locator("#ocrExportFormat").select_option("docx")
    with page.expect_download() as event:
        page.locator("#exportTextBtn").click()
    target = tmp_path / "ordered.docx"
    event.value.save_as(target)
    assert [p.text for p in Document(target).paragraphs] == ["First", "Second", "Third"]
    page.reload()
    expect(result).to_have_value("First\nSecond\nThird")
    draft_id = page.url.split("draft=")[1]
    model = page.request.get(f"{e2e_server}/api/recognize/drafts/{draft_id}/export?revision=2&format=model").json()
    from textalchemy.core.document_codec import document_from_dict

    assert [b.plain_text for b in document_from_dict(model).sections[0].blocks] == ["First", "Second", "Third"]
    page.locator("#paragraphTools summary").click()
    # Repeated text and empty paragraphs must move by position without loss.
    result.fill("Same\n\nSame\nLast")
    result.evaluate("e => {e.focus(); e.setSelectionRange(0, 6); e.dispatchEvent(new Event('select'));}")
    page.locator("#paragraphDown").click()
    expect(result).to_have_value("Same\nSame\n\nLast")
    page.locator("#paragraphDown").click()
    expect(result).to_have_value("Same\nLast\nSame\n")
    expect(page.locator("#paragraphDown")).to_be_disabled()
    page.locator("#paragraphUp").click()
    expect(result).to_have_value("Same\nSame\n\nLast")
    page.locator("#saveTextBtn").click()
    expect(page.locator("#editStatus")).to_have_text("Текст сохранён")
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    assert page.e2e_errors == []


def test_e2e_ocr_order_before_trailing_empty_paragraph(e2e_server, page, task_store, tmp_path):
    from playwright.sync_api import expect

    page.goto(f"{e2e_server}/recognize")
    page.locator("#scenario").select_option("fast")
    page.set_input_files("#fileInput", _make_pdf(tmp_path / "order-empty.pdf", "Original"))
    page.locator("#processBtn").click()
    result = page.locator("#result")
    expect(result).to_have_value("Original", timeout=30000)
    result.fill("A\nB\n")
    page.locator("#paragraphTools summary").click()
    result.evaluate("e => {e.setSelectionRange(0, 0); e.dispatchEvent(new Event('select'));}")
    page.locator("#paragraphDown").click()
    expect(result).to_have_value("B\nA\n")
    expect(page.locator("#paragraphDown")).to_be_enabled()
    page.locator("#paragraphDown").click()
    expect(result).to_have_value("B\n\nA")
    expect(page.locator("#paragraphDown")).to_be_disabled()
    page.locator("#saveTextBtn").click()
    expect(page.locator("#editStatus")).to_have_text("Текст сохранён")
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_pdf_region_order(e2e_server, page, task_store, tmp_path, width):
    import fitz
    from docx import Document
    from playwright.sync_api import expect

    source = tmp_path / 'regions.pdf'
    pdf = fitz.open()
    for words in [('Second', 'First'), ('Unchanged', 'Page two')]:
        sheet = pdf.new_page()
        sheet.insert_text((70, 150), words[0], fontsize=14)
        sheet.insert_text((70, 250), words[1], fontsize=18)
    pdf.save(source)
    pdf.close()
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/pdf-order')
    page.set_input_files('#orderFile', source)
    expect(page.locator('#orderBlocks option')).to_have_count(2)
    page.wait_for_function("document.getElementById('orderImage').naturalWidth > 0")
    expect(page.locator('#orderRegion')).to_be_visible()
    page.locator('#orderBlocks').focus()
    page.keyboard.press('ArrowDown')
    page.locator('#orderUp').focus()
    page.keyboard.press('Enter')
    expect(page.locator('#orderBlocks option').first).to_have_text('1. First')
    page.locator('#orderPage').select_option(index=1)
    assert page.e2e_errors == [], page.locator('#orderPage').input_value()
    assert page.locator('#orderImage').get_attribute('src').endswith('/pages/1')
    expect(page.locator('#orderBlocks option').first).to_have_text('1. Unchanged')
    page.locator('#orderFormat').select_option('docx')
    with page.expect_download() as event:
        page.locator('#orderExport').click()
    output = tmp_path / 'ordered.docx'
    event.value.save_as(output)
    text = [p.text for p in Document(output).paragraphs]
    assert text.index('First') < text.index('Second') < text.index('Unchanged')
    page.reload()
    expect(page.locator('#orderBlocks option').first).to_have_text('1. First')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    page.screenshot(path=str(tmp_path / f'pdf-order-{width}.png'), full_page=True)
    assert page.e2e_errors == []


def test_e2e_pdf_order_conflict(e2e_server, page, task_store, tmp_path):
    from playwright.sync_api import expect

    page.goto(f'{e2e_server}/pdf-order')
    import fitz

    source = tmp_path / 'separate.pdf'
    pdf = fitz.open()
    sheet = pdf.new_page()
    sheet.insert_text((70, 150), 'Second')
    sheet.insert_text((70, 250), 'First')
    pdf.save(source)
    pdf.close()
    page.set_input_files('#orderFile', source)
    expect(page.locator('#orderBlocks option')).to_have_count(2)
    draft_id = page.url.split('draft=')[1]
    url = f'{e2e_server}/api/pdf-order/{draft_id}'
    stored = page.request.get(url).json()
    order = [[block['id'] for block in section['blocks']] for section in stored['pages']]
    assert page.request.put(url, data={'revision': 1, 'order': order}).ok
    page.locator('#orderDown').click()
    page.locator('#blockClassification').select_option('heading3')
    page.locator('#blockClassification').select_option('heading3')
    page.locator('#orderSave').click()
    expect(page.locator('#orderStatus')).to_contain_text('другой вкладке')
    expect(page.locator('#orderBlocks option').first).to_have_text('1. First')
    expect(page.locator('#orderSave')).to_be_enabled()
    expect(page.locator('#blockClassification')).to_have_value('heading3')
    expect(page.locator('#blockClassification')).to_have_value('heading3')
    assert page.request.get(url).json()['pages'][0]['blocks'][0]['label'] == 'Second'
    assert all('409' in error for error in page.e2e_errors)


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_pdf_classification(e2e_server, page, task_store, tmp_path, width):
    import fitz
    from docx import Document
    from playwright.sync_api import expect

    source = tmp_path / 'classification.pdf'
    pdf = fitz.open()
    sheet = pdf.new_page()
    sheet.insert_text((70, 150), 'Chapter title', fontsize=12)
    sheet.insert_text((70, 250), 'Body text', fontsize=12)
    pdf.save(source)
    pdf.close()
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/pdf-order')
    page.set_input_files('#orderFile', source)
    expect(page.locator('#orderBlocks option')).to_have_count(2)
    page.locator('#blockClassification').select_option('heading2')
    page.locator('#orderDown').click()
    expect(page.locator('#blockClassification')).to_have_value('heading2')
    page.locator('#orderFormat').select_option('docx')
    with page.expect_download() as event:
        page.locator('#orderExport').click()
    output = tmp_path / 'heading.docx'
    event.value.save_as(output)
    title = next(p for p in Document(output).paragraphs if p.text == 'Chapter title')
    assert title.style.name == 'Heading 2'
    page.reload()
    expect(page.locator('#orderBlocks option')).to_have_count(2)
    page.locator('#orderBlocks').select_option(index=1)
    expect(page.locator('#blockClassification')).to_have_value('heading2')
    page.locator('#blockClassification').select_option('paragraph')
    page.locator('#orderFormat').select_option('docx')
    with page.expect_download() as event:
        page.locator('#orderExport').click()
    event.value.save_as(output)
    title = next(p for p in Document(output).paragraphs if p.text == 'Chapter title')
    assert title.style.name == 'Normal'
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_pdf_table_bounds(e2e_server, page, task_store, tmp_path, width):
    from docx import Document
    from playwright.sync_api import expect

    from tests.test_web_pdf_tables import make_table_pdf

    source = tmp_path / 'table.pdf'
    source.write_bytes(make_table_pdf())
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/pdf-order')
    page.set_input_files('#orderFile', source)
    page.locator('#orderBlocks option').filter(has_text='Таблица').wait_for(state='attached')
    table_id = page.locator('#orderBlocks option').filter(has_text='Таблица').get_attribute('value')
    page.locator('#orderBlocks').select_option(table_id)
    expect(page.locator('#tableBounds')).to_be_visible()
    expect(page.locator('#tablePreview tr')).to_have_count(3)
    page.locator('#tableFirstRow').fill('2')
    page.locator('#tableLastColumn').fill('2')
    expect(page.locator('#tablePreview tr')).to_have_count(2)
    expect(page.locator('#tablePreview tr').first.locator('td')).to_have_text(['R2C1', 'R2C2'])
    page.locator('#orderFormat').select_option('docx')
    with page.expect_download() as event:
        page.locator('#orderExport').click()
    output = tmp_path / 'cropped.docx'
    event.value.save_as(output)
    doc = Document(output)
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == [['R2C1', 'R2C2'], ['R3C1', 'R3C2']]
    page.reload()
    expect(page.locator('#orderBlocks option').filter(has_text='Таблица')).to_have_count(1)
    page.locator('#orderBlocks').select_option(table_id)
    expect(page.locator('#tableFirstRow')).to_have_value('2')
    expect(page.locator('#tableLastColumn')).to_have_value('2')
    page.locator('#tableReset').click()
    expect(page.locator('#tablePreview tr')).to_have_count(3)
    page.locator('#orderSave').click()
    expect(page.locator('#orderDirty')).to_have_text('Правки сохранены')
    page.locator('#orderBlocks').select_option(table_id)
    expect(page.locator('#tableFirstRow')).to_have_value('1')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_pdf_warning_navigation(e2e_server, page, task_store, tmp_path, width):
    from playwright.sync_api import expect

    from tests.test_web_pdf_diagnostics import make_warning_pdf

    source = tmp_path / 'warning.pdf'
    source.write_bytes(make_warning_pdf())
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/pdf-order')
    page.set_input_files('#orderFile', source)
    expect(page.locator('#orderEditor')).to_be_visible()
    page.locator('#orderCheck').click()
    warning = page.locator('#pdfIssueList button').filter(has_text='Vector drawing').first
    expect(warning).to_be_visible()
    warning.focus()
    page.keyboard.press('Enter')
    expect(page.locator('#orderPage')).to_have_value('1')
    expect(page.locator('#orderRegion')).to_be_visible()
    expect(page.locator('#orderBlocks option:checked')).to_contain_text('Vector drawing')
    page.wait_for_function("document.getElementById('orderImage').complete && "
                           "document.getElementById('orderImage').naturalWidth > 0")
    assert page.locator('#orderImage').get_attribute('src').endswith('/pages/1')
    block_id = page.locator('#orderBlocks').input_value()
    assert 'block=' + block_id in page.url
    expect(page.locator('#orderBlocks')).to_be_focused()
    page.reload()
    expect(page.locator('#orderBlocks')).to_have_value(block_id)
    expect(page.locator('#orderPage')).to_have_value('1')
    expect(page.locator('#orderRegion')).to_be_visible()
    page.locator('#orderCheck').click()
    expect(warning).to_be_visible()
    page.locator('#orderUp').click()
    expect(page.locator('#pdfDiagnostics')).to_be_hidden()
    page.locator('#orderCheck').click()
    expect(warning).to_be_visible()
    warning.click()
    expect(page.locator('#orderBlocks')).to_have_value(block_id)
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_generator_draft_restores_and_generates(e2e_server, page, width):
    import io

    from docx import Document
    from playwright.sync_api import expect

    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    expect(page.locator('#field-title')).to_be_visible(timeout=30000)
    page.locator('#field-title').fill('Сохранённый заголовок')
    page.locator('#field-author').fill('Автор черновика')
    page.locator('#field-body').fill('Текст после восстановления')
    page.locator('#output').fill('мой-документ.docx')
    expect(page.locator('#draftStatus')).to_contain_text('сохранён')
    page.reload()
    expect(page.locator('#field-title')).to_have_value('')
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-title')).to_have_value('Сохранённый заголовок')
    expect(page.locator('#output')).to_have_value('мой-документ.docx')
    expect(page.locator('#draftStatus')).to_contain_text('восстановлен')
    with page.expect_download(timeout=60000) as event:
        page.locator('#genBtn').click()
    doc = Document(io.BytesIO(event.value.path().read_bytes()))
    assert 'Сохранённый заголовок' in '\n'.join(p.text for p in doc.paragraphs)
    assert 'Текст после восстановления' in '\n'.join(p.text for p in doc.paragraphs)
    # A second tab in the same browser session must not overwrite the first tab's draft.
    other = page.context.new_page()
    other.goto(f'{e2e_server}/generate')
    expect(other.locator('#field-title')).to_be_visible(timeout=30000)
    other.locator('#field-title').fill('Другая вкладка')
    page.reload()
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-title')).to_have_value('Сохранённый заголовок')
    other.close()
    page.once('dialog', lambda dialog: dialog.accept())
    page.locator('#savedData > summary').click()
    page.locator('#clearDraft').click()
    expect(page.locator('#field-title')).to_have_value('')
    page.reload()
    expect(page.locator('#field-title')).to_have_value('')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


def test_e2e_generator_invalid_json_draft(e2e_server, page):
    from playwright.sync_api import expect

    schema = {'fields': [{'name': 'items', 'type': 'array'}, {'name': 'enabled', 'type': 'boolean', 'default': True}]}
    page.route('**/api/generate/templates', lambda route: route.fulfill(json=[{'name': 'draft-test'}]))
    page.route('**/api/generate/templates/draft-test/schema', lambda route: route.fulfill(
        json={'schema': schema, 'description': 'Test', 'source': 'sidecar'}))
    page.route('**/api/generate/templates/draft-test/preview/meta', lambda route: route.fulfill(json={'available': False}))
    page.goto(f'{e2e_server}/generate')
    page.locator('[data-field-name=items]').get_by_role('button', name='JSON', exact=True).click()
    page.locator('#field-items').fill('[unfinished')
    page.locator('#field-enabled').uncheck()
    page.locator('[data-format="html"]').click()
    page.reload()
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-items')).to_have_value('[unfinished')
    expect(page.locator('#field-enabled')).not_to_be_checked()
    expect(page.locator('[data-format="html"]')).to_have_class('format-option active')
    expect(page.locator('#output')).to_have_value('output.html')
    page.locator('#genBtn').click()
    expect(page.locator('#field-error-items')).to_contain_text('некорректный JSON')
    page.reload()
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-items')).to_have_value('[unfinished')
    assert page.e2e_errors == []


def test_e2e_generator_template_isolation_images_and_schema_change(e2e_server, page):
    import base64

    from playwright.sync_api import expect

    schema = {'fields': [{'name': 'title', 'type': 'string'}, {'name': 'picture', 'type': 'image'}]}
    page.route('**/api/generate/templates', lambda route: route.fulfill(json=[{'name': 'A'}, {'name': 'B'}]))
    page.route('**/api/generate/templates/*/schema', lambda route: route.fulfill(
        json={'schema': schema, 'description': 'Test', 'source': 'sidecar'}))
    page.route('**/api/generate/templates/*/preview/meta', lambda route: route.fulfill(json={'available': False}))
    page.goto(f'{e2e_server}/generate')
    page.locator('#field-title').fill('Template A')
    image = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jN1EAAAAASUVORK5CYII=')
    page.locator('#field-picture').set_input_files({'name': 'pixel.png', 'mimeType': 'image/png', 'buffer': image})
    expect(page.locator('#draftStatus')).to_contain_text('Черновик сохранён')
    page.locator('#template').select_option('B')
    expect(page.locator('#field-title')).to_have_value('')
    page.locator('#field-title').fill('Template B')
    page.reload()
    expect(page.locator('#template')).to_have_value('B')
    expect(page.locator('#field-title')).to_have_value('')
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-title')).to_have_value('Template B')
    page.locator('#template').select_option('A')
    expect(page.locator('#field-title')).to_have_value('')
    assert not page.locator('#field-picture').evaluate('e => e.dataset.imageData')
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-title')).to_have_value('Template A')
    assert page.locator('#field-picture').evaluate('e => e.dataset.imageData').startswith('data:image/png;base64,')
    expect(page.get_by_text('Изображение восстановлено из черновика.', exact=False)).to_be_visible()
    saved = page.evaluate("sessionStorage.getItem('textalchemy.generator-draft.v1.A')")
    schema['fields'][0]['type'] = 'integer'
    page.reload()
    page.locator('#restoreDraft').click()
    expect(page.locator('#draftStatus')).to_contain_text('схема шаблона изменилась')
    page.locator('#field-title').fill('42')
    assert page.evaluate("sessionStorage.getItem('textalchemy.generator-draft.v1.A')") == saved
    assert page.e2e_errors == []


def test_e2e_generator_storage_failure_keeps_input(e2e_server, page):
    from playwright.sync_api import expect

    page.add_init_script("""const original = Storage.prototype.setItem;
        Storage.prototype.setItem = function(key, value) {
            if (key.startsWith('textalchemy.generator-draft')) throw new DOMException('Full', 'QuotaExceededError');
            return original.call(this, key, value);
        };""")
    page.goto(f'{e2e_server}/generate')
    page.locator('#field-title').fill('Не потерять ввод')
    expect(page.locator('#draftStatus')).to_contain_text('Черновик не сохранён')
    expect(page.locator('#field-title')).to_have_value('Не потерять ввод')
    expect(page.locator('#genBtn')).to_be_enabled()
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_generator_named_dataset(e2e_server, page, tmp_path, monkeypatch, width):
    import importlib
    import io

    from docx import Document
    from playwright.sync_api import expect

    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path)
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    page.locator('#field-title').fill('Сохранённый набор')
    page.locator('#field-author').fill('Автор набора')
    page.locator('#field-body').fill('Содержимое набора')
    page.locator('#output').fill('dataset.docx')
    page.locator('#savedData > summary').click()
    page.locator('#datasetName').fill('Исследование')
    page.locator('#datasetCreate').click()
    expect(page.locator('#datasetStatus')).to_contain_text('версия 1')
    dataset_id = page.locator('#datasetList').input_value()
    template = page.locator('#template').input_value()
    # Fresh tab has no session draft but can load the persistent server dataset.
    other = page.context.new_page()
    other.goto(f'{e2e_server}/generate')
    expect(other.locator('#field-title')).to_have_value('')
    expect(other.locator('#datasetList option')).to_have_count(2)
    other.locator('#savedData > summary').click()
    other.locator('#datasetList').select_option(dataset_id)
    expect(other.locator('#datasetUpdate')).to_be_disabled()
    other.once('dialog', lambda dialog: dialog.accept())
    other.locator('#datasetLoad').click()
    expect(other.locator('#field-title')).to_have_value('Сохранённый набор')
    expect(other.locator('#output')).to_have_value('dataset.docx')
    other.locator('#field-title').fill('Обновлённый набор')
    other.locator('#datasetUpdate').click()
    expect(other.locator('#datasetStatus')).to_contain_text('версия 2')
    with other.expect_download(timeout=60000) as event:
        other.locator('#genBtn').click()
    doc = Document(io.BytesIO(event.value.path().read_bytes()))
    assert 'Обновлённый набор' in '\n'.join(p.text for p in doc.paragraphs)
    other.close()
    # The first tab still holds revision 1: its update must preserve both copies.
    page.locator('#field-title').fill('Мои несохранённые правки')
    page.locator('#datasetUpdate').click()
    expect(page.locator('#datasetStatus')).to_contain_text('изменён в другой вкладке')
    expect(page.locator('#field-title')).to_have_value('Мои несохранённые правки')
    saved = page.request.get(f'{e2e_server}/api/generate/datasets/{dataset_id}').json()
    assert saved['revision'] == 2
    assert {item['name']: item['value'] for item in saved['snapshot']['values']}['title'] == 'Обновлённый набор'
    page.locator('#datasetName').fill('Моя копия')
    page.locator('#datasetCreate').click()
    expect(page.locator('#datasetStatus')).to_contain_text('версия 1')
    assert page.locator('#datasetList').input_value() != dataset_id
    assert len(page.request.get(f'{e2e_server}/api/generate/datasets', params={'template': template}).json()) == 2
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert all('409' in error for error in page.e2e_errors)


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_template_variable_copy(e2e_server, page, tmp_path, monkeypatch, width):
    import importlib
    import io

    from docx import Document
    from playwright.sync_api import expect

    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path)
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    expect(page.locator('#field-title')).to_be_visible()
    page.locator('#editMode').click()
    page.get_by_text('Изменить переменную шаблона', exact=True).click()
    page.locator('#variableInspect').click()
    expect(page.locator('#variableSelect option').filter(has_text='title')).to_have_count(1)
    page.locator('#variableSelect').select_option('title')
    expect(page.locator('#variableContext')).to_contain_text('title')
    page.locator('#variableName').fill('subject')
    page.locator('#variableSave').click()
    expect(page.locator('#field-subject')).to_be_visible()
    copied = page.locator('#template').input_value()
    assert copied.startswith('edited-')
    page.reload()
    expect(page.locator('#template')).to_have_value(copied)
    page.locator('#field-subject').fill('Изменённая переменная')
    page.locator('#field-author').fill('Автор')
    page.locator('#field-body').fill('Текст')
    with page.expect_download(timeout=60000) as event:
        page.locator('#genBtn').click()
    doc = Document(io.BytesIO(event.value.path().read_bytes()))
    assert 'Изменённая переменная' in '\n'.join(p.text for p in doc.paragraphs)
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_template_condition_copy(e2e_server, page, tmp_path, monkeypatch, width):
    import importlib
    import io

    from docx import Document
    from playwright.sync_api import expect

    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path)
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    expect(page.locator('#field-title')).to_be_visible()
    page.locator('#editMode').click()
    page.get_by_text('Настроить условный абзац', exact=True).click()
    page.locator('#conditionInspect').click()
    option = page.locator('#conditionBlock option').filter(has_text='body')
    expect(option).to_have_count(1)
    page.locator('#conditionBlock').select_option(option.get_attribute('value'))
    expect(page.locator('#conditionContext')).to_contain_text('body')
    page.locator('#conditionSave').click()
    expect(page.locator('#field-show_section')).to_be_visible()
    name = page.locator('#template').input_value()
    page.reload()
    expect(page.locator('#template')).to_have_value(name)
    page.locator('#field-title').fill('Условный документ')
    page.locator('#field-author').fill('Автор')
    page.locator('#field-body').fill('Текст по условию')
    for enabled in (False, True):
        page.locator('#field-show_section').set_checked(enabled)
        with page.expect_download(timeout=60000) as event:
            page.locator('#genBtn').click()
        doc = Document(io.BytesIO(event.value.path().read_bytes()))
        assert ('Текст по условию' in '\n'.join(p.text for p in doc.paragraphs)) is enabled
    page.locator('#editMode').click()
    page.get_by_text('Настроить условный абзац', exact=True).click()
    page.locator('#conditionInspect').click()
    option = page.locator('#conditionBlock option').filter(has_text='body')
    expect(option).to_contain_text('с условием')
    page.locator('#conditionBlock').select_option(option.get_attribute('value'))
    expect(page.locator('#conditionField')).to_have_value('show_section')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_template_loop_copy(e2e_server, page, tmp_path, monkeypatch, width):
    import importlib
    import io

    from docx import Document
    from playwright.sync_api import expect

    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path)
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    expect(page.locator('#field-title')).to_be_visible()
    page.locator('#editMode').click()
    page.get_by_text('Повторять абзац по списку', exact=True).click()
    page.locator('#loopInspect').click()
    option = page.locator('#loopBlock option').filter(has_text='body')
    expect(option).to_have_count(1)
    page.locator('#loopBlock').select_option(option.get_attribute('value'))
    page.locator('#loopVariable').select_option('body')
    page.locator('#loopSave').click()
    expect(page.locator('[data-field-name=items] .list-editor')).to_be_visible()
    name = page.locator('#template').input_value()
    page.reload()
    expect(page.locator('#template')).to_have_value(name)
    page.locator('#field-title').fill('Документ со списком')
    page.locator('#field-author').fill('Автор')
    page.locator('#field-body').fill('Прежнее поле')
    for data, expected in [('[]', []), ('["Первый", "Второй", "Первый"]', ['Первый', 'Второй', 'Первый'])]:
        page.locator('[data-field-name=items]').get_by_role('button', name='JSON', exact=True).click()
        page.locator('#field-items').fill(data)
        with page.expect_download(timeout=60000) as event:
            page.locator('#genBtn').click()
        doc = Document(io.BytesIO(event.value.path().read_bytes()))
        values = [p.text for p in doc.paragraphs]
        assert [value for value in values if value in {'Первый', 'Второй'}] == expected
        assert 'Прежнее поле' not in values
    page.locator('#editMode').click()
    page.get_by_text('Повторять абзац по списку', exact=True).click()
    page.locator('#loopInspect').click()
    option = page.locator('#loopBlock option').filter(has_text='повторяется')
    expect(option).to_have_count(1)
    page.locator('#loopBlock').select_option(option.get_attribute('value'))
    expect(page.locator('#loopField')).to_have_value('items')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('width', [375, 1280])
def test_e2e_template_row_copy(e2e_server, page, tmp_path, monkeypatch, width):
    import importlib
    import io
    import json

    from docx import Document
    from playwright.sync_api import expect

    from tests.test_web_template_rows import make_table

    templates = tmp_path / 'templates'
    templates.mkdir()
    make_table(templates / 'table.docx')
    (templates / 'table.schema.json').write_text(json.dumps({'fields': [
        {'name': 'title', 'type': 'string', 'required': True}], 'allow_extra': False}), encoding='utf-8')
    monkeypatch.setattr('textalchemy.generate.template.TEMPLATES_DIR', templates)
    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path / 'data')
    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    expect(page.locator('#field-title')).to_be_visible()
    page.locator('#editMode').click()
    page.get_by_text('Повторять строку таблицы', exact=True).click()
    with page.expect_response(lambda response: response.url.endswith('/table-loops')) as inspection:
        page.locator('#rowLoopInspect').click()
    assert inspection.value.status == 200, inspection.value.text()
    assert len(inspection.value.json()['blocks']) == 1
    expect(page.locator('#rowLoopBlock option')).to_have_count(1)
    expect(page.locator('#rowLoopContext')).to_contain_text('kg')
    page.locator('#rowLoopSave').click()
    expect(page.locator('[data-field-name=items] .list-editor')).to_be_visible()
    name = page.locator('#template').input_value()
    page.reload()
    expect(page.locator('#template')).to_have_value(name)
    page.locator('#field-title').fill('Global')
    for data, expected in [('[]', []), ('["A", "B", "A"]', [['A', 'kg'], ['B', 'kg'], ['A', 'kg']])]:
        page.locator('[data-field-name=items]').get_by_role('button', name='JSON', exact=True).click()
        page.locator('#field-items').fill(data)
        with page.expect_download(timeout=60000) as event:
            page.locator('#genBtn').click()
        doc = Document(io.BytesIO(event.value.path().read_bytes()))
        assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == [['Name', 'Unit'], *expected, ['Total', 'End']]
    page.locator('#editMode').click()
    page.get_by_text('Повторять строку таблицы', exact=True).click()
    page.locator('#rowLoopInspect').click()
    expect(page.locator('#rowLoopField')).to_have_value('items')
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.e2e_errors == []


@pytest.mark.parametrize('endpoint,backend', [('preview/meta', 'cached_page_count'), ('preview', 'cached_page_png')])
def test_template_preview_does_not_block_web(e2e_server, monkeypatch, endpoint, backend):
    import concurrent.futures

    import httpx

    entered, release = threading.Event(), threading.Event()

    def slow_preview(*args, **kwargs):
        entered.set()
        release.wait(timeout=10)
        return 0 if backend == 'cached_page_count' else None

    monkeypatch.setattr(f'textalchemy.web.routes.generate.{backend}', slow_preview)
    with concurrent.futures.ThreadPoolExecutor() as executor:
        pending = executor.submit(httpx.get, f'{e2e_server}/api/generate/templates/abstract/{endpoint}', timeout=15)
        try:
            assert entered.wait(timeout=5)
            response = httpx.get(f'{e2e_server}/api/generate/templates', timeout=3)
            assert response.status_code == 200
            assert not release.is_set()
        finally:
            release.set()
            pending.result(timeout=15)
