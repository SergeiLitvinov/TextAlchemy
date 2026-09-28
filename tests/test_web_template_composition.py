"""Different editor actions compose in one persisted template and real export."""

import base64
import io
import json
from zipfile import ZipFile

from docx import Document
from PIL import Image

from tests import test_web_m4_completion as acceptance

m4_client, import_sample = acceptance.m4_client, acceptance.import_sample


def action(client, name, endpoint, prefix, **params):
    inspected = client.get(f"/api/generate/templates/{name}/{endpoint}")
    assert inspected.status_code == 200, inspected.text
    data = inspected.json()
    block = next(item for item in data["blocks"] if item["text"].startswith(prefix))
    response = client.post(
        f"/api/generate/templates/{name}/{endpoint}", json={"block": block["id"], "revision": data["revision"], **params}
    )
    assert response.status_code == 200, response.text
    return response.json()["name"]


def rich(client, name, prefix, kind, field):
    data = client.get(f"/api/generate/templates/{name}/source").json()
    block = next(item for item in data["blocks"] if item["text"].startswith(prefix))
    response = client.post(
        f"/api/generate/templates/{name}/rich",
        json={"block": block["id"], "kind": kind, "field": field, "expected": data["revision"]},
    )
    assert response.status_code == 200, response.text
    return response.json()["name"]


def test_composed_template_conditions_loops_rich_content_and_rename(m4_client):
    client, _ = m4_client
    doc = Document()
    for text in [
        "Title",
        "Optional {{body}}",
        "Repeat {{body}}",
        "Image slot",
        "Formula slot",
        "Bibliography slot",
        "Target paragraph",
        "Reference slot",
    ]:
        doc.add_paragraph(text)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Cell {{body}}"
    out = io.BytesIO()
    doc.save(out)
    name = import_sample(client, out.getvalue())
    name = action(client, name, "conditions", "Optional", field="show_section", negate=False)
    name = action(client, name, "loops", "Repeat", variable="body", field="items")
    name = action(client, name, "table-loops", "Cell", variable="body", field="rows")
    for prefix, kind, field in [
        ("Image", "image", "picture"),
        ("Formula", "formula", "equation"),
        ("Bibliography", "bibliography", "bibliography"),
        ("Target", "id", "sec:target"),
        ("Reference", "ref", "sec:target"),
    ]:
        name = rich(client, name, prefix, kind, field)
    data = client.get(f"/api/generate/templates/{name}/variables")
    assert data.status_code == 200, data.text
    renamed = client.post(
        f"/api/generate/templates/{name}/variables", json={"old": "body", "new": "subject", "revision": data.json()["revision"]}
    )
    assert renamed.status_code == 200, renamed.text
    name = renamed.json()["name"]
    picture = io.BytesIO()
    Image.new("RGB", (20, 10), "blue").save(picture, format="PNG")
    params = {
        "subject": "Shared",
        "show_section": True,
        "items": ["One", "Two"],
        "rows": ["Row1", "Row2"],
        "equation": '<math xmlns="http://www.w3.org/1998/Math/MathML"><msup><mi>x</mi><mn>2</mn></msup></math>',
        "picture": "data:image/png;base64," + base64.b64encode(picture.getvalue()).decode(),
        "bibliography": ["Первый источник", "Второй источник"],
    }
    for show in (True, False):
        params["show_section"] = show
        response = client.post("/api/generate", data={"template": name, "params": json.dumps(params)})
        assert "application/vnd.openxmlformats" in response.headers["content-type"], response.text
        result = Document(io.BytesIO(response.content))
        text = "\n".join(p.text for p in result.paragraphs)
        assert ("Optional Shared" in text) is show
        assert "Repeat One" in text and "Repeat Two" in text and "Первый источник" in text
        assert "Reference slot" not in text and "{{" not in text
        assert [row.cells[0].text for row in result.tables[0].rows] == ["Cell Row1", "Cell Row2"]
        with ZipFile(io.BytesIO(response.content)) as package:
            assert any(name.startswith("word/media/") for name in package.namelist())
            assert b"oMath" in package.read("word/document.xml")


def test_multiline_scalar_uses_word_line_break(m4_client):
    client, _ = m4_client
    doc = Document()
    doc.add_paragraph("{{body}}")
    out = io.BytesIO()
    doc.save(out)
    name = import_sample(client, out.getvalue())
    response = client.post("/api/generate", data={"template": name, "params": json.dumps({"body": "Line1\nLine2"})})
    assert Document(io.BytesIO(response.content)).paragraphs[0].text == "Line1\nLine2"
    with ZipFile(io.BytesIO(response.content)) as archive:
        assert b"<w:br" in archive.read("word/document.xml")
