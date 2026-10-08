"""Own native PPTX fixture, independent of the adapter under acceptance."""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.util import Inches, Pt


def write_chart_presentation(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(5), Inches(.5)).text = "Own chart acceptance"
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Own series", [3, 7])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data)
    presentation.save(path)


def write_visual_presentation(path, case):
    """Isolate native visual properties using only own text and shapes."""
    if case not in {"master-background", "nested-group", "table-style", "freeform-line",
                    "freeform-axis-lines", "table-theme"}:
        raise ValueError(f"Unknown own PPTX case: {case}")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(5), Inches(.5)).text = "Own visual acceptance"
    if case == "master-background":
        fill = presentation.slide_master.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(24, 72, 96)
    elif case == "nested-group":
        outer = slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        rectangle = inner.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                           Inches(1), Inches(1), Inches(2), Inches(1))
        rectangle.fill.solid()
        rectangle.fill.fore_color.rgb = RGBColor(24, 72, 96)
        rectangle.text = "Own group"
        inner.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(4), Inches(1), Inches(1), Inches(1))
        outer.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3), Inches(1.5),
                                   Inches(4), Inches(1.5))
    elif case == "freeform-line":
        builder = slide.shapes.build_freeform(0, 0, scale=Inches(1) / 1000)
        builder.add_line_segments([(2000, 0), (2000, 1500), (4000, 1500)], close=False)
        contour = builder.convert_to_shape(Inches(1), Inches(1))
        contour.fill.background()
        contour.line.color.rgb = RGBColor(24, 72, 96)
        contour.line.width = Pt(6)
    elif case == "freeform-axis-lines":
        for x, y, origin_y in ((4000, 0, 1), (0, 2000, 2)):
            builder = slide.shapes.build_freeform(0, 0, scale=Inches(1) / 1000)
            builder.add_line_segments([(x, y)], close=False)
            line = builder.convert_to_shape(Inches(1), Inches(origin_y))
            line.fill.background()
            line.line.color.rgb = RGBColor(24, 72, 96)
            line.line.width = Pt(6)
    else:
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2)).table
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.text = f"Own cell {row_index + 1}:{column_index + 1}"
                if case == "table-theme":
                    continue
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(24, 72, 96)
                cell.margin_left = cell.margin_right = Pt(18)
                run = cell.text_frame.paragraphs[0].runs[0]
                run.font.name = "Arial"
                run.font.size = Pt(20)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    presentation.save(path)
