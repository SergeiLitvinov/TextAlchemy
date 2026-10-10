"""Finite opaque DOCX comments and readOnly enforcement via the application."""

import json
import os
import shutil
import subprocess
import sys
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from textalchemy.convert.executor import ConversionExecutor, ConversionRequest
from textalchemy.core.quality_policy import QualityPolicy
from textalchemy.core.types import DocFormat


def own_document(path, kind):
    document = Document()
    paragraph = document.add_paragraph("Own commented paragraph." if kind == "comments" else "Own protected paragraph.")
    if kind == "comments":
        document.add_comment(paragraph.runs[0], text="Own comment content", author="QA Author", initials="QA")
        document.save(path)
    else:
        buffer = BytesIO()
        document.save(buffer)
        with ZipFile(buffer) as source, ZipFile(path, "w", ZIP_DEFLATED) as output:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/settings.xml":
                    assert b"</w:settings>" in data
                    data = data.replace(
                        b"</w:settings>", b'<w:documentProtection w:edit="readOnly" w:enforcement="1"/></w:settings>'
                    )
                output.writestr(info, data)
    return path


def cycle(source, model, output):
    executor = ConversionExecutor()
    imported = executor.execute(ConversionRequest(source, model, DocFormat.DOCX, DocFormat.MODEL))
    exported = executor.execute(ConversionRequest(model, output, DocFormat.MODEL, DocFormat.DOCX))
    assert imported.success and exported.success
    return imported, exported


@pytest.mark.parametrize("kind", ["comments", "protection"])
def test_opaque_docx_two_cycles_and_strict_atomic_refusal(tmp_path, kind):
    source = own_document(tmp_path / "source.docx", kind)
    original = source.read_bytes()
    current = source
    feature = "docx." + kind
    for index in range(2):
        model = tmp_path / f"cycle-{index}.json"
        output = tmp_path / f"cycle-{index}.docx"
        imported, exported = cycle(current, model, output)
        assert not imported.lossless and not exported.lossless
        assert any(i.feature == feature and i.location for i in imported.issues)
        assert any(i.feature == feature and i.location for i in exported.issues)
        assert [p.text for p in Document(output).paragraphs] == [p.text for p in Document(source).paragraphs]
        if kind == "comments":
            comments = list(Document(output).comments)
            assert len(comments) == 1
            assert (comments[0].text, comments[0].author, comments[0].initials) == ("Own comment content", "QA Author", "QA")
        else:
            from xml.etree import ElementTree as ET

            with ZipFile(output) as package:
                protection = ET.fromstring(package.read("word/settings.xml")).find(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}documentProtection"
                )
                assert protection.attrib == {
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}edit": "readOnly",
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}enforcement": "1",
                }
        current = output
    previous = tmp_path / "previous.docx"
    previous.write_bytes(b"previous artifact")
    rejected = ConversionExecutor().execute(
        ConversionRequest(model, previous, DocFormat.MODEL, DocFormat.DOCX, quality_policy=QualityPolicy(0))
    )
    assert not rejected.success
    assert any(i.feature == feature and i.location for i in rejected.issues)
    assert previous.read_bytes() == b"previous artifact"
    assert source.read_bytes() == original


@pytest.mark.skipif(
    sys.platform != "win32" or os.getenv("TEXTALCHEMY_WORD_ACCEPTANCE") != "1",
    reason="Requires installed Word and explicit local acceptance run",
)
@pytest.mark.parametrize("kind", ["comments", "protection"])
def test_native_word_properties_after_two_cycles(tmp_path, kind):
    source = own_document(tmp_path / "source.docx", kind)
    original = source.read_bytes()
    current = source
    expected_comment = "Own comment content"
    observations = []
    script = Path(__file__).resolve().parents[1] / "tools/acceptance/word-complex.ps1"
    for index in range(2):
        model = tmp_path / f"native-{index}.json"
        output = tmp_path / f"native-{index}.docx"
        cycle(current, model, output)
        arguments = [
            "powershell",
            "-NoProfile",
            "-NonInteractive",
            "-File",
            str(script),
            "-InputPath",
            str(output),
            "-Kind",
            kind,
            "-ExpectedComment",
            expected_comment,
        ]
        if kind == "comments" and index == 0:
            edited = tmp_path / "native-edited.docx"
            arguments += ["-EditedPath", str(edited)]
        result = subprocess.run(arguments, capture_output=True, text=True, encoding="utf-8", timeout=90, check=True)
        observed = json.loads(result.stdout)
        assert all(value is True for value in observed["checks"].values())
        if kind == "comments" and index == 0:
            assert observed["native_edit_reopened"] is True
            current = edited
            expected_comment = "Own native comment edit"
        else:
            assert observed["native_edit_reopened"] is None
            current = output
        observations.append(observed)
    assert source.read_bytes() == original
    (tmp_path / "word-evidence.json").write_text(
        json.dumps({"source_sha256": sha256(original).hexdigest(), "scope": kind, "cycles": observations}, indent=2),
        encoding="utf-8",
    )


@pytest.mark.skipif(
    sys.platform != "win32" or os.getenv("TEXTALCHEMY_WORD_ACCEPTANCE") != "1",
    reason="Requires installed Word/Excel and explicit local acceptance run",
)
@pytest.mark.parametrize("kind,feature", [("smartart", "docx.smartart"), ("ole", "docx.embedded-ole")])
def test_native_complex_sources_refused_by_zero_loss_budget(tmp_path, kind, feature):
    """Prove the strict application gate; permissive native export is not accepted."""
    script = Path(__file__).resolve().parents[1] / "tools/acceptance/docx-objects.ps1"
    completed = subprocess.run(
        [
            shutil.which("pwsh") or "powershell",
            "-NoProfile",
            "-NonInteractive",
            "-File",
            str(script),
            "-OutputDirectory",
            str(tmp_path),
            "-Kind",
            kind,
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=90,
        check=True,
    )
    generated = json.loads(completed.stdout)
    assert generated["source_reopened"] is True
    source = tmp_path / f"own-{kind}.docx"
    original = source.read_bytes()
    assert sha256(original).hexdigest() == generated["source_sha256"]
    model = tmp_path / "model.json"
    executor = ConversionExecutor()
    imported = executor.execute(ConversionRequest(source, model, DocFormat.DOCX, DocFormat.MODEL))
    assert imported.success and not imported.lossless
    assert any(issue.feature == feature and issue.location for issue in imported.issues)
    output = tmp_path / "previous.docx"
    output.write_bytes(b"previous artifact")
    rejected = executor.execute(
        ConversionRequest(
            model,
            output,
            DocFormat.MODEL,
            DocFormat.DOCX,
            quality_policy=QualityPolicy(0),
        )
    )
    assert not rejected.success
    assert any(issue.feature == feature and issue.location for issue in rejected.issues)
    assert output.read_bytes() == b"previous artifact"
    assert source.read_bytes() == original
