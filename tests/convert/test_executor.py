"""Tests for execution of planned conversion routes."""

from pathlib import Path

from docx import Document

from textalchemy.convert.base import ConversionResult
from textalchemy.convert.executor import ConversionExecutor, ConversionRequest, _convert_pdf, infer_format
from textalchemy.core.conversion_graph import (
    CapabilityRegistry,
    ConverterCapabilities,
    DocumentFeature,
    FeatureSupport,
)
from textalchemy.core.diagnostics import ConversionReport, IssueSeverity
from textalchemy.core.document_codec import load_document
from textalchemy.core.document_model import ConversionMode
from textalchemy.core.types import DocFormat


def test_executor_converts_docx_to_model_without_intermediate_file(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "model.json"
    document = Document()
    document.add_paragraph("Universal route")
    document.save(source)

    report = ConversionExecutor().execute(ConversionRequest(source, output, DocFormat.DOCX, DocFormat.MODEL))

    assert report.success
    assert report.metrics["executed_steps"] == ["docx.model"]
    assert load_document(output).sections[0].blocks[0].plain_text == "Universal route"


def test_executor_runs_docx_model_html_route(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "result.html"
    document = Document()
    document.add_heading("Result", level=1)
    document.add_paragraph("Editable body")
    document.save(source)

    report = ConversionExecutor().execute(
        ConversionRequest(
            source,
            output,
            DocFormat.DOCX,
            DocFormat.HTML,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert report.success
    assert report.metrics["executed_steps"] == ["docx.model", "model.html"]
    assert "Editable body" in output.read_text(encoding="utf-8")


def test_executor_skips_route_with_missing_requirement(tmp_path):
    registry = CapabilityRegistry()
    modes = frozenset({ConversionMode.BALANCED})
    registry.register(
        ConverterCapabilities(
            "missing",
            DocFormat.TXT,
            DocFormat.HTML,
            modes,
            {DocumentFeature.TEXT: FeatureSupport.EXACT},
            requirements=("missing-engine",),
        )
    )
    registry.register(
        ConverterCapabilities(
            "fallback",
            DocFormat.TXT,
            DocFormat.HTML,
            modes,
            {DocumentFeature.TEXT: FeatureSupport.EDITABLE},
            base_cost=2,
        )
    )
    source = tmp_path / "source.txt"
    output = tmp_path / "output.html"
    source.write_text("body", encoding="utf-8")

    def fallback(value, target):
        target.write_text(Path(value).read_text(encoding="utf-8"), encoding="utf-8")
        return target, None

    executor = ConversionExecutor(
        registry=registry,
        handlers={"missing": fallback, "fallback": fallback},
        requirement_checker=lambda requirement: requirement != "missing-engine",
    )
    report = executor.execute(
        ConversionRequest(
            source,
            output,
            DocFormat.TXT,
            DocFormat.HTML,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert report.success
    assert report.metrics["executed_steps"] == ["fallback"]
    assert output.read_text(encoding="utf-8") == "body"


def test_executor_reports_unavailable_theoretical_route(tmp_path):
    registry = CapabilityRegistry()
    registry.register(
        ConverterCapabilities(
            "missing",
            DocFormat.TXT,
            DocFormat.HTML,
            frozenset({ConversionMode.BALANCED}),
            {DocumentFeature.TEXT: FeatureSupport.EXACT},
            requirements=("missing-engine",),
        )
    )
    source = tmp_path / "source.txt"
    source.write_text("body", encoding="utf-8")
    executor = ConversionExecutor(
        registry=registry,
        handlers={"missing": lambda value, output: (output, None)},
        requirement_checker=lambda _requirement: False,
    )

    report = executor.execute(
        ConversionRequest(
            source,
            tmp_path / "output.html",
            DocFormat.TXT,
            DocFormat.HTML,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert not report.success
    assert report.issues[0].feature == "route"
    assert "missing-engine" in report.issues[0].message


def test_executor_stops_after_failed_backend(tmp_path):
    registry = CapabilityRegistry()
    modes = frozenset({ConversionMode.BALANCED})
    registry.register(
        ConverterCapabilities(
            "first",
            DocFormat.TXT,
            DocFormat.DOCX,
            modes,
            {DocumentFeature.TEXT: FeatureSupport.EXACT},
        )
    )
    registry.register(
        ConverterCapabilities(
            "second",
            DocFormat.DOCX,
            DocFormat.HTML,
            modes,
            {DocumentFeature.TEXT: FeatureSupport.EXACT},
        )
    )
    source = tmp_path / "source.txt"
    output = tmp_path / "output.html"
    source.write_text("body", encoding="utf-8")
    calls: list[str] = []

    def first(value, target):
        calls.append("first")
        report = ConversionReport(target)
        report.add(IssueSeverity.ERROR, "converter", "failed")
        return Path(value), report

    def second(value, target):
        calls.append("second")
        return target, None

    report = ConversionExecutor(
        registry=registry,
        handlers={"first": first, "second": second},
    ).execute(
        ConversionRequest(
            source,
            output,
            DocFormat.TXT,
            DocFormat.HTML,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert not report.success
    assert calls == ["first"]
    assert report.metrics["executed_steps"] == ["first"]


def test_empty_backend_mapping_disables_all_routes(tmp_path):
    source = tmp_path / "source.docx"
    source.write_bytes(b"not used")

    report = ConversionExecutor(handlers={}).execute(
        ConversionRequest(source, tmp_path / "model.json", DocFormat.DOCX, DocFormat.MODEL)
    )

    assert not report.success
    assert report.issues[0].feature == "route"


def test_infer_format_supports_model_and_latex():
    assert infer_format("document.json") is DocFormat.MODEL
    assert infer_format("document.tex") is DocFormat.LATEX


def test_executor_exposes_only_runtime_available_plan():
    executor = ConversionExecutor(requirement_checker=lambda _requirement: True)

    plan = executor.plan(
        DocFormat.PDF,
        DocFormat.DOCX,
        mode=ConversionMode.FAITHFUL,
        features=frozenset({DocumentFeature.PAGE_GEOMETRY}),
    )

    assert plan is not None
    assert plan.steps[0].id == "pdf.docx.pymupdf"


def test_pdf_backend_falls_back_after_runtime_failure(tmp_path, monkeypatch):
    source = tmp_path / "source.pdf"
    output = tmp_path / "output.docx"
    source.write_bytes(b"pdf")

    class FakeConverter:
        def __init__(self, name):
            self.name = name

        def convert(self, input_path, output_path):
            if self.name == "pdf2docx":
                return ConversionResult(input_path, output_path, False, "broken parser")
            Path(output_path).write_bytes(b"docx")
            return ConversionResult(input_path, output_path, True)

    monkeypatch.setattr(
        "textalchemy.convert.pdf_to_docx.create_converter",
        lambda name: FakeConverter(name),
    )

    report = _convert_pdf(source, output, "pdf2docx")

    assert report.success
    assert output.read_bytes() == b"docx"
    assert report.metrics == {"engine": "pymupdf", "requested_engine": "pdf2docx"}
    assert report.issues[0].feature == "engine-fallback"


def _rich_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.6))
    box.text_frame.text = "Model-based PPTX body"
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2), Inches(2), Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    presentation.save(path)
    return path


def test_executor_converts_pptx_to_model(tmp_path):
    source = _rich_pptx(tmp_path / "source.pptx")
    output = tmp_path / "model.json"

    report = ConversionExecutor().execute(ConversionRequest(source, output, DocFormat.PPTX, DocFormat.MODEL))

    assert report.success
    assert report.metrics["executed_steps"] == ["pptx.model"]
    model = load_document(output)
    assert model.source_format == "pptx"


def test_executor_runs_pptx_model_html_route(tmp_path):
    source = _rich_pptx(tmp_path / "source.pptx")
    output = tmp_path / "result.html"

    report = ConversionExecutor().execute(
        ConversionRequest(
            source,
            output,
            DocFormat.PPTX,
            DocFormat.HTML,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert report.success
    assert report.metrics["executed_steps"] == ["pptx.model", "model.html"]
    html = output.read_text(encoding="utf-8")
    assert "Model-based PPTX body" in html
    assert "background-image:url(&quot;data:image/svg+xml" in html
    assert "fill%3D%22%234472C4%22" in html


def test_executor_runs_pptx_model_docx_route(tmp_path):
    source = _rich_pptx(tmp_path / "source.pptx")
    output = tmp_path / "result.docx"

    report = ConversionExecutor().execute(
        ConversionRequest(
            source,
            output,
            DocFormat.PPTX,
            DocFormat.DOCX,
            features=frozenset({DocumentFeature.TEXT}),
        )
    )

    assert report.success
    assert report.metrics["executed_steps"] == ["pptx.model", "model.docx"]
    assert output.is_file()
    document = Document(str(output))
    assert "Model-based PPTX body" in "\n".join(p.text for p in document.paragraphs)
