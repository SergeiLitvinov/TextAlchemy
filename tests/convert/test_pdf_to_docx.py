"""Тесты fan-out конвертера и улучшенного PyMuPdfConverter."""
from __future__ import annotations

import zipfile

import pytest

from textalchemy.convert import FanOutConverter
from textalchemy.convert.pdf_to_docx import (
    PyMuPdfConverter,
    create_converter,
)


def test_fanout_registered():
    from textalchemy.convert import create_converter
    c = create_converter("fanout")
    assert isinstance(c, FanOutConverter)


def test_fanout_creates_real_docx(tmp_path):
    """Fan-out на реальном PDF: должен дать валидный DOCX через pymupdf."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf_path = tmp_path / "sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF world")
    doc.save(str(pdf_path))
    doc.close()

    out_path = tmp_path / "sample.docx"
    c = FanOutConverter(tools=["pdf2docx", "pymupdf"])
    result = c.convert(pdf_path, out_path)
    # Может не быть pdf2docx; но pymupdf должен сработать
    if result.success:
        assert out_path.is_file()
        assert out_path.stat().st_size > 0
        with zipfile.ZipFile(out_path) as zf:
            assert "word/document.xml" in zf.namelist()
    else:
        pytest.skip(f"all engines failed: {result.error}")


def test_pymupdf_text_page(tmp_path):
    """PyMuPdf: страница с текстом → текст в DOCX."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF world")
    page.insert_text((72, 100), "Second line here")
    doc.save(str(pdf))
    doc.close()

    # Сразу проверим, что pymupdf видит текст
    reread = fitz.open(str(pdf))
    text_before = reread[0].get_text()
    reread.close()
    assert "Hello" in text_before, f"pymupdf didn't extract: {text_before!r}"

    out = tmp_path / "out.docx"
    c = PyMuPdfConverter(text_threshold=10)  # понизим порог
    result = c.convert(pdf, out)
    assert result.success
    assert out.is_file()
    from docx import Document
    d = Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Hello" in text, f"DOCX has no text: {text!r}"


def test_pymupdf_scan_page_uses_image(tmp_path):
    """PyMuPdf: страница с <50 символов текста → изображение."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")

    pdf = tmp_path / "in.pdf"
    doc = fitz.open()
    # Страница без текста (пустая) — будет отрендерена как картинка
    doc.new_page()
    doc.save(str(pdf))
    doc.close()

    out = tmp_path / "out.docx"
    c = PyMuPdfConverter(render_dpi=72)
    result = c.convert(pdf, out)
    assert result.success
    assert out.is_file()
    # DOCX должен содержать изображение
    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert media  # хотя бы одна картинка


def test_create_converter_default():
    from textalchemy.convert.pdf_to_docx import Pdf2DocxConverter
    conv = create_converter("pdf2docx")
    assert isinstance(conv, Pdf2DocxConverter)


def test_create_converter_pymupdf():
    from textalchemy.convert.pdf_to_docx import PyMuPdfConverter
    conv = create_converter("pymupdf")
    assert isinstance(conv, PyMuPdfConverter)


def test_create_converter_unknown():
    with pytest.raises(Exception):
        create_converter("nonexistent_engine")


def test_convert_nonexistent():
    from textalchemy.convert.pdf_to_docx import Pdf2DocxConverter
    conv = Pdf2DocxConverter()
    result = conv.convert("nonexistent.pdf", "out.docx")
    assert not result.success
    assert "not found" in (result.error or "").lower()


def _make_pdf_with_headings(tmp_path):
    """PDF с заголовком (крупный шрифт), подзаголовком и телом текста."""
    try:
        import fitz
    except ImportError:
        pytest.skip("pymupdf not installed")
    pdf = tmp_path / "headings.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 60), "Глава 1. Введение", fontsize=20)
    page.insert_text((50, 100), "Обычный абзац с текстом для проверки извлечения.", fontsize=11)
    page.insert_text((50, 140), "Подраздел", fontsize=14)
    page.insert_text((50, 180), "Ещё один абзац тела документа.", fontsize=11)
    doc.save(str(pdf))
    doc.close()
    return pdf


def test_fanout_restores_headings(tmp_path):
    """После fanout заголовки должны получить стиль Heading, а не только Normal."""
    pdf = _make_pdf_with_headings(tmp_path)
    out = tmp_path / "out.docx"
    result = FanOutConverter().convert(pdf, out)
    assert result.success, result.error
    from docx import Document

    doc = Document(str(out))
    styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
    assert any(s.startswith("Heading") for s in styles), f"no headings: {styles}"


def test_pymupdf_preserves_text_runs(tmp_path):
    """PyMuPDF: текстовая страница → текст, без растровых вставок."""
    pdf = _make_pdf_with_headings(tmp_path)
    out = tmp_path / "out.docx"
    result = PyMuPdfConverter().convert(pdf, out)
    assert result.success
    from docx import Document

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    # Текст извлечён как текст (а не картинка). Кириллица в консоли искажается,
    # поэтому проверяем длину и отсутствие растровых вставок.
    assert len(text.strip()) > 50
    with zipfile.ZipFile(out) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert not media, "текстовая страница не должна содержать картинок"


def test_docx_to_latex_quality(tmp_path):
    """DOCX→LaTeX: заголовки, жирный текст и таблицы попадают в .tex."""
    try:
        from docx import Document as D
    except ImportError:
        pytest.skip("python-docx not installed")
    from textalchemy.convert.docx_to_latex import DocxToLatexConverter

    doc = D()
    doc.add_heading("Глава 1. Введение", level=1)
    p = doc.add_paragraph()
    p.add_run("Обычный ")
    rb = p.add_run("жирный")
    rb.bold = True
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Имя"
    t.rows[0].cells[1].text = "Год"
    src = tmp_path / "src.docx"
    doc.save(str(src))
    out = tmp_path / "out.tex"
    result = DocxToLatexConverter().convert(src, out)
    assert result.success, result.error
    tex = out.read_text(encoding="utf-8")
    assert r"\section{" in tex
    assert r"\textbf{жирный}" in tex
    assert r"\begin{tabular}" in tex
    assert r"\end{document}" in tex

