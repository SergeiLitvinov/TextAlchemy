"""Тесты DocumentModel → DOCX и диагностического отчёта."""

import zipfile

from docx import Document
from PIL import Image as PillowImage

from textalchemy.convert.docx_writer import write_docx_model
from textalchemy.core.diagnostics import IssueSeverity
from textalchemy.core.document_model import (
    Box,
    DocumentModel,
    Formula,
    FormulaFormat,
    Image,
    ImageCrop,
    Length,
    Paragraph,
    Resource,
    ResourceKind,
    Section,
    Table,
    TableCell,
    TableRow,
    TextRun,
    TextStyle,
)
from textalchemy.formats.docx import read_docx_model


def _png_bytes(tmp_path):
    path = tmp_path / "source.png"
    PillowImage.new("RGB", (12, 8), "blue").save(path)
    return path.read_bytes()


def test_write_docx_model_roundtrip_preserves_core_content(tmp_path):
    output = tmp_path / "roundtrip.docx"
    document = DocumentModel(
        metadata={"title": "Round trip", "author": "TextAlchemy"},
        resources={"picture": Resource("picture", ResourceKind.RASTER_IMAGE, "image/png", data=_png_bytes(tmp_path))},
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        content=[
                            TextRun(
                                "Linked text",
                                style=TextStyle(font_family="Arial", font_size=Length(13), bold=True, color="#336699"),
                                link="https://example.com",
                            ),
                            Image("picture", alt_text="blue image", box=Box(0, 0, 72, 48)),
                        ],
                        alignment="center",
                    ),
                    Table(
                        rows=[
                            TableRow(
                                cells=[
                                    TableCell(blocks=[Paragraph(content=[TextRun("wide")])], column_span=2),
                                    TableCell(blocks=[Paragraph(content=[TextRun("single")])]),
                                ]
                            )
                        ]
                    ),
                ],
                headers=[Paragraph(content=[TextRun("Header")])],
                footers=[Paragraph(content=[TextRun("Footer")])],
            )
        ],
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)

    assert report.success is True
    assert report.lossless is True
    assert output.is_file()
    assert restored.metadata["title"] == "Round trip"
    assert restored.metadata["author"] == "TextAlchemy"
    assert restored.sections[0].headers[0].plain_text == "Header"
    assert restored.sections[0].footers[0].plain_text == "Footer"
    assert len([resource for resource in restored.resources.values() if resource.kind is ResourceKind.RASTER_IMAGE]) == 1
    assert report.metrics["images"] == 1
    assert report.metrics["tables"] == 1
    with zipfile.ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "Linked text" in document_xml
    assert "https://example.com" in relationships


def test_write_docx_model_preserves_native_omml(tmp_path):
    output = tmp_path / "formula.docx"
    omml = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'
    document = DocumentModel(sections=[Section(blocks=[Formula(omml, FormulaFormat.OMML)])])

    report = write_docx_model(document, output)

    assert report.lossless is True
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<m:oMath" in xml
    assert "x+1" in xml


def test_write_docx_model_reports_formula_fallback(tmp_path):
    output = tmp_path / "latex.docx"
    document = DocumentModel(sections=[Section(blocks=[Formula("E=mc^2", FormulaFormat.LATEX, fallback_text="E = mc²")])])

    report = write_docx_model(document, output)

    assert report.success is True
    assert report.lossless is False
    assert any(issue.severity is IssueSeverity.LOSS and issue.feature == "formula" for issue in report.issues)
    restored = read_docx_model(output)
    assert restored.sections[0].blocks[0].plain_text == "E = mc²"


def test_write_docx_model_reports_missing_resource_as_error(tmp_path):
    output = tmp_path / "missing.docx"
    document = DocumentModel(sections=[Section(blocks=[Image("missing", alt_text="missing image")])])

    report = write_docx_model(document, output)

    assert report.success is False
    assert any(issue.severity is IssueSeverity.ERROR and issue.feature == "image" for issue in report.issues)


def test_write_docx_model_preserves_svg_as_vector_resource(tmp_path):
    output = tmp_path / "vector.docx"
    svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="100" height="50"><rect width="100" height="50"/></svg>'
    document = DocumentModel(
        resources={"vector": Resource("vector", ResourceKind.VECTOR_IMAGE, "image/svg+xml", data=svg)},
        sections=[Section(blocks=[Paragraph(content=[Image("vector", alt_text="vector diagram", box=Box(0, 0, 100, 50))])])],
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)

    assert report.lossless
    vector_resources = [resource for resource in restored.resources.values() if resource.kind is ResourceKind.VECTOR_IMAGE]
    assert len(vector_resources) == 1
    resource = vector_resources[0]
    assert resource.kind is ResourceKind.VECTOR_IMAGE
    assert resource.media_type == "image/svg+xml"
    assert resource.data == svg
    restored_image = restored.sections[0].blocks[0].content[0]
    assert isinstance(restored_image, Image)
    assert restored_image.alt_text == "vector diagram"


def test_write_docx_model_recreates_custom_styles_and_continuous_sections(tmp_path):
    source_path = tmp_path / "source.docx"
    output = tmp_path / "roundtrip.docx"
    from docx.enum.section import WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    source = Document()
    base = source.styles.add_style("Scientific Base", WD_STYLE_TYPE.PARAGRAPH)
    base.font.name = "Arial"
    base.font.size = Pt(12)
    child = source.styles.add_style("Scientific Result", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    child.font.bold = True
    source.add_paragraph("Styled content", style=child)
    source.sections[0].header.paragraphs[0].text = "Shared header"
    second = source.add_section(WD_SECTION.CONTINUOUS)
    second.header.is_linked_to_previous = True
    source.add_paragraph("Continuous content")
    source.save(source_path)
    model = read_docx_model(source_path)

    report = write_docx_model(model, output)
    restored = read_docx_model(output)

    assert report.lossless
    assert not [issue for issue in report.issues if issue.feature == "paragraph-style"]
    assert any(style.properties["style_name"] == "Scientific Result" for style in restored.styles.values())
    assert restored.sections[0].blocks[0].plain_text == "Styled content"
    assert restored.sections[0].blocks[0].content[0].style.font_family == "Arial"
    assert restored.sections[0].blocks[0].content[0].style.bold is True
    assert restored.sections[1].properties["start_type"] == "CONTINUOUS"
    assert restored.sections[1].properties["header_linked_to_previous"] is True


def test_write_docx_model_roundtrips_floating_image_anchor(tmp_path):
    output = tmp_path / "floating.docx"
    document = DocumentModel(
        resources={"picture": Resource("picture", ResourceKind.RASTER_IMAGE, "image/png", data=_png_bytes(tmp_path))},
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        content=[
                            Image(
                                "picture",
                                alt_text="floating chart",
                                box=Box(36, 54, 144, 96),
                                properties={
                                    "placement": "anchor",
                                    "horizontal_relative_from": "page",
                                    "vertical_relative_from": "page",
                                    "wrap": "square",
                                    "wrap_text": "bothSides",
                                    "relative_height": 5,
                                    "allow_overlap": False,
                                },
                            )
                        ]
                    )
                ]
            )
        ],
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)

    assert report.lossless
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "<wp:anchor" in xml
    assert '<wp:positionH relativeFrom="page"><wp:posOffset>457200</wp:posOffset>' in xml
    restored_image = restored.sections[0].blocks[0].content[0]
    assert isinstance(restored_image, Image)
    assert restored_image.properties["placement"] == "anchor"
    assert restored_image.properties["wrap"] == "square"
    assert restored_image.properties["horizontal_relative_from"] == "page"
    assert restored_image.properties["vertical_relative_from"] == "page"
    assert restored_image.box == Box(36, 54, 144, 96)


def test_write_docx_model_roundtrips_image_crop_rotation_and_wrap_polygon(tmp_path):
    output = tmp_path / "transformed-image.docx"
    polygon = {
        "edited": True,
        "points": [
            {"x": 0, "y": 0},
            {"x": 18000, "y": 1200},
            {"x": 21600, "y": 18000},
            {"x": 2400, "y": 21600},
        ],
    }
    document = DocumentModel(
        resources={"picture": Resource("picture", ResourceKind.RASTER_IMAGE, "image/png", data=_png_bytes(tmp_path))},
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        content=[
                            Image(
                                "picture",
                                box=Box(36, 54, 144, 96, rotation=12.5),
                                crop=ImageCrop(left=0.1, top=0.025, right=0.05, bottom=0.075),
                                properties={
                                    "placement": "anchor",
                                    "wrap": "tight",
                                    "wrap_text": "largest",
                                    "wrap_polygon": polygon,
                                },
                            )
                        ]
                    )
                ]
            )
        ],
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)

    assert report.lossless
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert '<a:srcRect l="10000" t="2500" r="5000" b="7500"/>' in xml
    assert 'rot="750000"' in xml
    assert '<wp:wrapTight wrapText="largest"><wp:wrapPolygon edited="1">' in xml
    restored_image = restored.sections[0].blocks[0].content[0]
    assert isinstance(restored_image, Image)
    assert restored_image.crop == ImageCrop(left=0.1, top=0.025, right=0.05, bottom=0.075)
    assert restored_image.box == Box(36, 54, 144, 96, rotation=12.5)
    assert restored_image.properties["wrap_text"] == "largest"
    assert restored_image.properties["wrap_polygon"] == polygon


def test_write_docx_model_preserves_first_and_even_running_content(tmp_path):
    output = tmp_path / "alternate-running-content.docx"
    document = DocumentModel(
        sections=[
            Section(
                blocks=[Paragraph(content=[TextRun("Body")])],
                headers=[Paragraph(content=[TextRun("Odd header")])],
                footers=[Paragraph(content=[TextRun("Odd footer")])],
                first_page_headers=[Paragraph(content=[TextRun("First header")])],
                first_page_footers=[Paragraph(content=[TextRun("First footer")])],
                even_page_headers=[Paragraph(content=[TextRun("Even header")])],
                even_page_footers=[Paragraph(content=[TextRun("Even footer")])],
                properties={
                    "different_first_page_header_footer": True,
                    "odd_and_even_pages_header_footer": True,
                },
            )
        ]
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)
    section = restored.sections[0]

    assert report.lossless
    assert section.properties["different_first_page_header_footer"] is True
    assert section.properties["odd_and_even_pages_header_footer"] is True
    assert section.headers[0].plain_text == "Odd header"
    assert section.footers[0].plain_text == "Odd footer"
    assert section.first_page_headers[0].plain_text == "First header"
    assert section.first_page_footers[0].plain_text == "First footer"
    assert section.even_page_headers[0].plain_text == "Even header"
    assert section.even_page_footers[0].plain_text == "Even footer"


def test_write_docx_model_preserves_inline_formula_order_and_simple_field(tmp_path):
    output = tmp_path / "ordered-content.docx"
    omml = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x+1</m:t></m:r></m:oMath>'
    document = DocumentModel(
        sections=[
            Section(
                footers=[
                    Paragraph(
                        content=[
                            TextRun("Page "),
                            TextRun("", properties={"field_instruction": "PAGE"}),
                        ]
                    )
                ],
                blocks=[
                    Paragraph(
                        content=[
                            TextRun("Score "),
                            Formula(omml, FormulaFormat.OMML, fallback_text="x+1"),
                            TextRun(" is final"),
                        ]
                    )
                ],
            )
        ]
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)

    assert report.lossless
    content = restored.sections[0].blocks[0].content
    assert [type(item) for item in content] == [TextRun, Formula, TextRun]
    assert content[0].text == "Score "
    assert content[1].fallback_text == "x+1"
    assert content[2].text == " is final"
    footer = restored.sections[0].footers[0].content
    assert footer[1].properties["field_instruction"] == "PAGE"


def test_write_docx_model_preserves_table_geometry_fill_and_repeating_header(tmp_path):
    output = tmp_path / "table-formatting.docx"
    document = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Table(
                        properties={
                            "style_name": "Table Grid",
                            "autofit": False,
                            "alignment": "center",
                            "grid_widths_twips": [2200, 3600],
                        },
                        rows=[
                            TableRow(
                                properties={"repeat_header": True},
                                cells=[
                                    TableCell(
                                        properties={
                                            "fill": "D9EAF7",
                                            "width_twips": 2200,
                                            "vertical_alignment": "center",
                                            "margins_twips": {"top": 80, "start": 120, "bottom": 80, "end": 120},
                                        },
                                        blocks=[Paragraph(content=[TextRun("Header")])],
                                    ),
                                    TableCell(
                                        properties={"fill": "D9EAF7", "width_twips": 3600},
                                        blocks=[Paragraph(content=[TextRun("Value")])],
                                    ),
                                ],
                            )
                        ],
                    )
                ]
            )
        ]
    )

    report = write_docx_model(document, output)
    restored = read_docx_model(output)
    table = restored.sections[0].blocks[0]

    assert report.lossless
    assert isinstance(table, Table)
    assert table.properties["autofit"] is False
    assert table.properties["alignment"] == "center"
    assert table.properties["grid_widths_twips"] == [2200, 3600]
    assert table.rows[0].properties["repeat_header"] is True
    assert table.rows[0].cells[0].properties["fill"] == "D9EAF7"
    assert table.rows[0].cells[0].properties["width_twips"] == 2200
    assert table.rows[0].cells[0].properties["vertical_alignment"] == "center"
    assert table.rows[0].cells[0].properties["margins_twips"] == {
        "top": 80,
        "start": 120,
        "bottom": 80,
        "end": 120,
    }
