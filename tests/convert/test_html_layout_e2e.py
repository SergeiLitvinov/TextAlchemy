"""Геометрия настоящего PPTX → HTML измеряется в Chromium."""

import pytest
from PIL import Image as PillowImage
from pptx import Presentation
from pptx.util import Pt

from tests import test_browser_e2e as fixtures
from textalchemy.convert.html_writer import write_html_model
from textalchemy.formats.pptx import read_pptx_model

browser, page = fixtures.browser, fixtures.page


@pytest.mark.parametrize("width", [375, 1280])
def test_pptx_origin_images_and_flipped_picture_keep_slide_coordinates(page, tmp_path, width):
    image = tmp_path / "picture.png"
    PillowImage.new("RGB", (120, 60), "red").save(image)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for name, x, y, w, h, flipped in [
        ("origin-first", 0, 0, 120, 60, False),
        ("origin-second", 0, 0, 60, 30, False),
        ("flipped", 180, 90, 120, 60, True),
    ]:
        picture = slide.shapes.add_picture(str(image), Pt(x), Pt(y), Pt(w), Pt(h))
        picture._element.nvPicPr.cNvPr.set("descr", name)
        if flipped:
            picture._element.spPr.xfrm.set("flipH", "1")
    table = slide.shapes.add_table(1, 1, Pt(0), Pt(0), Pt(120), Pt(30)).table
    table.cell(0, 0).text = "Origin table"
    source = tmp_path / "geometry.pptx"
    presentation.save(source)
    model = read_pptx_model(source)
    output = tmp_path / "geometry.html"
    report = write_html_model(model, output)
    assert report.success
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(output.as_uri())
    origin = page.locator(".ta-section").bounding_box()
    for name, x, y, w, h in [
        ("origin-first", 0, 0, 120, 60),
        ("origin-second", 0, 0, 60, 30),
        ("flipped", 180, 90, 120, 60),
    ]:
        bounds = page.locator(f'img[alt="{name}"]').bounding_box()
        assert bounds["x"] - origin["x"] == pytest.approx(x * 4 / 3, abs=0.5)
        assert bounds["y"] - origin["y"] == pytest.approx(y * 4 / 3, abs=0.5)
        assert bounds["width"] == pytest.approx(w * 4 / 3, abs=0.5)
        assert bounds["height"] == pytest.approx(h * 4 / 3, abs=0.5)
    bounds = page.locator(".ta-main > table").bounding_box()
    assert bounds["x"] - origin["x"] == pytest.approx(0, abs=0.5)
    assert bounds["y"] - origin["y"] == pytest.approx(0, abs=0.5)
    assert page.e2e_errors == []
