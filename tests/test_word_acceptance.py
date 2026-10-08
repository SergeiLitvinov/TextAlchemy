"""Opt-in acceptance in the installed Microsoft Word, through real Web routes."""

import io
import json
import os
import shutil
import subprocess
import sys
from hashlib import sha256
from importlib.metadata import version
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from tests import test_web_m4_completion as api_fixtures
from textalchemy.web.queue import task_queue

m4_client = api_fixtures.m4_client
SCRIPT = Path(__file__).resolve().parents[1] / "tools/acceptance/word-edit.ps1"


def fixture_document():
    document = Document()
    document.add_heading("Acceptance heading", level=2)
    document.add_paragraph().add_run("Editable body").bold = True
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip([table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)],
                          ["Item", "Value", "Own fixture", "9"], strict=True):
        cell.text = text
    image = io.BytesIO()
    Image.new("RGB", (20, 20), "green").save(image, format="PNG")
    image.seek(0)
    document.add_picture(image)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.skipif(sys.platform != "win32" or os.getenv("TEXTALCHEMY_WORD_ACCEPTANCE") != "1",
                    reason="Requires installed Word and explicit local acceptance run")
def test_real_word_edits_and_reopens_two_web_roundtrips(m4_client, tmp_path):
    client, store = m4_client
    original = fixture_document()
    source = original
    observations = []
    previous_text = "Editable body"
    for cycle in range(1, 3):
        imported = client.post("/api/convert", files={"file": ("source.docx", source)}, data={"target_format": "model"})
        assert imported.status_code == 200, imported.text
        assert task_queue.wait_idle(timeout=30)
        imported_id = imported.json()["task_id"]
        model = client.get(imported.json()["result"])
        assert model.status_code == 200, store.get(imported_id)
        exported = client.post("/api/convert", files={"file": ("model.json", model.content)}, data={"target_format": "docx"})
        assert exported.status_code == 200, exported.text
        assert task_queue.wait_idle(timeout=30)
        created = exported.json()
        task = store.get(created["task_id"])
        assert task["status"] == "done", task
        result = client.get(created["result"]).content
        input_path, edited_path = tmp_path / f"cycle-{cycle}.docx", tmp_path / f"edited-{cycle}.docx"
        input_path.write_bytes(result)
        replacement = f"Edited cycle {cycle}"
        completed = subprocess.run([
            shutil.which("pwsh") or "powershell", "-NoProfile", "-NonInteractive", "-File", str(SCRIPT),
            "-InputPath", str(input_path), "-OutputPath", str(edited_path),
            "-FindText", previous_text, "-Replacement", replacement,
        ], capture_output=True, text=True, encoding="utf-8", timeout=90, check=True)
        evidence = json.loads(completed.stdout.strip())
        assert evidence["program"]["name"] == "Microsoft Word"
        assert all(evidence["checks"][key] for key in (
            "text_edit", "table_cell_edit", "save_reopen", "inline_image", "bold")), evidence
        # Published adapters preserve the formal heading after Word edits and reimport.
        assert evidence["checks"]["heading_outline"] is True, evidence
        assert evidence["artifact_sha256"] == sha256(result).hexdigest()
        assert evidence["edited_sha256"] == sha256(edited_path.read_bytes()).hexdigest()
        assert input_path.read_bytes() == result
        assert store.store_target_program_check(created["task_id"], expected=task, evidence=evidence)
        assert store.get(created["task_id"]) == task
        payload = client.get(created["status"]).json()
        witness = payload["report"]["metrics"]["target_program_checks"][0]
        assert witness["editability_verified"] is True and witness["visual_score"] is None
        assert witness["program"] == evidence["program"]
        forecast = task["report"]["metrics"]["plan"]["editability_score"]
        observations.append({"cycle": cycle, "import_task_id": imported_id, "export_task_id": created["task_id"],
                             "prediction": forecast, "evidence": witness, "report": payload["report"]})
        source = edited_path.read_bytes()
        previous_text = replacement
    calibration = {"scope": "one_own_composite_docx_two_word_cycles", "source_sha256": sha256(original).hexdigest(),
                   "versions": {name: version(name) for name in ("textalchemy", "opendoc-model", "opendoc-formats")},
                   "observations": observations, "manual_gui_acceptance": None, "statistical_generalization": None}
    (tmp_path / "word-acceptance.json").write_text(json.dumps(calibration, ensure_ascii=False, indent=2), encoding="utf-8")


@pytest.mark.skipif(sys.platform != "win32" or os.getenv("TEXTALCHEMY_WORD_ACCEPTANCE") != "1",
                    reason="Requires installed Word and explicit local acceptance run")
def test_corpus_word_driver_verifies_new_edits_and_never_modifies_input(tmp_path):
    from tools.acceptance.corpus import word_check

    source = tmp_path / "own-fixture.docx"
    source.write_bytes(fixture_document())
    baseline = word_check(source)
    assert baseline["inventory"]["tables"] == 1
    for cycle in range(1, 3):
        before = source.read_bytes()
        edited = tmp_path / f"edited-{cycle}.docx"
        record = word_check(source, edited)
        assert all(record["checks"].values())
        assert record["before"]["normalized_text_sha256"] != record["after"]["normalized_text_sha256"]
        assert record["artifact_sha256"] == sha256(before).hexdigest()
        assert record["edited_sha256"] == sha256(edited.read_bytes()).hexdigest()
        assert record["visual_score"] is record["full_semantic_acceptance"] is None
        assert source.read_bytes() == before
        source = edited
