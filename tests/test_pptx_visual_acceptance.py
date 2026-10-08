"""Own PPTX witnesses distinguish observed loss from conservative diagnostics."""

import json
import os
import sys
import uuid

import pytest
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from tests.corpus.pptx_acceptance import write_visual_presentation
from textalchemy.convert.executor import ConversionExecutor, ConversionRequest
from textalchemy.core.quality_policy import QualityPolicy
from textalchemy.core.types import DocFormat
from tools.acceptance.corpus import ROOT, digest, target_program_check
from tools.acceptance.powerpoint_visual import compare


def group_count(shapes):
    return sum(1 + group_count(shape.shapes) for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)


def two_cycles(directory, case):
    source = directory / "source.pptx"
    write_visual_presentation(source, case)
    original = digest(source)
    executor = ConversionExecutor()
    current = source
    results = []
    for cycle in (1, 2):
        model = directory / f"cycle-{cycle}.json"
        imported = executor.execute(ConversionRequest(current, model, DocFormat.PPTX, DocFormat.MODEL))
        assert imported.success, imported.to_dict()
        result = directory / f"cycle-{cycle}.pptx"
        exported = executor.execute(ConversionRequest(model, result, DocFormat.MODEL, DocFormat.PPTX))
        assert exported.success, exported.to_dict()
        results.append((result, imported, exported))
        current = result
    assert digest(source) == original
    return source, results


@pytest.mark.parametrize("case", ["master-background", "nested-group"])
def test_two_cycles_record_actual_structure_separately_from_loss_labels(tmp_path, case):
    source, results = two_cycles(tmp_path, case)
    original = Presentation(source)
    if case == "master-background":
        assert original.slide_master.background.fill.type == MSO_FILL_TYPE.SOLID
        assert str(original.slide_master.background.fill.fore_color.rgb) == "184860"
    else:
        assert group_count(original.slides[0].shapes) == 2
    for result, imported, exported in results:
        output = Presentation(result)
        assert len(output.slides) == 1
        assert (output.slide_width, output.slide_height) == (original.slide_width, original.slide_height)
        for report in (imported, exported):
            assert any(i.feature == "pptx.package" and i.severity.value == "loss" for i in report.issues)
        if case == "master-background":
            # A green witness records the known loss; it does not accept the result's appearance.
            assert output.slide_master.background.fill.type == MSO_FILL_TYPE.BACKGROUND
        else:
            assert group_count(output.slides[0].shapes) == 2
            assert any(i.feature == "pptx.group" and i.severity.value == "loss" for i in imported.issues)


@pytest.mark.parametrize("case", ["freeform-axis-lines", "freeform-line", "table-theme", "table-style"])
def test_lines_and_table_controls_keep_observations_separate_from_diagnostics(tmp_path, case):
    source, results = two_cycles(tmp_path, case)
    original = Presentation(source).slides[0]
    if case == "freeform-axis-lines":
        contours = list(original.shapes)[1:]
        assert len(contours) == 2 and all(s.shape_type == MSO_SHAPE_TYPE.FREEFORM for s in contours)
        assert contours[0].height == 0 and contours[0].width > 0
        assert contours[1].width == 0 and contours[1].height > 0
    for cycle, (result, imported, exported) in enumerate(results, 1):
        output = Presentation(result).slides[0]
        if case == "freeform-axis-lines":
            assert len(output.shapes) == 3
            for before, after in zip(contours, list(output.shapes)[1:], strict=True):
                assert after.shape_type == MSO_SHAPE_TYPE.FREEFORM
                assert (after.left, after.top, after.width, after.height) == (
                    before.left, before.top, before.width, before.height)
            assert not any(i.feature == "pptx.shape" and i.severity.value == "loss" for i in imported.issues)
        elif case == "freeform-line":
            assert output.shapes[1].shape_type == MSO_SHAPE_TYPE.FREEFORM
            assert not any(i.feature == "pptx.shape" for i in imported.issues)
        else:
            table = output.shapes[1].table
            assert len(table.rows) == len(table.columns) == 2
            assert [cell.text for row in table.rows for cell in row.cells] == [
                "Own cell 1:1", "Own cell 1:2", "Own cell 2:1", "Own cell 2:2"]
            assert any(i.feature == "styles" and i.severity.value == "loss" for i in exported.issues)
            if case == "table-style":
                cell = table.cell(0, 0)
                run = cell.text_frame.paragraphs[0].runs[0]
                assert str(cell.fill.fore_color.rgb) == "184860"
                assert run.font.name == "Arial" and run.font.size.pt == 20 and run.font.bold is True
                assert str(run.font.color.rgb) == "FFFFFF"


@pytest.mark.parametrize("case", ["master-background", "nested-group", "freeform-axis-lines", "table-theme"])
def test_strict_budget_keeps_previous_native_result_and_source(tmp_path, case):
    source = tmp_path / "source.pptx"
    write_visual_presentation(source, case)
    original = source.read_bytes()
    result = tmp_path / "result.pptx"
    result.write_bytes(b"Previous accepted result")
    executor = ConversionExecutor()
    model = tmp_path / "model.json"
    imported = executor.execute(ConversionRequest(source, model, DocFormat.PPTX, DocFormat.MODEL))
    assert imported.success, imported.to_dict()
    report = executor.execute(ConversionRequest(
        model, result, DocFormat.MODEL, DocFormat.PPTX, quality_policy=QualityPolicy(0)))
    assert not report.success
    assert report.metrics["quality_gate"]["loss_issues"] > 0
    assert result.read_bytes() == b"Previous accepted result"
    assert source.read_bytes() == original


@pytest.mark.skipif(sys.platform != "win32" or os.getenv("TEXTALCHEMY_POWERPOINT_ACCEPTANCE") != "1",
                    reason="Requires installed PowerPoint and explicit local acceptance run")
@pytest.mark.parametrize("case", ["master-background", "nested-group", "freeform-axis-lines", "freeform-line",
                                  "table-theme", "table-style"])
def test_native_visual_two_cycles_on_own_minimal_cases(tmp_path, case):
    source, results = two_cycles(tmp_path, case)
    directory = ROOT / ".textalchemy/acceptance/minimal-pptx-tests" / uuid.uuid4().hex
    directory.mkdir(parents=True)
    for label, path in [("source", source), *[(f"cycle-{i}", r[0]) for i, r in enumerate(results, 1)]]:
        before = digest(path)
        record = target_program_check(path, program="powerpoint", render_directory=directory / label)
        record["input_unchanged"] = before == digest(path) == record["artifact_sha256"]
        assert record["input_unchanged"]
        (directory / f"{label}.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    first = compare(directory, "source", "cycle-1")
    second = compare(directory, "cycle-1", "cycle-2")
    assert first["all_slides_compared"] and second["all_slides_compared"]
    assert first["all_pixels_equal"] is (case in {"nested-group", "freeform-line", "freeform-axis-lines", "table-style"})
    assert second["all_pixels_equal"] is True
    assert first["full_visual_acceptance"] is None
    (directory / "visual-comparison.json").write_text(json.dumps([first, second], indent=2), encoding="utf-8")
