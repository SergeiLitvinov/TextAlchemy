"""PDF block order must preserve all imported payloads and resources."""
import copy
import importlib
import io

import fitz
import pytest
from fastapi.testclient import TestClient

from textalchemy.web.services import pdf_order
from textalchemy.web.tasks import TaskStore


@pytest.fixture
def pdf_draft(tmp_path, monkeypatch):
    web = importlib.import_module('textalchemy.web.app')
    store = TaskStore(tmp_path / 'tasks')
    monkeypatch.setattr(web, 'tasks_store', store)
    pdf = fitz.open()
    from PIL import Image

    raster = io.BytesIO()
    Image.new('RGB', (20, 20), 'red').save(raster, format='PNG')
    for words in [('Second', 'First'), ('Other page', 'End')]:
        page = pdf.new_page()
        page.insert_text((70, 150), words[0], fontsize=14)
        page.insert_text((70, 250), words[1], fontsize=18)
        page.insert_image(fitz.Rect(70, 350, 130, 410), stream=raster.getvalue())
    content = pdf.tobytes()
    pdf.close()
    with TestClient(web.app) as client:
        response = client.post('/api/pdf-order', files={'file': ('order.pdf', content, 'application/pdf')})
        assert response.status_code == 200, response.text
        data = response.json()
        yield client, '/api/pdf-order/' + data['draft_id'], pdf_order.store_for(store), data


def test_reorder_preserves_model_and_exports_docx(pdf_draft):
    from docx import Document

    client, url, store, data = pdf_draft
    before = copy.deepcopy(pdf_order.load(store, data['draft_id']))
    order = copy.deepcopy(before['order'])
    order[0].reverse()
    saved = client.put(url, json={'revision': 1, 'order': order})
    assert saved.status_code == 200
    after = pdf_order.load(TaskStore(store.root), data['draft_id'])
    expected = copy.deepcopy(before['model'])
    expected['document']['sections'][0]['blocks'].reverse()
    assert after['model'] == expected
    assert expected['document']['resources']
    assert after['order'] == order
    assert client.get(url).json()['revision'] == 2
    assert client.get(url + '/export?revision=2&format=model').json() == expected
    exported = client.get(url + '/export?revision=2&format=docx')
    assert exported.status_code == 200
    paragraphs = [p.text for p in Document(io.BytesIO(exported.content)).paragraphs]
    assert paragraphs.index('First') < paragraphs.index('Second') < paragraphs.index('Other page')
    image = client.get(url + '/pages/0')
    assert image.status_code == 200
    assert image.content.startswith(b'\x89PNG')
    assert client.get(url + '/pages/2').status_code == 404
    assert client.get(url + '/export?revision=1').status_code == 409
    assert client.put(url, json={'revision': 1, 'order': before['order']}).status_code == 409
    assert pdf_order.load(store, data['draft_id'])['model'] == expected


@pytest.mark.parametrize('invalid', ['missing', 'duplicate', 'foreign', 'page', 'pages'])
def test_invalid_permutation_is_atomic(pdf_draft, invalid):
    client, url, store, data = pdf_draft
    before = pdf_order.load(store, data['draft_id'])
    order = copy.deepcopy(before['order'])
    if invalid == 'missing':
        order[0].pop()
    elif invalid == 'duplicate':
        order[0][0] = order[0][1]
    elif invalid == 'foreign':
        order[0][0] = 'unknown'
    elif invalid == 'page':
        order[0], order[1] = order[1], order[0]
    else:
        order.pop()
    assert client.put(url, json={'revision': 1, 'order': order}).status_code == 422
    assert pdf_order.load(store, data['draft_id']) == before


def test_expiry_and_failed_write(pdf_draft, monkeypatch):
    client, url, store, data = pdf_draft
    before = pdf_order.load(store, data['draft_id'])
    order = copy.deepcopy(before['order'])
    order[0].reverse()

    def fail(*args, **kwargs):
        raise OSError('disk unavailable')

    with monkeypatch.context() as patch:
        patch.setattr(TaskStore, 'set', fail)
        assert client.put(url, json={'revision': 1, 'order': order}).status_code == 503
    assert pdf_order.load(store, data['draft_id']) == before
    with monkeypatch.context() as patch:
        patch.setattr('textalchemy.web.routes.pdf_order._store', lambda: TaskStore(store.root, ttl_seconds=-1))
        assert client.get(url).status_code == 404
        assert client.put(url, json={'revision': 1, 'order': order}).status_code == 404


def test_concurrent_orders_and_bad_uploads(pdf_draft):
    from concurrent.futures import ThreadPoolExecutor

    client, url, store, data = pdf_draft
    order = pdf_order.load(store, data['draft_id'])['order']
    other = copy.deepcopy(order)
    other[0].reverse()
    with ThreadPoolExecutor(2) as pool:
        responses = list(pool.map(lambda value: client.put(url, json={'revision': 1, 'order': value}), [order, other]))
    assert sorted(response.status_code for response in responses) == [200, 409]
    assert client.get(url).json()['revision'] == 2
    for name, content in [('broken.pdf', b'not a PDF'), ('note.txt', b'text')]:
        assert client.post('/api/pdf-order', files={'file': (name, content)}).status_code == 422
    assert client.get(url + '/export?revision=2&format=exe').status_code == 422


@pytest.mark.parametrize('role,style,level', [('heading1', 'Heading 1', 1), ('heading2', 'Heading 2', 2),
                                            ('heading3', 'Heading 3', 3), ('paragraph', 'Normal', None)])
def test_classification_survives_export_and_preserves_objects(pdf_draft, role, style, level):
    from docx import Document
    from opendoc_model import DocumentModel, TextStyle, document_from_dict, document_to_dict, get_heading

    client, url, store, data = pdf_draft
    before = pdf_order.load(store, data['draft_id'])
    block_id = before['order'][0][0]
    order = copy.deepcopy(before['order'])
    order[0].reverse()
    response = client.put(url, json={'revision': 1, 'order': order, 'classifications': {block_id: role}})
    assert response.status_code == 200, response.text
    expected = copy.deepcopy(before['model'])
    definition = TextStyle(properties={'style_name': style, 'style_type': 'paragraph'})
    encoded_style = document_to_dict(DocumentModel(styles={style: definition}))['document']['styles'][style]
    expected['document']['styles'].setdefault(style, encoded_style)
    block = expected['document']['sections'][0]['blocks'][0]
    block['style_id'] = style
    block['properties'].update(style_name=style, legacy_type='heading' if level else 'paragraph', pdf_editor_role=role)
    if level:
        block['properties'].update(heading_level=level, level=level)
        block['properties']['opendoc.heading'] = {'format': 'opendoc.heading', 'version': 1, 'level': level}
    else:
        block['properties'].pop('heading_level', None)
        block['properties'].pop('level', None)
    expected['document']['sections'][0]['blocks'].reverse()
    assert client.get(url + '/export?revision=2&format=model').json() == expected
    assert not document_from_dict(expected).validate()
    heading = get_heading(document_from_dict(expected).sections[0].blocks[-1])
    assert (heading.level if heading else None) == level
    restored = client.get(url).json()
    assert next(b for b in restored['pages'][0]['blocks'] if b['id'] == block_id)['classification'] == role
    exported = client.get(url + '/export?revision=2&format=docx')
    assert exported.status_code == 200
    paragraph = next(p for p in Document(io.BytesIO(exported.content)).paragraphs if p.text == 'Second')
    assert paragraph.style.name == style
    if level:
        from docx.oxml.ns import qn

        assert paragraph.style.element.xpath('./w:pPr/w:outlineLvl')[0].get(qn('w:val')) == str(level - 1)
    # Demotion must remove heading semantics, even after reordering and reloading.
    response = client.put(url, json={'revision': 2, 'order': order, 'classifications': {block_id: 'paragraph'}})
    assert response.status_code == 200
    model = client.get(url + '/export?revision=3&format=model').json()
    block = model['document']['sections'][0]['blocks'][-1]
    assert 'heading_level' not in block['properties']
    assert 'opendoc.heading' not in block['properties']
    assert block['style_id'] == 'Normal'


@pytest.mark.parametrize('invalid', ['unknown', 'image', 'role'])
def test_invalid_classification_keeps_entire_revision(pdf_draft, invalid):
    client, url, store, data = pdf_draft
    before = pdf_order.load(store, data['draft_id'])
    order = copy.deepcopy(before['order'])
    order[0].reverse()
    changes = {before['order'][0][0]: 'heading1'}
    if invalid == 'unknown':
        changes['unknown'] = 'heading2'
    elif invalid == 'image':
        image = next(b for b in data['pages'][0]['blocks'] if not b['classifiable'])
        changes[image['id']] = 'heading2'
    else:
        changes[before['order'][0][0]] = 'table'
    assert client.put(url, json={'revision': 1, 'order': order, 'classifications': changes}).status_code == 422
    assert pdf_order.load(store, data['draft_id']) == before
