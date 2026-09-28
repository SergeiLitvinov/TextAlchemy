"""Conditional paragraph editing, persistent copies and both output branches."""
import importlib
import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from textalchemy.web.services.template_variables import custom_path


@pytest.fixture
def condition_client(tmp_path, monkeypatch):
    templates = tmp_path / 'templates'
    templates.mkdir()
    path = templates / 'example.docx'
    doc = Document()
    doc.add_paragraph('Always visible')
    doc.add_paragraph().add_run('Optional {{ title }}').bold = True
    doc.add_paragraph('Final paragraph')
    doc.save(path)
    schema = {'fields': [{'name': 'title', 'type': 'string', 'required': True}], 'allow_extra': False}
    path.with_suffix('.schema.json').write_text(json.dumps(schema), encoding='utf-8')
    monkeypatch.setattr('textalchemy.generate.template.TEMPLATES_DIR', templates)
    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path / 'data')
    with TestClient(importlib.import_module('textalchemy.web.app').app) as client:
        yield client, path


def save(client, name='example', **overrides):
    inspection = client.get(f'/api/generate/templates/{name}/conditions').json()
    block = next(item for item in inspection['blocks'] if item['text'].startswith('Optional'))
    payload = {'block': block['id'], 'field': 'show_section', 'negate': False, 'revision': inspection['revision'], **overrides}
    return client.post(f'/api/generate/templates/{name}/conditions', json=payload)


def test_condition_copy_reopen_change_and_generate(condition_client):
    client, path = condition_client
    original = path.read_bytes()
    response = save(client)
    assert response.status_code == 200, response.text
    name = response.json()['name']
    assert path.read_bytes() == original
    schema = client.get(f'/api/generate/templates/{name}/schema').json()['schema']
    assert schema['fields'][-1]['type'] == 'boolean'
    assert schema['fields'][-1]['default'] is False
    inspection = client.get(f'/api/generate/templates/{name}/conditions').json()
    assert inspection['blocks'][1]['field'] == 'show_section'
    for value in (False, True):
        result = client.post('/api/generate', data={'template': name, 'params': json.dumps({
            'title': 'Result', 'show_section': value})})
        assert result.headers['content-type'].startswith('application/vnd.openxmlformats'), result.text
        doc = Document(io.BytesIO(result.content))
        texts = [p.text for p in doc.paragraphs]
        assert ('Optional Result' in texts) is value
        assert texts[0] == 'Always visible' and texts[-1] == 'Final paragraph'
        if value:
            assert next(p for p in doc.paragraphs if p.text == 'Optional Result').runs[0].bold
    changed = save(client, name, negate=True).json()['name']
    paragraphs = [p.text for p in Document(custom_path(changed)).paragraphs]
    assert paragraphs.count('{% if not show_section %}') == 1
    assert paragraphs.count('{% endif %}') == 1
    for value in (False, True):
        result = client.post('/api/generate', data={'template': changed, 'params': json.dumps({
            'title': 'Inverted', 'show_section': value})})
        assert ('Optional Inverted' in [p.text for p in Document(io.BytesIO(result.content)).paragraphs]) is not value


@pytest.mark.parametrize('overrides', [{'field': 'title'}, {'field': 'if'}, {'field': '../path'}, {'block': 999},
                                       {'revision': 'stale'}])
def test_invalid_condition_does_not_publish(condition_client, overrides):
    client, path = condition_client
    original = path.read_bytes()
    assert save(client, **overrides).status_code == 422
    assert path.read_bytes() == original
    assert len(client.get('/api/generate/templates').json()) == 1


def test_complex_condition_rejected(condition_client):
    client, path = condition_client
    doc = Document(path)
    doc.add_paragraph('{% for item in items %}')
    doc.add_paragraph('{{ item }}')
    doc.add_paragraph('{% endfor %}')
    doc.save(path)
    assert client.get('/api/generate/templates/example/conditions').status_code == 422
