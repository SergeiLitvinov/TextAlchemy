"""Тесты структурной инспекции документов."""

from textalchemy.core.document_model import (
    Box,
    DocumentModel,
    Formula,
    FormulaFormat,
    Image,
    Paragraph,
    Provenance,
    Resource,
    ResourceKind,
    Section,
    Table,
    TableCell,
    TableRow,
    TextRun,
    TextStyle,
)
from textalchemy.core.inspection import DocumentInspection, compare_inspections, inspect_document_model, inspect_path


def test_inspect_txt_has_text_structure_but_no_invented_page_geometry(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("Первая\n\nПоследняя", encoding="utf-8")
    report = inspect_path(source)
    assert report.source_format == "txt"
    assert report.valid
    assert report.metrics["paragraphs"] == 3
    assert len(report.objects) == 3
    assert report.pages == []
    assert "pages" not in report.metrics


def test_inspect_document_model_counts_nested_features_and_resources():
    document = DocumentModel(
        resources={
            "used": Resource("used", ResourceKind.VECTOR_IMAGE, "image/svg+xml", data=b"<svg/>"),
            "unused": Resource("unused", ResourceKind.ATTACHMENT, "application/octet-stream", data=b"x"),
        },
        sections=[
            Section(
                blocks=[
                    Paragraph(
                        content=[
                            TextRun("Linked", TextStyle(font_family="Arial"), link="https://example.com"),
                            Image("used"),
                            Formula("x^2", FormulaFormat.LATEX),
                        ]
                    ),
                    Table(
                        rows=[
                            TableRow(
                                cells=[
                                    TableCell(
                                        blocks=[Paragraph(content=[TextRun("Cell")])],
                                        column_span=2,
                                    )
                                ]
                            )
                        ]
                    ),
                ]
            )
        ],
    )

    report = inspect_document_model(document)
    payload = report.to_dict()

    assert report.valid is True
    assert report.has_warnings is True
    assert report.metrics["paragraphs"] == 2
    assert report.metrics["tables"] == 1
    assert report.metrics["table_cells"] == 1
    assert report.metrics["merged_cells"] == 1
    assert report.metrics["images"] == 1
    assert report.metrics["formulas"] == 1
    assert report.metrics["hyperlinks"] == 1
    assert report.fonts == {"Arial": 1}
    assert report.formula_formats == {"latex": 1}
    assert payload["resources"][0]["sha256"]
    assert {issue.feature for issue in report.issues} == {
        "formula-fallback",
        "image-alt-text",
        "unused-resource",
    }


def test_inspect_document_model_reports_invalid_reference_and_geometry():
    document = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Image("missing", alt_text="Missing", box=Box(-2, 0, -1, 10)),
                ]
            )
        ]
    )

    report = inspect_document_model(document)

    assert report.valid is False
    assert any(issue.feature == "model-validation" for issue in report.issues)
    assert any(issue.feature == "model-validation" and "box.width" in issue.location for issue in report.issues)
    assert not report.objects  # Invalid models have no trustworthy inventory.


def test_inspect_pdf_reports_text_vector_fonts_and_page_geometry(tmp_path):
    import fitz

    source = tmp_path / "source.pdf"
    document = fitz.open()
    page = document.new_page(width=400, height=300)
    page.insert_text((40, 60), "Inspectable PDF")
    page.draw_rect(fitz.Rect(40, 80, 180, 140))
    document.save(source)
    document.close()

    report = inspect_path(source)

    assert report.valid is True
    assert report.source_format == "pdf"
    assert report.metrics["pages"] == 1
    assert report.metrics["characters"] >= len("Inspectable PDF")
    assert report.metrics["vector_drawings"] >= 1
    assert report.pages[0]["width_pt"] == 400
    assert report.fonts


def test_inspect_pptx_uses_shared_document_model(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Inspectable presentation"
    presentation.save(source)

    report = inspect_path(source)

    assert report.valid is True
    assert report.source_format == "pptx"
    assert report.metrics["pages"] == 1
    assert report.metrics["paragraphs"] == 1
    assert report.metrics["characters"] == len("Inspectable presentation")
    assert report.pages[0]["width_pt"] > 0


def test_compare_inspections_reports_retention_and_page_geometry_loss():
    source = DocumentModel(
        sections=[
            Section(
                blocks=[
                    Paragraph(content=[TextRun("Long source text")]),
                    Table(rows=[TableRow(cells=[TableCell(blocks=[Paragraph(content=[TextRun("Cell")])])])]),
                ]
            )
        ]
    )
    target = DocumentModel(
        sections=[
            Section(
                blocks=[Paragraph(content=[TextRun("Short")])],
            )
        ]
    )
    target.sections[0].page.width.pt += 10
    target.sections[0].page.margin_top.pt += 3

    comparison = compare_inspections(inspect_document_model(source), inspect_document_model(target))

    assert comparison.valid is True
    assert comparison.has_losses is True
    assert comparison.retention["tables"]["ratio"] == 0
    assert comparison.retention["characters"]["target"] < comparison.retention["characters"]["source"]
    assert comparison.page_geometry[0]["same_size"] is False
    assert comparison.page_geometry[0]["same_margins"] is False
    assert comparison.page_geometry[0]["dimension_error_pt"] == 10
    assert comparison.geometry_summary["max_dimension_error_pt"] == 10
    assert comparison.geometry_summary["max_margin_error_pt"] == 3
    assert comparison.geometry_summary["rms_margin_error_pt"] == 1.5
    assert any(issue.feature == "page-geometry" for issue in comparison.issues)
    assert any(issue.feature == "page-margins" for issue in comparison.issues)


def test_compare_inspections_reports_object_level_changes_and_recommendations():
    origin = Provenance(source_format="pptx", source_path="slides.pptx", page=1, object_id="shape-7")
    lost_origin = Provenance(source_format="pptx", source_path="slides.pptx", page=1, object_id="shape-8")
    source = DocumentModel(sections=[Section(blocks=[
        Paragraph(content=[TextRun("Title")], box=Box(10, 20, 100, 30), provenance=origin),
        Paragraph(content=[TextRun("Lost")], provenance=lost_origin),
    ])])
    target = DocumentModel(sections=[Section(blocks=[
        Paragraph(content=[TextRun("Title changed")], box=Box(14, 20, 100, 30), provenance=origin),
    ])])

    comparison = compare_inspections(inspect_document_model(source), inspect_document_model(target))

    object_diff = comparison.object_diff
    assert object_diff["source_count"] == 2
    assert object_diff["target_count"] == 1
    assert object_diff["retention_ratio"] == 0.5
    assert object_diff["changed"][0]["changes"] == ["content_hash", "geometry"]
    assert object_diff["lost"][0]["provenance"]["object_id"] == "shape-8"
    assert {item["code"] for item in object_diff["recommendations"]} == {
        "restore-lost-objects", "review-object-geometry",
    }
    assert any(issue.feature == "object-loss" for issue in comparison.issues)


def test_inspect_html_extracts_structure_resources_and_fonts(tmp_path):
    import base64

    png = base64.b64encode(b"\x89PNG\r\n\x1a\nfakepixelbytes").decode("ascii")
    html = f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>Page</title><style>body {{ font-family: Arial, sans-serif; }}</style></head>
<body><section class="ta-section ta-section-0"><main class="ta-main">
<h1 style="font-family:'Cambria Math'">Заголовок</h1>
<p>Первый абзац с <a href="https://example.com">ссылкой</a> и <a href="#bookmark">внутренней</a>.</p>
<table><tbody><tr><td>Ячейка 1</td><td>Ячейка 2</td></tr></tbody></table>
<p><img src="data:image/png;base64,{png}" alt="pic"></p>
<p><span class="ta-formula"><math xmlns="http://www.w3.org/1998/Math/MathML"><mi>x</mi></math></span></p>
</main></section></body></html>"""
    source = tmp_path / "page.html"
    source.write_text(html, encoding="utf-8")

    report = inspect_path(source)

    assert report.valid is True
    assert report.source_format == "html"
    assert report.metrics["sections"] == 1
    assert report.metrics["headings"] == 1
    assert report.metrics["paragraphs"] == 3
    assert report.metrics["tables"] == 1
    assert report.metrics["table_rows"] == 1
    assert report.metrics["table_cells"] == 2
    assert report.metrics["images"] == 1
    assert report.metrics["formulas"] == 1
    assert report.metrics["hyperlinks"] == 1
    assert report.metrics["internal_hyperlinks"] == 1
    assert report.metrics["characters"] > 0
    assert report.metrics["text_runs"] > 0
    assert report.resources[0]["media_type"] == "image/png"
    assert report.resources[0]["size_bytes"] == 22
    assert report.resources[0]["sha256"]
    assert report.fonts == {"Cambria Math": 1}
    assert report.formula_formats == {"mathml": 1}


def test_inspect_latex_extracts_structure_and_papersize(tmp_path):
    tex = r"""\documentclass[12pt,a4paper]{article}
\title{Документ}
\begin{document}
\section{Введение}
Текст введения с формулой $x^2 + y^2 = z^2$.

\subsection{Раздел}
\href{https://example.com}{Ссылка} и цитата \cite{key1}.

\begin{equation}
E = mc^2
\end{equation}

\begin{itemize}
\item Первый пункт
\item Второй пункт
\end{itemize}

\begin{tabular}{|l|l|}
\hline
A & B \\
\hline
\end{tabular}

Картинка: \includegraphics[width=3cm]{figure.png}.
\end{document}
"""
    source = tmp_path / "paper.tex"
    source.write_text(tex, encoding="utf-8")

    report = inspect_path(source)

    assert report.valid is True
    assert report.source_format == "latex"
    assert report.metrics["pages"] == 1
    assert report.metrics["headings"] == 2
    assert report.metrics["tables"] == 1
    assert report.metrics["images"] == 1
    assert report.metrics["formulas"] == 2
    assert report.metrics["hyperlinks"] == 1
    assert report.metrics["internal_hyperlinks"] == 0
    assert report.metrics["citations"] == 1
    assert report.metrics["list_items"] == 2
    assert report.metrics["characters"] > 0
    assert report.metrics["paragraphs"] > 0
    assert report.pages == [{"index": 0, "width_pt": 595.28, "height_pt": 841.89}]
    assert report.formula_formats == {"latex": 2}


def test_compare_pageless_html_target_avoids_page_geometry_loss(tmp_path):
    source = DocumentModel(sections=[Section(blocks=[Paragraph(content=[TextRun("Hello")])])])
    html = "<html><body><section><main><p>Hello</p></main></section></body></html>"
    target_file = tmp_path / "out.html"
    target_file.write_text(html, encoding="utf-8")

    comparison = compare_inspections(inspect_document_model(source), inspect_path(target_file))

    assert comparison.retention["pages"]["ratio"] == 1.0
    assert comparison.retention["characters"]["ratio"] == 1.0
    assert comparison.page_geometry == []
    assert not any(issue.feature == "page-geometry" for issue in comparison.issues)
    assert comparison.geometry_summary["available"] is False


def test_compare_latex_target_papersize_geometry(tmp_path):
    source = DocumentModel(sections=[Section()])
    tex = "\\documentclass[12pt,a4paper]{article}\n\\begin{document}\nПривет, мир.\n\\end{document}\n"
    target_file = tmp_path / "out.tex"
    target_file.write_text(tex, encoding="utf-8")

    comparison = compare_inspections(inspect_document_model(source), inspect_path(target_file))

    assert comparison.retention["pages"]["ratio"] == 1.0
    assert comparison.target.pages[0]["width_pt"] == 595.28
    assert comparison.target.pages[0]["height_pt"] == 841.89
    assert comparison.page_geometry == []
    assert comparison.geometry_summary["available"] is False
    assert comparison.geometry_summary["max_dimension_error_pt"] is None


def test_compare_inspections_reports_exact_resource_and_font_losses():
    source = DocumentInspection(
        None,
        "source",
        resources=[
            {
                "id": "diagram",
                "kind": "vector_image",
                "media_type": "image/svg+xml",
                "size_bytes": 100,
                "sha256": "a" * 64,
            },
            {
                "id": "photo",
                "kind": "raster_image",
                "media_type": "image/png",
                "size_bytes": 300,
                "sha256": "b" * 64,
            },
        ],
        fonts={"ABCDEF+ArialMT": 2, "Calibri": 3},
    )
    target = DocumentInspection(
        None,
        "target",
        resources=[
            {
                "id": "diagram",
                "kind": "raster_image",
                "media_type": "image/png",
                "size_bytes": 80,
                "sha256": "c" * 64,
            },
            {
                "id": "photo-copy",
                "kind": "raster_image",
                "media_type": "image/png",
                "size_bytes": 300,
                "sha256": "b" * 64,
            },
        ],
        fonts={"Arial": 2, "Liberation Sans": 3},
    )

    # Complete public model inspection supplies the evidence required by OpenDoc.
    def complete_inspection(report):
        model = DocumentModel(sections=[Section()])
        for resource in report.resources:
            model.add_resource(Resource(
                resource["id"], ResourceKind(resource["kind"]), resource["media_type"],
                data=resource["sha256"][0].encode() * resource["size_bytes"],
            ))
            model.sections[0].blocks.append(Image(resource["id"], alt_text="Resource"))
        for family, count in report.fonts.items():
            model.sections[0].blocks.append(Paragraph([
                TextRun("Text", TextStyle(font_family=family)) for _ in range(count)
            ]))
        return inspect_document_model(model)

    comparison = compare_inspections(complete_inspection(source), complete_inspection(target))

    resources = comparison.resource_comparison
    assert resources["exact_hash_matches"] == 1
    assert resources["exact_hash_retention_ratio"] == 0.5
    assert resources["exact_bytes_retained"] == 300
    assert resources["exact_byte_retention_ratio"] == 0.75
    assert resources["media_type_retention_ratio"] == 0.5
    assert resources["lost_resources"][0]["id"] == "diagram"
    assert resources["changed_ids"][0]["id"] == "diagram"
    fonts = comparison.font_comparison
    assert fonts["preserved_families"] == ["Arial"]
    assert fonts["missing_families"] == {"Calibri": 3}
    assert fonts["added_families"] == {"liberation sans": 3}
    assert fonts["exact_run_retention_ratio"] == 0.4
    assert fonts["possible_substitutions"] == [
        {"source": "Calibri", "target": "liberation sans", "runs": 3}
    ]
    assert any(issue.feature == "resource-loss" for issue in comparison.issues)
    assert any(issue.feature == "font-substitution" for issue in comparison.issues)
