"""Интеграционные тесты PPTX → DocumentModel (импортёр на общей модели)."""

from pathlib import Path

import pytest
from lxml import etree
from PIL import Image as PillowImage
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from textalchemy.core.color import ColorValue
from textalchemy.core.document_model import Formula, Image, Paragraph, ResourceKind, Table
from textalchemy.formats import pptx as pptx_mod
from textalchemy.formats.pptx import read_pptx, read_pptx_model

CORPUS_PPTX = Path(__file__).resolve().parent.parent / "corpus" / "office" / "libreoffice-scientific-slides.pptx"


def _png_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "_probe.png"
    PillowImage.new("RGB", (12, 8), "green").save(path)
    return path.read_bytes()


def _build_rich_pptx(path: Path, *, image_bytes: bytes | None = None) -> None:
    """Собрать PPTX с текстом, фигурами, группой, формулой, таблицей, диаграммой."""
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = "Rich deck"

    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2)

    box = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(6), Inches(0.6))
    first = box.text_frame.paragraphs[0]
    run = first.add_run()
    run.text = "Mixed formatting "
    run.font.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)
    link_run = first.add_run()
    link_run.text = "and link"
    link_run.font.italic = True
    link_run.hyperlink.address = "https://example.com"

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1.4), Inches(3), Inches(0.8))
    shape.text = "Filled shape"
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x5B, 0x9B, 0xD5)

    group = slide.shapes.add_group_shape()
    group.left = Emu(457200)
    group.top = Emu(457200)
    group.width = Emu(1828800)
    group.height = Emu(914400)
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(457200), Emu(457200), Emu(914400), Emu(457200)).text = "in group"

    math_box = slide.shapes.add_textbox(Inches(7), Inches(0.5), Inches(4), Inches(0.6))
    omath = etree.fromstring(
        b'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        b"<m:r><m:t>x</m:t></m:r>"
        b"<m:sSup>"
        b"<m:e><m:r><m:t>y</m:t></m:r></m:e>"
        b"<m:sup><m:r><m:t>2</m:t></m:r></m:sup>"
        b"</m:sSup></m:oMath>"
    )
    math_box.text_frame.paragraphs[0]._p.append(omath)

    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2.6), Inches(4), Inches(1.2)).table
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    chart_data = ChartData()
    chart_data.categories = ["Alpha", "Beta"]
    chart_data.add_series("Score", (10, 20))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(6),
        Inches(2.6),
        Inches(5),
        Inches(3),
        chart_data,
    ).chart
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor(0x44, 0x72, 0xC4)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.category_axis.has_title = True
    chart.category_axis.axis_title.text_frame.text = "Categories"
    chart.value_axis.has_title = True
    chart.value_axis.axis_title.text_frame.text = "Score axis"

    slide.notes_slide.notes_text_frame.text = "Speaker notes for test."

    if image_bytes is not None:
        image_path = path.parent / "_probe_pic.png"
        image_path.write_bytes(image_bytes)
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(4), Inches(2), Inches(1))
        image_path.unlink(missing_ok=True)

    presentation.save(path)


def _build_transformed_group_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(0.5)).text = "matrix child"
    group.left = Inches(1)
    group.top = Inches(1)
    group.width = Inches(2)
    group.height = Inches(1)
    xfrm = group._element.grpSpPr.xfrm
    xfrm.set("rot", "5400000")
    xfrm.set("flipH", "1")
    presentation.save(path)


def _build_nested_transformed_group_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(0.5)).text = "nested child"
    inner.left = Inches(0.5)
    inner.top = Inches(0.5)
    inner.width = Inches(2)
    inner.height = Inches(1)
    inner._element.grpSpPr.xfrm.set("rot", "900000")
    inner._element.grpSpPr.xfrm.set("flipV", "1")
    outer.left = Inches(1)
    outer.top = Inches(1)
    outer.width = Inches(4)
    outer.height = Inches(2)
    outer._element.grpSpPr.xfrm.set("rot", "1800000")
    outer._element.grpSpPr.xfrm.set("flipH", "1")
    presentation.save(path)


def _paragraphs(blocks):
    return [block for block in blocks if isinstance(block, Paragraph)]


class TestReadPptxModel:
    def test_builds_sections_metadata_and_blocks(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)

        assert model.source_format == "pptx"
        assert len(model.sections) == 1
        assert model.metadata["title"] == "Rich deck"
        assert model.metadata["slide_width_pt"] == pytest.approx(13.333333 * 72, abs=1.0)
        assert model.validate() == []

    def test_preserves_run_formatting_and_hyperlink(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        paragraphs = _paragraphs(model.sections[0].blocks)
        text_box = next(p for p in paragraphs if any(run.text.startswith("Mixed") for run in p.content))
        runs = [run for run in text_box.content if run.text.strip()]

        assert runs[0].style.bold is True
        assert runs[0].style.font_size.pt == 20
        assert isinstance(runs[0].style.color, ColorValue)
        assert runs[0].style.color.to_hex() == "#AA0000"
        assert runs[1].style.italic is True
        assert runs[1].link == "https://example.com"

    def test_keeps_shape_geometry_and_fill(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        shape_block = next(
            p for p in _paragraphs(model.sections[0].blocks) if p.properties["pptx"]["shape"].get("prst") == "roundRect"
        )
        assert shape_block.box.x == pytest.approx(72.0)
        assert shape_block.properties["pptx"]["shape"]["fill"] == "#5B9BD5"
        assert ColorValue.from_dict(shape_block.properties["pptx"]["shape"]["fill_color"]).to_hex() == "#5B9BD5"

    def test_shape_gradient_keeps_canonical_stops_and_alpha(self):
        element = etree.fromstring(
            f'<p:sp xmlns:p="{pptx_mod._P_NS}" xmlns:a="{pptx_mod._A_NS}"><p:spPr>'
            '<a:prstGeom prst="rect"/><a:gradFill><a:gsLst>'
            '<a:gs pos="0"><a:srgbClr val="000000"/></a:gs>'
            '<a:gs pos="100000"><a:schemeClr val="accent1"><a:alpha val="50000"/></a:schemeClr></a:gs>'
            "</a:gsLst></a:gradFill></p:spPr></p:sp>"
        )

        metadata = pptx_mod._shape_metadata(element, pptx_mod.Box(0, 0, 10, 10), {"accent1": "204060"})

        assert metadata is not None
        stops = metadata["gradient_colors"]
        assert stops[0]["position"] == 0.0
        assert ColorValue.from_dict(stops[1]["color"]).to_hex(include_alpha=True) == "#20406080"

    def test_group_child_resolves_absolute_geometry(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        group_block = next(
            p
            for p in _paragraphs(model.sections[0].blocks)
            if any(run.text == "in group" for run in p.content if hasattr(run, "text"))
        )
        assert group_block.box.x == pytest.approx(36.0)
        assert group_block.box.y == pytest.approx(36.0)
        assert group_block.box.width == pytest.approx(72.0)

    def test_group_rotation_and_flip_produce_exact_affine_matrix(self, tmp_path):
        path = tmp_path / "transformed-group.pptx"
        _build_transformed_group_pptx(path)

        model = read_pptx_model(path)
        child = next(
            paragraph
            for paragraph in _paragraphs(model.sections[0].blocks)
            if any(run.text == "matrix child" for run in paragraph.content if hasattr(run, "text"))
        )
        transform = child.properties["pptx"]["transform"]

        assert transform["matrix"] == pytest.approx([0, -2, -2, 0, 180, 180], abs=0.01)
        assert transform["width_pt"] == pytest.approx(72.0)
        assert transform["height_pt"] == pytest.approx(36.0)
        assert transform["rotation"] == pytest.approx(90.0)
        assert child.box.x == pytest.approx(108.0)
        assert child.box.y == pytest.approx(36.0)
        assert child.box.width == pytest.approx(72.0)
        assert child.box.height == pytest.approx(144.0)

    def test_shape_flip_is_preserved_as_affine_matrix(self, tmp_path):
        path = tmp_path / "flipped-shape.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1))
        shape.text = "flipped"
        shape._element.spPr.xfrm.set("flipV", "1")
        presentation.save(path)

        model = read_pptx_model(path)
        paragraph = next(block for block in _paragraphs(model.sections[0].blocks) if block.plain_text == "flipped")
        transform = paragraph.properties["pptx"]["transform"]

        assert transform["matrix"] == pytest.approx([1, 0, 0, -1, 72, 216], abs=0.01)
        assert transform["flip_vertical"] is True

    def test_nested_group_transforms_are_composed(self, tmp_path):
        path = tmp_path / "nested-groups.pptx"
        _build_nested_transformed_group_pptx(path)

        model = read_pptx_model(path)
        paragraph = next(block for block in _paragraphs(model.sections[0].blocks) if block.plain_text == "nested child")
        transform = paragraph.properties["pptx"]["transform"]
        matrix = transform["matrix"]

        assert transform["rotation"] == pytest.approx(45.0)
        assert abs(matrix[1]) > 0.1
        assert abs(matrix[2]) > 0.1
        assert matrix[0] * matrix[3] - matrix[1] * matrix[2] > 0  # Two flips preserve orientation.
        assert paragraph.box.width > 0
        assert paragraph.box.height > 0

    def test_extracts_table(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        tables = [block for block in model.sections[0].blocks if isinstance(block, Table)]
        assert len(tables) == 1
        assert [cell_plain(c) for c in tables[0].rows[0].cells] == ["A1", "B1"]
        assert tables[0].box.x == pytest.approx(72.0)

    def test_extracts_chart_data(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        chart_block = next(p for p in _paragraphs(model.sections[0].blocks) if p.properties["pptx"]["shape"]["kind"] == "chart")
        chart = chart_block.properties["pptx"]["chart"]
        assert chart["categories"] == ["Alpha", "Beta"]
        assert chart["series"][0]["name"] == "Score"
        assert chart["series"][0]["values"] == ["10", "20"]
        assert chart["series"][0]["color"] == "#4472C4"
        assert chart["bar_direction"] == "col"
        assert chart["grouping"] == "clustered"
        assert chart["legend_position"] == "b"
        assert chart["category_axis_title"] == "Categories"
        assert chart["value_axis_title"] == "Score axis"
        axes = chart["axes"]
        assert axes["category"]["position"] == "b"
        assert axes["category"]["tick_label_position"] == "nextTo"
        assert axes["category"]["hidden"] is False
        assert axes["value"]["position"] == "l"
        assert axes["value"]["major_tick_mark"] == "out"

    def test_load_theme_colors_reads_actual_theme(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        presentation = Presentation(str(path))
        colors = pptx_mod._load_theme_colors(presentation.slides[0].part)
        assert colors["accent1"] == "4F81BD"
        assert colors["accent1"] != pptx_mod._THEME_COLORS["accent1"]
        assert colors["accent2"] == "C0504D"

    def test_chart_series_scheme_color_resolves_against_theme(self):
        xml = (
            f'<ser xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f'<c:spPr><a:solidFill><a:schemeClr val="accent2"/></a:solidFill></c:spPr>'
            f"</ser>"
        )
        element = etree.fromstring(xml)
        color = pptx_mod._chart_series_color(element, {"accent2": "C0504D"})
        assert color == "#C0504D"
        assert color != "#" + pptx_mod._THEME_COLORS["accent2"]

    def test_chart_color_keeps_alpha_and_theme_modifiers_in_canonical_value(self):
        element = etree.fromstring(
            f'<c:ser xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            '<c:spPr><a:solidFill><a:schemeClr val="accent2">'
            '<a:shade val="50000"/><a:alpha val="40000"/>'
            "</a:schemeClr></a:solidFill></c:spPr></c:ser>"
        )

        color = pptx_mod._chart_series_color_value(element, {"accent2": "C0504D"})

        assert color is not None
        assert color.to_hex(include_alpha=True) == "#60282666"

    def test_drawingml_theme_color_preserves_and_applies_modifiers(self):
        solid = etree.fromstring(
            f'<a:solidFill xmlns:a="{pptx_mod._A_NS}">'
            '<a:schemeClr val="accent1"><a:tint val="50000"/><a:alpha val="40000"/></a:schemeClr>'
            "</a:solidFill>"
        )

        color, properties = pptx_mod._resolve_color_value(solid, {"accent1": "204060"})

        assert color is not None
        assert color.to_hex(include_alpha=True) == "#90A0B066"
        assert properties["color_theme"] == "accent1"
        assert properties["color_modifiers"] == {"tint": 0.5, "alpha": 0.4}

    def test_read_chart_data_axes_and_theme_accents(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f"<c:ser>"
            f"<c:tx><c:strRef><c:strCache><c:pt><c:v>Score</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache>"
            f"<c:pt><c:v>A</c:v></c:pt><c:pt><c:v>B</c:v></c:pt>"
            f"</c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>10</c:v></c:pt><c:pt><c:v>20</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:barChart>"
            f"<c:valAx>"
            f'<c:scaling><c:autoMin val="0"/><c:autoMax val="0"/><c:max val="100"/></c:scaling>'
            f'<c:numFmt formatCode="0.0%" sourceLinked="0"/>'
            f'<c:tickLblPos val="none"/>'
            f"</c:valAx>"
            f'<c:catAx><c:scaling><c:orientation val="minMax"/></c:scaling>'
            f'<c:tickLblPos val="nextTo"/></c:catAx>'
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: {**pptx_mod._THEME_COLORS, "accent1": "4F81BD"})

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["series"][0]["name"] == "Score"
        assert data["series"][0]["color"] == "#4F81BD"
        value_axis = data["axes"]["value"]
        assert value_axis["num_format"] == "0.0%"
        assert value_axis["num_format_linked"] is False
        assert value_axis["tick_label_position"] == "none"
        assert value_axis["auto_min"] is False
        assert value_axis["auto_max"] is False
        assert value_axis["max"] == 100.0
        assert data["axes"]["category"]["tick_label_position"] == "nextTo"

    def test_read_chart_data_labels(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f"<c:pieChart>"
            f"<c:dLbls>"
            f'<c:showLegendKey val="0"/><c:showVal val="1"/>'
            f'<c:showCatName val="0"/><c:showSerName val="1"/>'
            f'<c:showPercent val="0"/>'
            f'<c:numFmt formatCode="0.0" sourceLinked="0"/>'
            f'<c:dLblPos val="ctr"/>'
            f"</c:dLbls>"
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache><c:pt><c:v>10</c:v></c:pt></c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:pieChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        labels = data["data_labels"]
        assert labels["show_value"] is True
        assert labels["show_percent"] is False
        assert labels["show_series"] is True
        assert labels["show_legend_key"] is False
        assert labels["num_format"] == "0.0"
        assert labels["position"] == "ctr"

    def test_read_chart_bar_gap_and_overlap(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f'<c:varyColors val="1"/><c:gapWidth val="200"/><c:overlap val="-27"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache><c:pt><c:v>10</c:v></c:pt></c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:barChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["gap_width"] == 200.0
        assert data["overlap"] == -27.0
        assert data["vary_colors"] is True

    def test_read_chart_3d_normalizes_chart_type(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:bar3DChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache><c:pt><c:v>10</c:v></c:pt></c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:bar3DChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["chart_type"] == "barChart"
        assert data["chart_3d"] is True
        assert data["series"][0]["chart_type"] == "barChart"

    def test_read_chart_3d_view_ser_axis_wall_floor(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:bar3DChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache><c:pt><c:v>10</c:v></c:pt></c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f'<c:axId val="1"/><c:axId val="2"/><c:axId val="3"/>'
            f"</c:bar3DChart>"
            f'<c:catAx><c:axId val="1"/><c:delete val="0"/><c:axPos val="b"/></c:catAx>'
            f'<c:serAx><c:axId val="2"/><c:delete val="0"/><c:axPos val="b"/></c:serAx>'
            f'<c:valAx><c:axId val="3"/><c:delete val="0"/><c:axPos val="l"/></c:valAx>'
            f"</c:plotArea>"
            f'<c:view3D><c:rotX val="15"/><c:rotY val="20"/><c:rAngAx val="1"/>'
            f'<c:perspective val="30"/><c:depthPercent val="130"/></c:view3D>'
            f'<c:sideWall><c:thickness val="5"/><c:spPr><a:solidFill>'
            f'<a:srgbClr val="C9C9C9"/></a:solidFill></c:spPr></c:sideWall>'
            f'<c:backWall><c:thickness val="3"/></c:backWall>'
            f'<c:floor><c:thickness val="5"/></c:floor>'
            f"</c:chart>"
            f"</c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["chart_3d"] is True
        assert data["chart_3d_type"] == "bar3DChart"
        assert data["view3d"] == {
            "rot_x": 15.0,
            "rot_y": 20.0,
            "right_angle_axes": True,
            "perspective": 30.0,
            "depth_percent": 130.0,
        }
        assert data["side_wall"]["thickness"] == 5.0
        assert data["side_wall"]["fill"] == "#C9C9C9"
        assert data["back_wall"]["thickness"] == 3.0
        assert data["floor"]["thickness"] == 5.0
        assert "series" in data["axes"]
        assert data["axes"]["series"]["ax_id"] == "2"

    def test_read_chart_3d_line_and_cone_types(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:line3DChart><c:grouping val="standard"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache><c:pt><c:v>10</c:v></c:pt></c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:line3DChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["chart_type"] == "lineChart"
        assert data["chart_3d"] is True
        assert data["chart_3d_type"] == "line3DChart"

        cone_xml = chart_xml.replace(
            '<c:line3DChart><c:grouping val="standard"/>',
            '<c:bar3DChart><c:barDir val="col"/><c:grouping val="standard"/><c:shape val="cylinder"/>',
        ).replace("</c:line3DChart>", "</c:bar3DChart>")

        class ConePart:
            blob = cone_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: ConePart())
        cone = pptx_mod._read_chart_data(object(), "rId1")
        assert cone["chart_type"] == "barChart"
        assert cone["chart_3d"] is True
        assert cone["chart_3d_shape"] == "cylinder"

    def test_read_chart_combo_marks_series_by_chart_node(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>Bars</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt><c:pt><c:v>B</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>10</c:v></c:pt><c:pt><c:v>20</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:barChart>"
            f'<c:lineChart><c:grouping val="standard"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>Trend</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>1</c:v></c:pt><c:pt><c:v>3</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:lineChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["chart_type"] == "barChart"
        assert data["combo_types"] == ["barChart", "lineChart"]
        assert data["series"][0]["name"] == "Bars"
        assert data["series"][0]["chart_type"] == "barChart"
        assert data["series"][1]["name"] == "Trend"
        assert data["series"][1]["chart_type"] == "lineChart"
        assert data["categories"] == ["A", "B"]

    def test_read_chart_secondary_axis_marks_series(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f'<c:axId val="1"/><c:axId val="2"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>Bars</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt><c:pt><c:v>B</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>10</c:v></c:pt><c:pt><c:v>20</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:barChart>"
            f'<c:lineChart><c:grouping val="standard"/>'
            f'<c:axId val="1"/><c:axId val="3"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>Growth</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>100</c:v></c:pt><c:pt><c:v>300</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"</c:ser>"
            f"</c:lineChart>"
            f'<c:catAx><c:axId val="1"/></c:catAx>'
            f'<c:valAx><c:axId val="2"/></c:valAx>'
            f'<c:valAx><c:axId val="3"/><c:axPos val="r"/>'
            f"<c:title><c:tx><c:rich><a:p><a:r><a:t>Right</a:t></a:r></a:p></c:rich></c:tx></c:title>"
            f"</c:valAx>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        assert data["axes"]["value"]["ax_id"] == "2"
        assert data["axes"]["secondary_value"]["ax_id"] == "3"
        assert data["axes"]["secondary_value"]["position"] == "r"
        assert data["series"][0].get("axis") is None
        assert data["series"][1]["axis"] == "secondary_value"
        assert data["secondary_value_axis_title"] == "Right"

    def test_read_chart_trendline(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:lineChart><c:grouping val="standard"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>1</c:v></c:pt><c:pt><c:v>3</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"<c:trendline>"
            f'<c:trendlineType val="poly"/>'
            f'<c:order val="2"/>'
            f'<c:dispRSqr val="1"/>'
            f'<c:dispEq val="0"/>'
            f'<c:spPr><a:solidFill><a:srgbClr val="FF8800"/></a:solidFill></c:spPr>'
            f"</c:trendline>"
            f"</c:ser>"
            f"</c:lineChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        trend = data["series"][0]["trendline"]
        assert trend["type"] == "poly"
        assert trend["order"] == 2
        assert trend["show_r_squared"] is True
        assert trend["show_equation"] is False
        assert trend["color"] == "#FF8800"

    def test_read_chart_error_bars(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:lineChart><c:grouping val="standard"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>1</c:v></c:pt><c:pt><c:v>3</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f"<c:errBars>"
            f'<c:errDir val="y"/>'
            f'<c:errBarType val="both"/>'
            f'<c:errValType val="fixedVal"/>'
            f'<c:val val="0.5"/>'
            f"</c:errBars>"
            f"</c:ser>"
            f"</c:lineChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        error = data["series"][0]["error_bars"]
        assert error["direction"] == "y"
        assert error["bar_type"] == "both"
        assert error["value_type"] == "fixedVal"
        assert error["value"] == 0.5

    def test_read_chart_data_points(self, monkeypatch):
        chart_xml = (
            f'<c:chartSpace xmlns:c="{pptx_mod._C_NS}" xmlns:a="{pptx_mod._A_NS}">'
            f"<c:chart><c:plotArea>"
            f'<c:barChart><c:barDir val="col"/><c:grouping val="clustered"/>'
            f"<c:ser><c:tx><c:strRef><c:strCache><c:pt><c:v>S</c:v></c:pt></c:strCache></c:strRef></c:tx>"
            f"<c:cat><c:strRef><c:strCache><c:pt><c:v>A</c:v></c:pt><c:pt><c:v>B</c:v></c:pt></c:strCache></c:strRef></c:cat>"
            f"<c:val><c:numRef><c:numCache>"
            f"<c:pt><c:v>10</c:v></c:pt><c:pt><c:v>20</c:v></c:pt>"
            f"</c:numCache></c:numRef></c:val>"
            f'<c:dPt><c:idx val="0"/>'
            f'<c:spPr><a:solidFill><a:srgbClr val="70AD47"/></a:solidFill></c:spPr>'
            f"</c:dPt>"
            f"</c:ser>"
            f"</c:barChart>"
            f"</c:plotArea></c:chart></c:chartSpace>"
        )

        class FakePart:
            blob = chart_xml.encode()

        monkeypatch.setattr(pptx_mod, "_related_part", lambda _slide, _rid: FakePart())
        monkeypatch.setattr(pptx_mod, "_load_theme_colors", lambda _slide: dict(pptx_mod._THEME_COLORS))

        data = pptx_mod._read_chart_data(object(), "rId1")
        points = data["series"][0]["data_points"]
        assert points[0]["color"] == "#70AD47"
        assert ColorValue.from_dict(points[0]["color_value"]).to_hex() == "#70AD47"

    def test_extracts_image_resource(self, tmp_path):
        path = tmp_path / "rich.pptx"
        image_bytes = _png_bytes(tmp_path)
        _build_rich_pptx(path, image_bytes=image_bytes)

        model = read_pptx_model(path)
        images = [block for block in model.sections[0].blocks if isinstance(block, Image)]
        assert len(images) == 1
        resource = model.resources[images[0].resource_id]
        assert resource.kind == ResourceKind.RASTER_IMAGE
        assert resource.data == image_bytes
        assert model.validate() == []

    def test_extracts_omml_formula(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        paragraphs = _paragraphs(model.sections[0].blocks)
        formulas = [item for p in paragraphs for item in p.content if isinstance(item, Formula)]
        assert len(formulas) == 1
        assert formulas[0].format == "omml"
        assert "oMath" in formulas[0].value
        assert "x" in formulas[0].fallback_text or "y" in formulas[0].fallback_text

    def test_background_fill_and_notes(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        model = read_pptx_model(path)
        section = model.sections[0]
        assert section.properties["background_fill"] == "#F2F2F2"
        assert section.properties["notes"] == "Speaker notes for test."
        assert section.page.margin_top.pt == 0
        assert section.page.margin_right.pt == 0
        assert section.page.margin_bottom.pt == 0
        assert section.page.margin_left.pt == 0

    def test_reads_plain_text(self, tmp_path):
        path = tmp_path / "rich.pptx"
        _build_rich_pptx(path)

        text = read_pptx(path)

        assert text.pages == 1
        assert text.source_format.value == "pptx"
        assert "Mixed formatting" in text.plain
        cells = [cell for table in text.tables for row in table.rows for cell in row]
        assert "A1" in cells and "B2" in cells

    def test_imports_corpus_presentation(self):
        model = read_pptx_model(CORPUS_PPTX)

        assert len(model.sections) == 2
        assert model.validate() == []
        tables = [block for block in model.sections[0].blocks if isinstance(block, Table)]
        assert len(tables) == 1
        chart_blocks = [p for p in _paragraphs(model.sections[0].blocks) if p.properties["pptx"]["shape"]["kind"] == "chart"]
        assert len(chart_blocks) == 1
        assert chart_blocks[0].properties["pptx"]["chart"]["chart_type"] == "barChart"

    def test_missing_file_raises(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            read_pptx_model(tmp_path / "missing.pptx")

    def test_placeholder_inherits_layout_geometry_and_style(self, tmp_path):
        path = tmp_path / "placeholders.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.placeholders[0].text_frame.text = "Slide title"
        slide.placeholders[1].text_frame.text = "Body text"
        presentation.save(path)

        model = read_pptx_model(path)

        blocks = _paragraphs(model.sections[0].blocks)
        title = next(p for p in blocks if any(run.text == "Slide title" for run in p.content))
        body = next(p for p in blocks if any(run.text == "Body text" for run in p.content))
        title_run = next(r for r in title.content if r.text == "Slide title")
        body_run = next(r for r in body.content if r.text == "Body text")
        assert title.box.x == pytest.approx(36.0)
        assert title.box.width == pytest.approx(648.0)
        assert body.box.y == pytest.approx(126.0)
        assert title_run.style.font_size.pt == pytest.approx(44.0)
        assert body_run.style.font_size.pt == pytest.approx(32.0)
        assert isinstance(title_run.style.color, ColorValue)
        assert title_run.style.color.to_hex() == "#000000"
        assert title.alignment == "center"


def cell_plain(cell) -> str:
    return "".join(run.text for block in cell.blocks for run in block.content if hasattr(run, "text"))
