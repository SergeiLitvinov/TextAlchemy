"""Интеграционные тесты DOCX → DocumentModel."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from PIL import Image as PillowImage

from textalchemy.core.document_model import Formula, Image, Paragraph, ResourceKind, Table
from textalchemy.formats.docx import read_docx_model


def test_read_docx_model_preserves_order_styles_tables_and_metadata(tmp_path):
    path = tmp_path / "structured.docx"
    source = Document()
    source.core_properties.title = "Structured document"
    paragraph = source.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Formatted")
    run.bold = True
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    source.add_paragraph("After table")
    source.save(path)

    model = read_docx_model(path)

    assert model.metadata["title"] == "Structured document"
    assert model.source_format == "docx"
    assert [type(block) for block in model.sections[0].blocks] == [Paragraph, Table, Paragraph]
    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Paragraph)
    assert imported.alignment == "center"
    assert imported.content[0].style.bold is True
    assert imported.content[0].style.italic is True
    assert imported.content[0].style.font_family == "Arial"
    assert imported.content[0].style.font_size.pt == 14


def test_read_docx_model_extracts_image_formula_header_and_page_geometry(tmp_path):
    path = tmp_path / "media.docx"
    image_path = tmp_path / "pixel.png"
    PillowImage.new("RGB", (8, 6), "red").save(image_path)

    source = Document()
    source.sections[0].header.paragraphs[0].text = "Header text"
    source.sections[0].page_width = Inches(9)
    paragraph = source.add_paragraph("Before ")
    picture = paragraph.add_run()
    picture.add_picture(str(image_path), width=Inches(1))
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "x+1"
    math_run.append(math_text)
    math.append(math_run)
    paragraph._p.append(math)
    source.save(path)

    model = read_docx_model(path)

    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Paragraph)
    assert any(isinstance(item, Image) for item in imported.content)
    assert any(isinstance(item, Formula) and item.fallback_text == "x+1" for item in imported.content)
    assert len(model.resources) == 1
    resource = next(iter(model.resources.values()))
    assert resource.kind is ResourceKind.RASTER_IMAGE
    assert resource.media_type == "image/png"
    assert resource.data
    assert model.sections[0].headers[0].plain_text == "Header text"
    assert model.sections[0].page.width.pt == 648
    assert model.validate() == []


def test_read_docx_model_preserves_horizontal_cell_merge(tmp_path):
    path = tmp_path / "merged.docx"
    source = Document()
    table = source.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "merged"
    table.cell(0, 2).text = "single"
    source.save(path)

    model = read_docx_model(path)

    imported = model.sections[0].blocks[0]
    assert isinstance(imported, Table)
    assert len(imported.rows[0].cells) == 2
    assert imported.rows[0].cells[0].column_span == 2
    assert imported.rows[0].cells[0].blocks[0].plain_text == "merged"
