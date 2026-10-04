from pathlib import Path
from zipfile import ZipFile

import pytest

pytest.importorskip("bs4")
pytest.importorskip("tinycss2")

from textalchemy.convert.executor import ConversionExecutor, ConversionRequest
from textalchemy.core.document_codec import document_from_json, document_to_json, load_document, save_document
from textalchemy.core.document_model import Formula, Table
from textalchemy.core.quality_policy import QualityPolicy
from textalchemy.core.types import DocFormat, Document
from textalchemy.formats.html import read_html, read_html_model
from textalchemy.pipeline.extract import extract_html_model

CORPUS = Path(__file__).parents[1] / "corpus/scientific-html.html"


def test_scientific_structure_css_and_order():
    model = read_html_model(CORPUS)
    assert not model.validate()
    assert model.metadata["html"]["warnings"] == []
    blocks = model.sections[0].blocks
    assert [block.plain_text for block in blocks[:7]] == [
        "Исследование",
        "До текста после ссылка.",
        "Важное",
        "Первый",
        "Вложенный",
        "Второй",
        "Измерения",
    ]
    assert blocks[0].style_id == "Heading1"
    assert blocks[1].content[0].style.color.to_hex() == "#778899"
    assert next(run for run in blocks[1].content if run.text == "текста").style.bold is False
    assert blocks[2].content[0].style.color.to_hex().lower() == "#aa0000"
    assert [block.properties["list_level"] for block in blocks[3:6]] == [0, 1, 0]
    assert blocks[3].properties["list_start"] == 3
    table = blocks[7]
    assert isinstance(table, Table)
    assert [p.plain_text for p in table.rows[1].cells[1].blocks] == ["12", "± 1"]
    assert table.rows[0].cells[0].blocks[0].content[0].style.bold
    assert any(isinstance(item, Formula) for item in blocks[8].content)
    assert len(model.resources) == 2
    assert 'viewBox="0 0 10 10"' in next(r.data.decode() for r in model.resources.values() if r.media_type == "image/svg+xml")
    assert blocks[-2].plain_text == " a  b\n c"
    assert read_html(CORPUS).source_format is DocFormat.HTML
    assert extract_html_model(doc=Document(CORPUS, DocFormat.HTML, 0, "")).source_format == "html"


def test_json_mutations_reach_html_and_docx(tmp_path):
    pytest.importorskip("docx")
    from bs4 import BeautifulSoup
    from docx import Document as WordDocument

    model = document_from_json(document_to_json(read_html_model(CORPUS)))
    blocks = model.sections[0].blocks
    blocks[0].content[0].text = "Изменённый заголовок"
    blocks[7].rows[1].cells[1].blocks[0].content[0].text = "24"
    formula = next(item for item in blocks[8].content if isinstance(item, Formula))
    formula.value = formula.value.replace(">2<", ">4<")
    source = tmp_path / "edited.json"
    save_document(model, source)
    executor = ConversionExecutor()
    html, docx = tmp_path / "result.html", tmp_path / "result.docx"
    assert executor.execute(ConversionRequest(source, html, DocFormat.MODEL, DocFormat.HTML)).success
    assert executor.execute(ConversionRequest(source, docx, DocFormat.MODEL, DocFormat.DOCX)).success
    soup = BeautifulSoup(html.read_text(encoding="utf-8"), "html.parser")
    assert soup.h1.get_text() == "Изменённый заголовок"
    assert soup.ol["start"] == "3" and len(soup.ol.find_all("li", recursive=False)) == 2
    assert soup.ol.li.ul.li.get_text() == "Вложенный"
    assert "24" in soup.table.get_text()
    assert soup.math is not None and soup.math.mfrac is not None
    assert "data:image/png;base64," in str(soup)
    word = WordDocument(docx)
    assert word.paragraphs[0].text == "Изменённый заголовок"
    assert word.tables[0].cell(1, 1).paragraphs[0].text == "24"
    with ZipFile(docx) as package:
        xml = package.read("word/document.xml").decode()
        assert "oMath" in xml and ">4<" in xml
        assert "numPr" in xml
        from textalchemy.convert.docx_html_links import bookmark_name

        assert f'w:name="{bookmark_name("title")}"' in xml
        assert f'w:anchor="{bookmark_name("title")}"' in xml
        assert 'w:val="3"' in package.read("word/numbering.xml").decode()


@pytest.mark.parametrize(
    "css,expected",
    [
        ("p {color:red} .x {color:blue}", "#0000FF"),
        ("#a {color:red} .x {color:blue}", "#FF0000"),
        (".x {color:red} .x {color:blue}", "#0000FF"),
        (".x {color:red!important}", "#FF0000"),
    ],
)
def test_cascade_precedence(tmp_path, css, expected):
    path = tmp_path / "style.html"
    path.write_text(f'<style>{css}</style><p id="a" class="x">Text</p>', encoding="utf-8")
    assert read_html_model(path).sections[0].blocks[0].content[0].style.color.to_hex().upper() == expected


def test_nested_relative_sizes_and_whitespace(tmp_path):
    path = tmp_path / "size.html"
    path.write_text(
        '<html style="font-size:10pt"><body><div style="font-size:2em"><p style="font-size:50%">A\n <b>B</b> C</p>'
        '<p style="font-size:2rem">D</p></div></body></html>'
    )
    a, b = read_html_model(path).sections[0].blocks
    assert a.plain_text == "A B C" and a.content[0].style.font_size.pt == 10
    assert b.content[0].style.font_size.pt == 20


def test_loss_report_and_strict_budget_preserve_previous_output(tmp_path):
    source, target = tmp_path / "bad.html", tmp_path / "result.json"
    source.write_text(
        "<style>@page {size:A4} p:hover {color:red} p {display:grid}</style><script>alert(1)</script>"
        '<p>Keep<img alt="Missing" src="https://example.org/image.png"></p>'
    )
    report = ConversionExecutor().execute(ConversionRequest(source, target, DocFormat.HTML, DocFormat.MODEL))
    assert report.success and not report.lossless
    assert {item.feature for item in report.issues} >= {"html-css", "html-content", "html-resource"}
    assert load_document(target).sections[0].blocks[0].plain_text == "KeepMissing"
    previous = target.read_bytes()
    request = ConversionRequest(source, target, DocFormat.HTML, DocFormat.MODEL, quality_policy=QualityPolicy(max_loss_issues=0))
    assert not ConversionExecutor().execute(request).success
    assert target.read_bytes() == previous


def test_external_assets_opt_in_and_svg_sanitizing(tmp_path):
    path = tmp_path / "input.html"
    (tmp_path / "asset.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script>'
        '<rect width="3" height="3" onclick="run()"/><image href="https://example.org/a"/></svg>'
    )
    path.write_text('<p><img src="asset.svg"></p>')
    assert not read_html_model(path).resources
    model = read_html_model(path, resource_root=tmp_path)
    data = next(iter(model.resources.values())).data.decode()
    assert "script" not in data and "onclick" not in data and "https://example.org/a" not in data
    assert model.metadata["html"]["warnings"]
    path.write_text('<p><img src="../outside.png"></p>')
    assert not read_html_model(path, resource_root=tmp_path).resources


def test_cli_pipeline_and_web_catalog(tmp_path):
    from textalchemy.__main__ import main
    from textalchemy.web.services.conversion_catalog import available_conversions

    output = tmp_path / "document.json"
    assert main(["convert-file", str(CORPUS), str(output), "--json"]) == 0
    assert load_document(output).source_format == "html"
    catalog = available_conversions(ConversionExecutor())
    source = next(item for item in catalog["sources"] if item["format"] == "html")
    assert {target["format"] for target in source["targets"]} >= {"docx", "model", "txt"}


def test_optional_dependencies_unavailable(monkeypatch):
    import builtins

    original = builtins.__import__

    def without_html(name, *args, **kwargs):
        if name == "bs4":
            raise ImportError("not installed")
        return original(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_html)
    assert read_html(CORPUS).warnings == ["Install opendoc-formats[html]"]
    executor = ConversionExecutor(requirement_checker=lambda name: name not in {"beautifulsoup4", "tinycss2"})
    assert executor.plan(DocFormat.HTML, DocFormat.MODEL) is None


def test_nested_underline_and_inline_preserved_spaces(tmp_path):
    path = tmp_path / "inline.html"
    path.write_text('<p><u><b>Decorated</b></u></p><p><span style="white-space:pre">  x  </span></p>')
    first, second = read_html_model(path).sections[0].blocks
    assert first.content[0].style.underline and first.content[0].style.bold
    assert second.plain_text == "  x  "


def test_unsupported_active_attributes_and_background_are_reported(tmp_path):
    path = tmp_path / "attributes.html"
    path.write_text('<p onclick="run()" style="background-color:red;font-weight:BOLD">Visible</p>')
    model = read_html_model(path)
    assert model.sections[0].blocks[0].content[0].style.bold
    assert {item["feature"] for item in model.metadata["html"]["warnings"]} == {"html-css", "html-content"}
