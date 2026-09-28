import json
from pathlib import Path

import pytest

from textalchemy import __version__
from textalchemy.__main__ import main


def test_cli_no_args(capsys):
    ret = main([])
    assert ret == 1
    captured = capsys.readouterr()
    assert "usage" in captured.out or "usage" in captured.err


def test_cli_version(capsys):
    with pytest.raises(SystemExit) as exc:
        main(["--version"])
    assert exc.value.code == 0
    assert __version__ in capsys.readouterr().out


@pytest.mark.parametrize("config_exists", [False, True])
@pytest.mark.parametrize("command", ["init", "match"])
def test_cli_config_rejected_before_side_effects(tmp_path, monkeypatch, capsys, config_exists, command):
    monkeypatch.chdir(tmp_path)
    config = tmp_path / "settings.json"
    if config_exists:
        config.write_text(json.dumps({"matching": {"threshold": 0.99}, "output": "from-config"}), encoding="utf-8")
    output = tmp_path / "result"
    if command == "init":
        arguments = ["init", "-o", str(output)]
    else:
        source = tmp_path / "source"
        source.mkdir()
        (source / "Ferroresonance_in_Power_Grids_2020.txt").write_text(
            "Иванов И.И. Ferroresonance in Power Grids 2020", encoding="utf-8",
        )
        bib = tmp_path / "bibliography.txt"
        bib.write_text("1. Иванов И.И. Ferroresonance in Power Grids. 2020.\n", encoding="utf-8")
        arguments = ["match", "-s", str(source), "-b", str(bib), "-o", str(output), "-t", "0.30", "--json"]
    assert main(["--config", str(config), *arguments]) == 1
    captured = capsys.readouterr()
    assert "--config" in captured.err
    assert "не поддерживается" in captured.err
    assert not output.exists()
    assert not (tmp_path / "from-config").exists()
    assert not (tmp_path / "matching_report.json").exists()
    assert captured.out == ""


def test_cli_short_config_without_command_is_rejected(capsys):
    assert main(["-c", "missing.json"]) == 1
    assert "не поддерживается" in capsys.readouterr().err


def test_cli_help_extract():
    with pytest.raises(SystemExit):
        main(["extract", "--help"])


def test_cli_help_convert():
    with pytest.raises(SystemExit):
        main(["convert", "--help"])


def test_cli_help_match():
    with pytest.raises(SystemExit):
        main(["match", "--help"])


def test_cli_help_gost():
    with pytest.raises(SystemExit):
        main(["gost", "--help"])


def test_cli_help_stats():
    with pytest.raises(SystemExit):
        main(["stats", "--help"])


def test_cli_help_export():
    with pytest.raises(SystemExit):
        main(["export", "--help"])


def test_cli_help_bibtex():
    with pytest.raises(SystemExit):
        main(["bibtex", "--help"])


def test_cli_help_init():
    with pytest.raises(SystemExit):
        main(["init", "--help"])


def test_cli_help_web():
    with pytest.raises(SystemExit):
        main(["web", "--help"])


def test_cli_help_generate():
    with pytest.raises(SystemExit):
        main(["generate", "--help"])


def test_cli_help_inspect():
    with pytest.raises(SystemExit):
        main(["inspect", "--help"])


def test_cli_help_recognize():
    with pytest.raises(SystemExit):
        main(["recognize", "--help"])


# ── init ───────────────────────────────────────────


def test_cli_init(tmp_path, capsys):
    out = tmp_path / "test_config.json"
    ret = main(["init", "-o", str(out)])
    assert ret == 0
    assert out.exists()
    cfg = json.loads(out.read_text(encoding="utf-8"))
    assert "mode" in cfg
    assert "CLI не применяет" in capsys.readouterr().out


# ── extract ────────────────────────────────────────


def test_cli_extract_no_input(capsys):
    ret = main(["extract", "nonexistent.docx"])
    assert ret != 0


# ── convert ────────────────────────────────────────


def test_cli_convert_no_input(capsys):
    ret = main(["convert", "-i", "nonexistent_dir"])
    assert ret != 0


def test_cli_convert_dry_run(tmp_path):
    src = tmp_path / "pdfs"
    src.mkdir()
    ret = main(["convert", "-i", str(src), "--dry-run"])
    assert ret == 0


def test_cli_convert_dry_run_json(tmp_path):
    src = tmp_path / "pdfs"
    src.mkdir()
    ret = main(["convert", "-i", str(src), "--dry-run", "--json"])
    assert ret == 0


# ── stats ──────────────────────────────────────────


def test_cli_stats(capsys):
    ret = main(["stats"])
    assert ret == 0


def test_cli_stats_json(capsys):
    ret = main(["stats", "--json"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert "total" in data
    assert "matched" in data
    assert "unmatched" in data
    assert "bibliography" in data


def test_cli_stats_with_bibliography(tmp_path):
    bib = tmp_path / "bib.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    ret = main(["stats", "-b", str(bib)])
    assert ret == 0


# ── export ─────────────────────────────────────────


def test_cli_export_json(tmp_path):
    bib = tmp_path / "input.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    out = tmp_path / "out.json"
    ret = main(["export", "-i", str(bib), "-o", str(out), "-f", "json"])
    assert ret == 0
    assert out.exists()


def test_cli_export_markdown(tmp_path):
    bib = tmp_path / "input.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    out = tmp_path / "out.md"
    ret = main(["export", "-i", str(bib), "-o", str(out), "-f", "markdown"])
    assert ret == 0
    assert out.exists()


def test_cli_export_gost(tmp_path):
    bib = tmp_path / "input.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    out = tmp_path / "out.txt"
    ret = main(["export", "-i", str(bib), "-o", str(out), "-f", "gost"])
    assert ret == 0
    assert out.exists()


# ── gost ───────────────────────────────────────────


def test_cli_gost(tmp_path):
    bib = tmp_path / "input.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    out = tmp_path / "out.txt"
    ret = main(["gost", "-i", str(bib), "-o", str(out)])
    assert ret == 0
    assert out.exists()


def test_cli_gost_json(tmp_path, capsys):
    bib = tmp_path / "input.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    out = tmp_path / "out.txt"
    ret = main(["gost", "-i", str(bib), "-o", str(out), "--json"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert data["entries"] == 1
    assert data["output"] == str(out)


# ── match ──────────────────────────────────────────


def test_cli_match_dry_run(tmp_path):
    src = tmp_path / "literature_files"
    src.mkdir()
    bib = tmp_path / "bibliography.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    ret = main(["match", "-s", str(src), "-b", str(bib), "--dry-run"])
    assert ret == 0


def test_cli_match_no_bibliography(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "literature_files"
    src.mkdir()
    ret = main(["match", "-s", str(src)])
    assert ret == 1


def test_cli_match_json_report(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    src = tmp_path / "literature_files"
    src.mkdir()
    bib = tmp_path / "bibliography.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    ret = main(["match", "-s", str(src), "-b", str(bib), "--json"])
    assert ret == 0
    report = json.loads((tmp_path / "matching_report.json").read_text(encoding="utf-8"))
    assert report == {"matched": [], "unmatched": []}


@pytest.mark.parametrize("relative_output", [False, True])
@pytest.mark.parametrize("copy_fails", [False, True])
def test_cli_match_json_reports_actual_copy(tmp_path, monkeypatch, relative_output, copy_fails):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "источники"
    source.mkdir()
    original = source / "Ferroresonance_in_Power_Grids_2020.TXT"
    original.write_text("Иванов И.И. Ferroresonance in Power Grids 2020", encoding="utf-8")
    (source / "unrelated.txt").write_text("completely unrelated", encoding="utf-8")
    bib = tmp_path / "bibliography.txt"
    bib.write_text("1. Иванов И.И. Ferroresonance in Power Grids. Электричество. 2020. с. 10-20.\n", encoding="utf-8")
    output = Path("результат") if relative_output else tmp_path / "результат"
    if copy_fails:
        import shutil

        def fail_copy(source, target):
            raise OSError("Copy unavailable")

        monkeypatch.setattr(shutil, "copy2", fail_copy)

    assert main(["match", "-s", str(source), "-b", str(bib), "-o", str(output), "--json"]) == 0
    report = json.loads((tmp_path / "matching_report.json").read_text(encoding="utf-8"))
    assert report["unmatched"] == ["unrelated.txt"]
    assert len(report["matched"]) == 1
    entry = report["matched"][0]
    assert entry["original"] == original.name
    assert entry["score"] >= 0.3
    assert entry["signals"]
    assert original.exists()
    copied = list(output.iterdir())
    if copy_fails:
        assert entry["new"] is None
        assert copied == []
    else:
        assert len(copied) == 1
        assert entry["new"] == str(copied[0])
        assert copied[0].name != original.name
        assert copied[0].suffix == ".txt"
        assert copied[0].read_bytes() == original.read_bytes()


def test_cli_match_no_files(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    bib = tmp_path / "bibliography.txt"
    bib.write_text("1. Author A. Title.", encoding="utf-8")
    ret = main(["match", "-s", str(tmp_path), "-b", str(bib)])
    assert ret == 0


# ── generate ───────────────────────────────────────


def test_cli_generate_list(capsys):
    ret = main(["generate", "--list"])
    assert ret == 0


def test_cli_generate_list_json(capsys):
    ret = main(["generate", "--list", "--json"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert isinstance(data, list)


def test_cli_generate_docx_from_data_and_schema(tmp_path, capsys):
    from docx import Document

    template = tmp_path / "report.docx"
    output = tmp_path / "result.docx"
    data_path = tmp_path / "data.json"
    schema_path = tmp_path / "schema.json"
    document = Document()
    document.add_heading("{{ title }}", level=1)
    document.add_paragraph("{% for item in items %}")
    document.add_paragraph("Item: {{ item }}")
    document.add_paragraph("{% endfor %}")
    document.save(template)
    data_path.write_text('{"title": "Report", "items": ["A", "B"]}', encoding="utf-8")
    schema_path.write_text(
        '{"fields": {"title": {"type": "string"}, "items": {"type": "array"}}, "allow_extra": false}',
        encoding="utf-8",
    )

    ret = main(
        [
            "generate",
            str(template),
            str(output),
            "--data",
            str(data_path),
            "--schema",
            str(schema_path),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    assert ret == 0
    assert payload["success"] is True
    assert payload["lossless"] is True
    assert [paragraph.text for paragraph in Document(output).paragraphs] == ["Report", "Item: A", "Item: B"]


def test_cli_generate_validation_error_is_json(tmp_path, capsys):
    from docx import Document

    template = tmp_path / "report.docx"
    output = tmp_path / "result.docx"
    data_path = tmp_path / "data.json"
    schema_path = tmp_path / "schema.json"
    document = Document()
    document.add_paragraph("{{ title }}")
    document.save(template)
    data_path.write_text('{"title": 42}', encoding="utf-8")
    schema_path.write_text('{"fields": {"title": {"type": "string"}}}', encoding="utf-8")

    ret = main(
        [
            "generate",
            str(template),
            str(output),
            "--data",
            str(data_path),
            "--schema",
            str(schema_path),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    assert ret == 1
    assert payload["success"] is False
    assert "expected string" in payload["error"]
    assert not output.exists()


def test_cli_generate_self_contained_html_with_typed_image(tmp_path, capsys):
    from docx import Document
    from PIL import Image as PillowImage

    template = tmp_path / "report.docx"
    output = tmp_path / "result.html"
    image = tmp_path / "diagram.png"
    data_path = tmp_path / "data.json"
    schema_path = tmp_path / "schema.json"
    document = Document()
    document.add_paragraph("{{ diagram }}")
    document.save(template)
    PillowImage.new("RGB", (5, 5), "blue").save(image)
    data_path.write_text(json.dumps({"diagram": {"source": str(image), "alt_text": "Diagram"}}), encoding="utf-8")
    schema_path.write_text('{"fields": {"diagram": {"type": "image"}}, "allow_extra": false}', encoding="utf-8")

    ret = main(
        [
            "generate",
            str(template),
            str(output),
            "--data",
            str(data_path),
            "--schema",
            str(schema_path),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    html = output.read_text(encoding="utf-8")
    assert ret == 0
    assert payload["success"] is True
    assert payload["lossless"] is True
    assert payload["metrics"]["embedded_resources"] == 1
    assert "data:image/png;base64," in html
    assert 'alt="Diagram"' in html


def test_cli_generate_pdf_from_same_template_data_and_schema(tmp_path, capsys):
    import fitz
    from docx import Document

    template = tmp_path / "report.docx"
    output = tmp_path / "result.pdf"
    data_path = tmp_path / "data.json"
    schema_path = tmp_path / "schema.json"
    document = Document()
    document.add_heading("{{ title }}", level=1)
    document.add_paragraph("{% for item in items %}")
    document.add_paragraph("Item: {{ item }}")
    document.add_paragraph("{% endfor %}")
    document.save(template)
    data_path.write_text('{"title": "PDF Report", "items": ["Alpha", "Beta"]}', encoding="utf-8")
    schema_path.write_text(
        '{"fields": {"title": {"type": "string"}, "items": {"type": "array"}}, "allow_extra": false}',
        encoding="utf-8",
    )

    ret = main(
        [
            "generate",
            str(template),
            str(output),
            "--data",
            str(data_path),
            "--schema",
            str(schema_path),
            "--json",
        ]
    )

    payload = json.loads(capsys.readouterr().out)
    assert ret == 0
    assert payload["success"] is True
    assert payload["lossless"] is True
    with fitz.open(output) as pdf:
        text = "".join(page.get_text() for page in pdf)
    assert "PDF Report" in text
    assert "Item: Alpha" in text
    assert "Item: Beta" in text


# ── inspect ────────────────────────────────────────


def test_cli_inspect_docx_json_and_report_file(tmp_path, capsys):
    from docx import Document

    source = tmp_path / "source.docx"
    output = tmp_path / "inspection.json"
    document = Document()
    document.add_heading("Inspection", level=1)
    document.add_paragraph("Body text")
    document.save(source)

    ret = main(["inspect", str(source), "--json", "--output", str(output)])

    payload = json.loads(capsys.readouterr().out)
    saved = json.loads(output.read_text(encoding="utf-8"))
    assert ret == 0
    assert payload["valid"] is True
    assert payload["source_format"] == "docx"
    assert payload["metrics"]["paragraphs"] >= 2
    assert saved == payload


def test_cli_plan_conversion_route_json(capsys):
    ret = main(["plan", "docx", "pdf", "--mode", "faithful", "--feature", "page_geometry", "--json"])

    payload = json.loads(capsys.readouterr().out)
    assert ret == 0
    assert payload["success"] is True
    assert [step["id"] for step in payload["steps"]] == ["docx.model", "model.pdf"]
    assert payload["feature_support"]["page_geometry"] == "visual"


def test_cli_plan_reports_missing_route(capsys):
    ret = main(["plan", "latex", "pptx", "--json"])

    payload = json.loads(capsys.readouterr().out)
    assert ret == 1
    assert payload["success"] is False
    assert payload["error"] == "conversion route not found"


def test_cli_convert_file_docx_to_model(tmp_path, capsys):
    from docx import Document

    source = tmp_path / "source.docx"
    output = tmp_path / "model.json"
    document = Document()
    document.add_paragraph("Executor CLI")
    document.save(source)

    ret = main(["convert-file", str(source), str(output), "--json"])

    payload = json.loads(capsys.readouterr().out)
    assert ret == 0
    assert payload["success"] is True
    assert payload["metrics"]["executed_steps"] == ["docx.model"]
    assert output.is_file()


def test_cli_inspect_missing_file_is_machine_readable(tmp_path, capsys):
    ret = main(["inspect", str(tmp_path / "missing.docx"), "--json"])

    payload = json.loads(capsys.readouterr().out)
    assert ret == 1
    assert payload["valid"] is False
    assert "missing.docx" in payload["error"]


def test_cli_inspect_compare_reports_structural_loss(tmp_path, capsys):
    from textalchemy.core.document_codec import save_document
    from textalchemy.core.document_model import DocumentModel, Paragraph, Section, Table, TableRow, TextRun

    source = tmp_path / "source.json"
    target = tmp_path / "target.json"
    save_document(
        DocumentModel(
            sections=[
                Section(
                    blocks=[
                        Paragraph(content=[TextRun("Source text")]),
                        Table(rows=[TableRow()]),
                    ]
                )
            ]
        ),
        source,
    )
    save_document(DocumentModel(sections=[Section(blocks=[Paragraph(content=[TextRun("Short")])])]), target)

    ret = main(["inspect", str(source), "--compare", str(target), "--json", "--strict"])

    payload = json.loads(capsys.readouterr().out)
    assert ret == 1
    assert payload["comparison"]["has_losses"] is True
    assert payload["comparison"]["retention"]["tables"]["ratio"] == 0
    assert payload["comparison"]["geometry_summary"]["max_dimension_error_pt"] == 0
    assert payload["comparison"]["resource_comparison"]["exact_hash_retention_ratio"] == 1
    assert payload["comparison"]["font_comparison"]["exact_run_retention_ratio"] == 1


# ── recognize ──────────────────────────────────────


def test_cli_recognize_no_file(capsys):
    ret = main(["recognize", "nonexistent.pdf"])
    assert ret == 1


def test_cli_recognize_json(capsys):
    ret = main(["recognize", "nonexistent.pdf", "--json"])
    assert ret == 1
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert "error" in data


def test_cli_recognize_with_output(tmp_path):
    out = tmp_path / "out.txt"
    ret = main(["recognize", "nonexistent.pdf", "--output", str(out)])
    assert ret == 1


def test_cli_recognize_pdf_fast_uses_text_layer(tmp_path, capsys):
    """fast scenario reads the PDF text layer without any OCR backend."""
    import fitz

    pdf = tmp_path / "layer.pdf"
    doc = fitz.open()
    page = doc.new_page(width=300, height=400)
    page.insert_text((36, 48), "Fast CLI text layer", fontname="helv", fontsize=12)
    doc.save(str(pdf))
    doc.close()

    ret = main(["recognize", str(pdf), "--scenario", "fast"])
    assert ret == 0
    assert "Fast CLI text layer" in capsys.readouterr().out


def test_cli_recognize_pdf_scan_requires_backend(tmp_path, capsys):
    """scan scenario without an OCR backend exits with an error."""
    import fitz

    pdf = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page(width=300, height=400)
    doc.save(str(pdf))
    doc.close()

    ret = main(["recognize", str(pdf), "--scenario", "scan"])
    captured = capsys.readouterr()
    assert ret == 1
    assert "OCR" in captured.err or "OCR" in captured.out


# ── bibtex ─────────────────────────────────────────


def test_cli_bibtex(tmp_path):
    src = tmp_path / "literature_files"
    src.mkdir()
    out = tmp_path / "out.bib"
    ret = main(["bibtex", "-s", str(src), "-o", str(out)])
    assert ret == 0
    assert out.exists()


def test_cli_bibtex_json(tmp_path, capsys):
    src = tmp_path / "literature_files"
    src.mkdir()
    out = tmp_path / "out.bib"
    ret = main(["bibtex", "-s", str(src), "-o", str(out), "--json"])
    assert ret == 0
    captured = capsys.readouterr()
    data = json.loads(captured.out)
    assert "source" in data
    assert "output" in data
    assert "entries" in data


# ── unknown command ────────────────────────────────


def test_cli_unknown_command():
    with pytest.raises(SystemExit) as exc:
        main(["nonexistent"])
    assert exc.value.code == 2


# ── completion ──────────────────────────────────────


def test_cli_completion_bash():
    ret = main(["completion", "bash"])
    assert ret == 0


def test_cli_completion_zsh():
    ret = main(["completion", "zsh"])
    assert ret == 0


def test_cli_completion_fish():
    ret = main(["completion", "fish"])
    assert ret == 0


def test_cli_completion_invalid_shell():
    with pytest.raises(SystemExit):
        main(["completion", "invalid"])
