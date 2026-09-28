"""M4 acceptance on real DOCX/ZIP output and uploaded pipeline inputs."""
import importlib
import io
import json
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from textalchemy.web.main import app
from textalchemy.web.services.template_variables import custom_path
from textalchemy.web.tasks import TaskStore


@pytest.fixture
def m4_client(tmp_path, monkeypatch):
    web = importlib.import_module('textalchemy.web.app')
    store = TaskStore(tmp_path / 'tasks')
    monkeypatch.setattr(web, 'data_dir', tmp_path / 'data')
    monkeypatch.setattr(web, 'tasks_store', store)
    monkeypatch.setattr('textalchemy.web.routes.convert.tasks_store', store)
    with TestClient(app) as client:
        yield client, store


def sample_bytes():
    doc = Document()
    doc.add_heading('Служебная записка', level=1)
    para = doc.add_paragraph('Руководитель: ')
    para.add_run('Иванов ').bold = True
    para.add_run('Иван Иванович').italic = True
    doc.add_paragraph('Исполнитель: Иванов Иван Иванович')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = 'Автомобиль: Лада'
    doc.sections[0].footer.paragraphs[0].text = 'Подпись: Иванов Иван Иванович'
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def import_sample(client, content=None):
    response = client.post('/api/generate/import', files={'file': ('sample.docx', content or sample_bytes())})
    assert response.status_code == 200, response.text
    return response.json()['name']


def test_sample_linked_occurrences_preserve_package(m4_client):
    client, _ = m4_client
    original = sample_bytes()
    name = import_sample(client, original)
    query = 'Иванов Иван Иванович'
    found = client.get(f'/api/generate/templates/{name}/source', params={'query': query}).json()
    assert len(found['occurrences']) == 3
    chosen = [item for item in found['occurrences'] if 'Исполнитель' not in item['text']]
    payload = {'query': query, 'label': 'ФИО руководителя', 'selected': chosen, 'expected': found['revision']}
    response = client.post(f'/api/generate/templates/{name}/source', json=payload)
    assert response.status_code == 200, response.text
    edited = response.json()['name']
    assert custom_path(name).read_bytes() == original
    schema = client.get(f'/api/generate/templates/{edited}/schema').json()['schema']
    key = schema['fields'][0]['name']
    response = client.post('/api/generate', data={'template': edited, 'params': json.dumps({key: 'Петров Пётр Петрович'})})
    result = Document(io.BytesIO(response.content))
    assert result.paragraphs[1].text == 'Руководитель: Петров Пётр Петрович'
    assert result.paragraphs[2].text == 'Исполнитель: ' + query
    assert result.paragraphs[1].runs[1].bold
    assert result.sections[0].footer.paragraphs[0].text == 'Подпись: Петров Пётр Петрович'
    assert result.tables[0].cell(0, 0).text == 'Автомобиль: Лада'
    with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(response.content)) as after:
        for part in before.namelist():
            if part not in {'word/document.xml', 'word/footer1.xml'}:
                assert before.read(part) == after.read(part), part
    assert client.post(f'/api/generate/templates/{edited}/source', json=payload).status_code == 422


def test_uploaded_pipeline_download_and_expiry(m4_client):
    client, store = m4_client
    uploaded = client.post('/api/pipeline/files', files={'file': ('input.txt', 'Текст сценария'.encode())}).json()
    spec = {'steps': [{'op': 'ingest.file', 'params': {'path': uploaded['value']}, 'output': 'doc'},
                      {'op': 'extract.text', 'input': 'doc', 'output': 'text'},
                      {'op': 'render.docx', 'input': 'text', 'params': {'output_path': '../result.docx'}}]}
    response = client.post('/api/pipeline/run', data={'spec': json.dumps(spec)}).json()
    assert response['result']['ok'], response
    file = client.get(response['downloads'][0]['url'])
    assert 'Текст сценария' in '\n'.join(p.text for p in Document(io.BytesIO(file.content)).paragraphs)
    assert len(response['downloads']) == 1
    store.delete(uploaded['value'].split(':')[1])
    expired = client.post('/api/pipeline/run', data={'spec': json.dumps(spec)}).json()
    assert not expired['success'] and 'заново' in expired['error']


def test_preview_download_is_filled_document(m4_client):
    client, _ = m4_client
    name = import_sample(client)
    response = client.post('/api/generate', data={'template': name, 'preview': 'true'}).json()
    assert response['success'], response
    download = client.get(response['download'])
    assert Document(io.BytesIO(download.content)).paragraphs[0].text == 'Служебная записка'
    if response['available']:
        assert client.get(f"{response['download']}/pages/1").content.startswith(b'\x89PNG')
    assert client.get(f"{response['download']}/pages/0").status_code == 404


@pytest.mark.parametrize('fmt', ['docx', 'html'])
def test_incomplete_preview_is_marked_draft_but_final_stays_strict(m4_client, fmt):
    from textalchemy.web.services.template_variables import save_template_copy

    client, _ = m4_client
    doc = Document()
    doc.add_paragraph('{{ title }}')
    doc.add_paragraph('Количество: {{ count }}')
    doc.add_paragraph('{{ picture }}')
    data = io.BytesIO()
    doc.save(data)
    name = save_template_copy(data.getvalue(), {'fields': [
        {'name': 'title', 'type': 'string', 'description': 'Заголовок'},
        {'name': 'count', 'type': 'integer'}, {'name': 'picture', 'type': 'image'},
    ]})
    params = {'title': ' ', 'count': 0}
    request = {'template': name, 'format': fmt, 'params': json.dumps(params)}
    preview = client.post('/api/generate', data={**request, 'preview': 'true'}).json()
    assert preview['success'] and preview['draft'], preview
    assert {item['name'] for item in preview['missing_fields']} == {'title', 'picture'}
    assert preview['filename'].startswith('Черновик')
    file = client.get(preview['download'])
    text = '\n'.join(p.text for p in Document(io.BytesIO(file.content)).paragraphs) if fmt == 'docx' else file.text
    assert '[Не заполнено: Заголовок]' in text and 'Количество: 0' in text
    assert '[Не заполнено: picture]' in text
    final = client.post('/api/generate', data=request).json()
    assert not final['success'] and 'picture' in final['errors']
    invalid = client.post('/api/generate', data={**request, 'preview': 'true',
        'params': json.dumps({'title': 'Заголовок', 'count': 'не число'})}).json()
    assert not invalid['success'] and 'count' in invalid['errors']


def test_selected_archive_excludes_unselected_and_active(m4_client, tmp_path):
    client, store = m4_client
    entries = []
    for index, status in enumerate(['done', 'done', 'queued']):
        task_id = f'file{index}'
        path = tmp_path / f'{index}.txt'
        path.write_text(str(index))
        artifact = store.store_artifact(task_id, path, path.name)
        store.set(task_id, {'status': status, 'artifact': artifact})
        entries.append({'task_id': task_id, 'name': path.name})
    store.set_job('batch', {'files': entries})
    response = client.get('/api/convert/jobs/batch/archive', params={'task_ids': '["file1"]'})
    assert response.status_code == 200, response.text
    with ZipFile(io.BytesIO(response.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))['files']
        assert len(manifest) == 1 and archive.read(manifest[0]['file']) == b'1'
    for invalid in ('[]', '["alien"]', '["file2"]', '["file1", "file1"]'):
        assert client.get('/api/convert/jobs/batch/archive', params={'task_ids': invalid}).status_code == 409


def test_selected_cancel_leaves_other_tasks(m4_client, monkeypatch):
    client, store = m4_client
    for task_id in ['one', 'two']:
        store.set(task_id, {'status': 'queued'})
    store.set_job('batch', {'files': [{'task_id': name, 'name': name} for name in ['one', 'two']]})
    calls = []

    class Service:
        def cancel(self, task_id):
            calls.append(task_id)
            return {'task_id': task_id}

    monkeypatch.setattr('textalchemy.web.routes.convert._task_service', Service)
    before = store.get('two')
    response = client.post('/api/convert/jobs/batch/cancel', data={'task_ids': '["one"]'})
    assert response.status_code == 200
    assert calls == ['one'] and store.get('two') == before
    assert client.post('/api/convert/jobs/batch/cancel', data={'task_ids': ''}).status_code == 400


@pytest.mark.parametrize('filename, data', [('bad.docx', b'not zip'), ('bad.txt', b'text')])
def test_invalid_upload_creates_no_template(m4_client, filename, data):
    client, _ = m4_client
    before = client.get('/api/generate/templates').json()
    assert client.post('/api/generate/import', files={'file': (filename, data)}).status_code == 422
    assert client.get('/api/generate/templates').json() == before


def test_cancel_race_preserves_newly_completed_result(m4_client, tmp_path):
    from textalchemy.web.services.conversion_tasks import ConversionTaskService

    _, store = m4_client
    path = tmp_path / 'ready.txt'
    path.write_text('Completed')
    store.set('race', {'status': 'running'})

    class FinishingQueue:
        def cancel(self, task_id):
            artifact = store.store_artifact(task_id, path, path.name)
            store.set(task_id, {'status': 'done', 'artifact': artifact})
            return False

    service = ConversionTaskService(store=store, queue=FinishingQueue(), workspace_factory=None)
    with pytest.raises(ValueError, match='готовый результат'):
        service.cancel('race')
    task = store.get('race')
    assert task['status'] == 'done'
    assert store.result_path('race', task['artifact']).read_text() == 'Completed'
