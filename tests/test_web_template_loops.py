"""Проверки реального DOCX после создания и повторной правки цикла."""
import io
import json

import pytest
from docx import Document

from tests import test_web_template_conditions as fixtures
from textalchemy.web.services.template_variables import custom_path

loop_client = fixtures.condition_client


def save(client, name='example', **overrides):
    inspection = client.get(f'/api/generate/templates/{name}/loops').json()
    block = next(item for item in inspection['blocks'] if item['text'].startswith('Optional'))
    return client.post(f'/api/generate/templates/{name}/loops', json={
        'block': block['id'], 'variable': block['variables'][0], 'field': 'items',
        'revision': inspection['revision'], **overrides})


def test_loop_copy_reopen_and_generate(loop_client):
    client, path = loop_client
    original = path.read_bytes()
    response = save(client)
    assert response.status_code == 200, response.text
    name = response.json()['name']
    assert path.read_bytes() == original
    schema = client.get(f'/api/generate/templates/{name}/schema').json()['schema']
    assert schema['fields'][-1]['type'] == 'array'
    assert schema['fields'][-1]['default'] == []
    for items in ([], ['First'], ['First', 'Second', 'First']):
        result = client.post('/api/generate', data={'template': name, 'params': json.dumps({'title': 'Global', 'items': items})})
        assert result.headers['content-type'].startswith('application/vnd.openxmlformats'), result.text
        doc = Document(io.BytesIO(result.content))
        assert [p.text for p in doc.paragraphs] == ['Always visible', *['Optional ' + item for item in items], 'Final paragraph']
        assert all(p.runs[0].bold for p in doc.paragraphs[1:-1])
    changed = save(client, name, field='entries').json()['name']
    texts = [p.text for p in Document(custom_path(changed)).paragraphs]
    assert texts.count('{% for __ta_item in entries %}') == 1
    assert texts.count('{% endfor %}') == 1
    result = client.post('/api/generate', data={'template': changed, 'params': json.dumps({
        'title': 'Global', 'items': ['Old'], 'entries': ['New']})})
    assert [p.text for p in Document(io.BytesIO(result.content)).paragraphs] == [
        'Always visible', 'Optional New', 'Final paragraph']


@pytest.mark.parametrize('overrides', [{'field': 'title'}, {'field': '__ta_item'}, {'field': 'for'},
                                       {'block': 999}, {'variable': 'unknown'}, {'revision': 'stale'}])
def test_invalid_loop_does_not_publish(loop_client, overrides):
    client, path = loop_client
    original = path.read_bytes()
    assert save(client, **overrides).status_code == 422
    assert path.read_bytes() == original
    assert len(client.get('/api/generate/templates').json()) == 1


def test_loop_schema_change_and_nested_directive(loop_client):
    client, path = loop_client
    inspection = client.get('/api/generate/templates/example/loops').json()
    schema_path = path.with_suffix('.schema.json')
    schema = json.loads(schema_path.read_text(encoding='utf-8'))
    schema['fields'][0]['description'] = 'Changed'
    schema_path.write_text(json.dumps(schema), encoding='utf-8')
    assert save(client, revision=inspection['revision']).status_code == 422
    doc = Document(path)
    doc.add_paragraph('{% if title %}')
    doc.add_paragraph('Nested')
    doc.add_paragraph('{% endif %}')
    doc.save(path)
    assert client.get('/api/generate/templates/example/loops').status_code == 422
