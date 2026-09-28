"""Editing lists preserves the stored values and the generated document."""
import importlib
import io
import json

import pytest

from tests import test_browser_e2e as fixtures

browser, e2e_server, page = fixtures.browser, fixtures.e2e_server, fixtures.page


@pytest.fixture
def list_template(tmp_path, monkeypatch):
    from docx import Document

    templates = tmp_path / 'templates'
    templates.mkdir()
    doc = Document()
    for text in ('{{ title }}', '{% for item in items %}', '{{ item }}', '{% endfor %}'):
        doc.add_paragraph(text)
    doc.save(templates / 'list.docx')
    schema = {'fields': [{'name': 'title', 'type': 'string'}, {'name': 'items', 'type': 'array', 'default': []}]}
    (templates / 'list.schema.json').write_text(json.dumps(schema), encoding='utf-8')
    monkeypatch.setattr('textalchemy.generate.template.TEMPLATES_DIR', templates)
    monkeypatch.setattr(importlib.import_module('textalchemy.web.app'), 'data_dir', tmp_path / 'data')


@pytest.mark.parametrize('width', [375, 1280])
def test_list_editor_restore_dataset_and_document(e2e_server, page, list_template, width):
    from axe_core_python.sync_playwright import Axe
    from docx import Document
    from playwright.sync_api import expect

    page.set_viewport_size({'width': width, 'height': 900})
    page.goto(f'{e2e_server}/generate')
    field = page.locator('[data-field-name=items]')
    rows = field.locator('.list-row textarea')
    expect(field.locator('.list-editor')).to_be_visible()
    expect(page.locator('#field-items')).to_be_hidden()
    page.locator('#field-title').fill('Список')
    for value in ('Первый', 'Второй', 'Первый'):
        field.get_by_role('button', name='Добавить строку', exact=True).click()
        rows.last.fill(value)
    field.get_by_role('button', name='Строка 2: выше', exact=True).click()
    expect(rows.first).to_have_value('Второй')
    field.get_by_role('button', name='Удалить строку 3', exact=True).click()
    expect(rows).to_have_count(2)
    page.reload()
    page.locator('#restoreDraft').click()
    expect(rows.first).to_have_value('Второй')
    expect(rows.last).to_have_value('Первый')
    page.locator('#savedData > summary').click()
    page.locator('#datasetName').fill('Список для документа')
    page.locator('#datasetCreate').click()
    expect(page.locator('#datasetStatus')).to_contain_text('версия 1')
    rows.first.fill('Несохранённое в наборе')
    page.once('dialog', lambda dialog: dialog.accept())
    page.locator('#datasetLoad').click()
    expect(rows.first).to_have_value('Второй')
    with page.expect_download(timeout=60000) as download:
        page.locator('#genBtn').click()
    doc = Document(io.BytesIO(download.value.path().read_bytes()))
    assert [p.text for p in doc.paragraphs if p.text] == ['Список', 'Второй', 'Первый']
    # Exact strings (including line breaks and numeric-looking text) survive both modes.
    rows.first.fill('001\nПродолжение')
    field.get_by_role('button', name='JSON', exact=True).click()
    assert json.loads(page.locator('#field-items').input_value()) == ['001\nПродолжение', 'Первый']
    field.get_by_role('button', name='По строкам', exact=True).click()
    expect(rows.first).to_have_value('001\nПродолжение')
    for _ in range(2):
        field.get_by_role('button', name='Удалить строку 1', exact=True).click()
    expect(page.locator('#field-items')).to_have_value('[]')
    expect(field.get_by_role('button', name='Добавить строку', exact=True)).to_be_focused()
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    result = Axe().run(page)
    violations = [v for v in result['violations'] if v.get('impact') in ('serious', 'critical')]
    assert not violations, json.dumps(violations, ensure_ascii=False, indent=2)
    assert page.e2e_errors == []


@pytest.mark.parametrize('value', [
    '[1, true, null, {"x": "y"}]', '[unfinished', '["a", "a", ""]',
    pytest.param(json.dumps(['value'] * 201), id='long-list'),
])
def test_json_mode_preserves_values(e2e_server, page, list_template, value):
    from playwright.sync_api import expect

    page.goto(f'{e2e_server}/generate')
    field = page.locator('[data-field-name=items]')
    field.get_by_role('button', name='JSON', exact=True).click()
    page.locator('#field-items').fill(value)
    field.get_by_role('button', name='По строкам', exact=True).click()
    expect(page.locator('#field-items')).to_have_value(value)
    if value != '["a", "a", ""]':
        expect(page.locator('#field-items')).to_be_visible()
        expect(field.locator('.list-rows')).to_be_hidden()
    page.reload()
    page.locator('#restoreDraft').click()
    expect(page.locator('#field-items')).to_have_value(value)
    if value == '["a", "a", ""]':
        expect(field.locator('.list-row textarea')).to_have_count(3)
    assert page.e2e_errors == []
