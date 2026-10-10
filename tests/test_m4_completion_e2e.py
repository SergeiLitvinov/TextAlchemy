"""Complete M4 paths through the actual HTTP application in Chromium."""

import json
from zipfile import ZipFile

import pytest
from docx import Document
from playwright.sync_api import expect

from tests import test_browser_e2e as fixtures
from tests import test_web_m4_completion as acceptance

browser, e2e_server, page = fixtures.browser, fixtures.e2e_server, fixtures.page
m4_client, sample_bytes = acceptance.m4_client, acceptance.sample_bytes


def accessible(page):
    from axe_core_python.sync_playwright import Axe

    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
    errors = [
        (v["id"], [n["html"] for n in v["nodes"]])
        for v in Axe().run(page)["violations"]
        if v.get("impact") in ("serious", "critical")
    ]
    assert not errors, errors
    assert page.e2e_errors == []


def test_auto_preview_keeps_latest_input_and_serializes_requests(e2e_server, page, m4_client):
    requests = []
    page.route('**/api/generate/live-preview', lambda route: requests.append(route))
    # Hold the initial preview deliberately: typing can occur before it finishes.
    with page.expect_request('**/api/generate/live-preview'):
        page.goto(e2e_server + '/generate')
        fixtures._generator_step(page)
    page.locator('#field-title').fill('Первый вариант')
    with page.expect_request(lambda request: request.url.endswith('/api/generate/live-preview')
                             and 'Первый вариант' in (request.post_data or '')):
        requests[0].fulfill(json={'success': True, 'html': '<html><head></head><body>Начальный ответ</body></html>'})
    page.locator('#field-title').fill('Второй вариант')
    page.locator('#field-title').fill('Последний вариант')
    # Let the debounce elapse while the server still holds the earlier request.
    page.wait_for_timeout(1800)
    assert len(requests) == 2
    requests[1].fulfill(json={'success': True, 'html': '<html><head></head><body>Устаревший ответ</body></html>'})
    page.wait_for_timeout(1800)
    assert len(requests) == 3
    assert 'Последний вариант' in requests[2].request.post_data
    assert 'Устаревший ответ' not in (page.locator('#livePreviewFrame').get_attribute('srcdoc') or '')
    requests[2].fulfill(json={'success': True, 'html': '<html><head></head><body>Последний вариант</body></html>'})
    expect(page.frame_locator('#livePreviewFrame').locator('body')).to_have_text('Последний вариант')
    expect(page.locator('#livePreviewFrame')).to_have_attribute('sandbox', '')
    assert "default-src 'none'" in page.locator('#livePreviewFrame').get_attribute('srcdoc')
    expect(page.locator('#filledDownload')).to_be_hidden()
    assert page.e2e_errors == []


@pytest.mark.parametrize("width", [375, 1280])
@pytest.mark.parametrize("unavailable", [False, True])
def test_sample_to_field_preview_download(e2e_server, page, m4_client, width, unavailable):
    if unavailable:
        def missing_pages(route):
            response = route.fetch()
            data = response.json()
            if data.get("success"):
                data.update(available=False, pages=0)
            route.fulfill(response=response, json=data)
        page.route("**/api/generate", missing_pages)
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(e2e_server + "/generate")
    fixtures._generator_step(page)
    fixtures._generator_step(page, 0)
    page.get_by_text("Создать шаблон из своего DOCX", exact=True).click()
    page.locator("#sourceUpload").set_input_files(
        {
            "name": "memo.docx",
            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "buffer": sample_bytes(),
        }
    )
    page.locator("#importTemplate").click()
    expect(page.locator("#sourceInspect")).to_be_visible(timeout=30000)
    page.locator("#sourceInspect").click()
    expect(page.locator("#sourceBlocks option")).to_have_count(5)
    page.locator("#sourceQuery").fill("Иванов Иван Иванович")
    page.locator("#sourceSearch").click()
    expect(page.locator("#sourceMatches input")).to_have_count(3)
    page.locator("#sourceMatches input").nth(0).check()
    page.locator("#sourceMatches input").nth(2).check()
    page.locator("#sourceLabel").fill("ФИО руководителя")
    page.locator("#sourceSave").click()
    expect(page.locator("#field-field_1")).to_be_visible(timeout=30000)
    page.locator("#field-field_1").fill("Петров Пётр Петрович")
    page.locator('#previewAccurate').click()
    expect(page.locator("#filledDownload")).to_be_visible(timeout=60000)
    with page.expect_download() as downloaded:
        page.locator("#filledDownload").click()
    result = Document(downloaded.value.path())
    assert result.paragraphs[1].text == "Руководитель: Петров Пётр Петрович"
    assert result.paragraphs[2].text == "Исполнитель: Иванов Иван Иванович"
    assert result.sections[0].footer.paragraphs[0].text == "Подпись: Петров Пётр Петрович"
    page.locator("#field-field_1").fill("")
    page.locator('#previewAccurate').click()
    expect(page.locator("#filledDownload")).to_have_text("Скачать черновик", timeout=60000)
    expect(page.locator("#filledPreviewStatus")).to_contain_text("Осталось заполнить: ФИО руководителя")
    if unavailable:
        expect(page.locator("#filledPreviewStatus")).to_contain_text("Просмотр страниц недоступен")
    with page.expect_download() as draft_download:
        page.locator("#filledDownload").click()
    draft = Document(draft_download.value.path())
    assert "[Не заполнено: ФИО руководителя]" in draft.paragraphs[1].text
    assert draft.paragraphs[2].text == "Исполнитель: Иванов Иван Иванович"
    accessible(page)
    page.locator("#field-field_1").fill("Новое значение")
    expect(page.locator("#filledDownload")).to_be_hidden()
    expect(page.locator("#livePreview")).to_be_visible()
    page.reload()
    fixtures._generator_step(page)
    page.locator("#restoreDraft").click()
    expect(page.locator("#field-field_1")).to_have_value("Новое значение")
    accessible(page)


@pytest.mark.parametrize("width", [375, 1280])
def test_pipeline_uploaded_input_to_download(e2e_server, page, m4_client, width):
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(e2e_server + "/pipeline")
    page.locator("[data-upload=path]").set_input_files(
        {"name": "sample.txt", "mimeType": "text/plain", "buffer": "Содержимое файла".encode()}
    )
    expect(page.locator("[data-upload-status]").first).to_contain_text("загружен")
    page.locator(".step-card .op-select").nth(2).select_option("render.docx")
    page.locator("[data-name=output_path]").fill("result.docx")
    page.locator("#runBtn").click()
    expect(page.locator("#pipeline-result-title")).to_have_text("Сценарий выполнен")
    with page.expect_download() as download:
        page.locator("#resultDownloads a").click()
    assert "Содержимое файла" in "\n".join(p.text for p in Document(download.value.path()).paragraphs)
    accessible(page)


@pytest.mark.parametrize("width", [375, 1280])
def test_selected_results_and_cancellation(e2e_server, page, m4_client, tmp_path, width):
    _, store = m4_client
    entries = []
    for index, status in enumerate(["done", "done", "queued", "queued"]):
        task_id = f"choice{index}"
        path = tmp_path / f"{index}.txt"
        path.write_text(str(index))
        artifact = store.store_artifact(task_id, path, path.name)
        store.set(task_id, {"status": status, "artifact": artifact, "filename": path.name})
        entries.append({"task_id": task_id, "name": path.name, "target_format": "txt", "mode": "balanced"})
    store.set_job("selection", {"job_id": "selection", "files": entries, "mode": "balanced", "target_format": "txt"})
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(e2e_server + "/convert")
    fixtures._wait_convert_ready(page)
    fixtures._open_batch_history(page, "selection")
    page.locator("[data-archive-task=choice1]").check()
    with page.expect_download() as download:
        page.locator("[data-archive-selected]").click()
    with ZipFile(download.value.path()) as archive:
        manifest = json.loads(archive.read("manifest.json"))["files"]
        assert len(manifest) == 1 and archive.read(manifest[0]["file"]) == b"1"
    page.locator("[data-cancel-task=choice2]").check()
    page.locator("[data-cancel-selected]").click()
    expect(page.locator("[data-retry-task=choice2]")).to_be_visible()
    assert store.get("choice2")["status"] == "cancelled"
    assert store.get("choice3")["status"] == "queued"
    accessible(page)


@pytest.mark.parametrize("width", [375, 1280])
def test_pdf_page_regions_select_block(e2e_server, page, m4_client, width):
    import fitz

    pdf = fitz.open()
    sheet = pdf.new_page()
    sheet.insert_text((60, 60), "First block")
    sheet.insert_text((60, 220), "Second block")
    data = pdf.tobytes()
    pdf.close()
    page.set_viewport_size({"width": width, "height": 900})
    page.goto(e2e_server + "/pdf-order")
    page.locator("#orderFile").set_input_files({"name": "blocks.pdf", "mimeType": "application/pdf", "buffer": data})
    expect(page.locator(".pdf-hotspot")).to_have_count(2)
    page.locator(".pdf-hotspot").last.click()
    expect(page.locator("#orderBlocks option:checked")).to_contain_text("Second block")
    page.locator("#blockClassification").select_option("heading1")
    page.locator("#orderSave").click()
    expect(page.locator("#orderStatus")).to_contain_text("сохранены")
    accessible(page)
