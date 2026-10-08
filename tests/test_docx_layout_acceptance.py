"""Independent own DOCX witnesses for widths, opaque smartTag and strict budgets."""

import json
import os
import sys
import uuid

import pytest
from docx import Document
from docx.oxml.ns import qn

from tests.corpus.docx_layout import write_percentage_table, write_smart_tag
from textalchemy.convert.executor import ConversionExecutor, ConversionRequest
from textalchemy.core.quality_policy import QualityPolicy
from textalchemy.core.types import DocFormat
from tools.acceptance.corpus import ROOT, digest, target_program_check

BUILDERS = {"percentage-table": write_percentage_table, "smart-tag": write_smart_tag}


def two_cycles(directory, case, *, max_loss_issues=0):
    source = directory / "source.docx"
    BUILDERS[case](source)
    original = digest(source)
    current = source
    paths = [source]
    executor = ConversionExecutor()
    for cycle in (1, 2):
        model = directory / f"cycle-{cycle}.json"
        imported = executor.execute(ConversionRequest(
            current, model, DocFormat.DOCX, DocFormat.MODEL, quality_policy=QualityPolicy(max_loss_issues)))
        assert imported.success, imported.to_dict()
        output = directory / f"cycle-{cycle}.docx"
        exported = executor.execute(ConversionRequest(
            model, output, DocFormat.MODEL, DocFormat.DOCX, quality_policy=QualityPolicy(max_loss_issues)))
        assert exported.success, exported.to_dict()
        for report in (imported, exported):
            losses = [issue for issue in report.issues if issue.severity.value == "loss"]
            if case == "smart-tag":
                assert [issue.feature for issue in losses] == ["docx.smart-tags"]
                assert all(issue.location for issue in losses)
            else:
                assert not losses
        paths.append(output)
        current = output
    assert digest(source) == original
    return paths


def native_text(path):
    # Include wrapped run text independently; python-docx paragraph.text omits smartTag runs.
    return " ".join("".join(Document(path)._element.xpath(".//w:t/text()")).split())


def first_row_widths(path):
    table = Document(path).tables[0]
    return [(cell._tc.tcPr.tcW.get(qn("w:type")), cell._tc.tcPr.tcW.get(qn("w:w")))
            for cell in table.rows[0].cells]


def test_percentage_table_two_cycles_preserve_preferred_widths(tmp_path):
    source, first, second = two_cycles(tmp_path, "percentage-table")
    assert first_row_widths(source) == [("pct", "500"), ("pct", "4000"), ("pct", "500")]
    for result in (first, second):
        widths = first_row_widths(result)
        assert widths == first_row_widths(source)
        assert native_text(result) == native_text(source)


def smart_tag_xml(path):
    nodes = Document(path)._element.xpath(".//w:smartTag")
    assert len(nodes) == 1
    return [(node.tag, sorted(node.attrib.items()), node.text) for node in nodes[0].iter()]


def test_smart_tag_two_cycles_preserve_text_and_opaque_wrapper(tmp_path):
    source, first, second = two_cycles(tmp_path, "smart-tag", max_loss_issues=1)
    assert native_text(source) == "Before OwnSmartToken After"
    assert native_text(first) == native_text(second) == native_text(source)
    assert smart_tag_xml(first) == smart_tag_xml(second) == smart_tag_xml(source)


def test_smart_tag_strict_budget_preserves_previous_result(tmp_path):
    source = tmp_path / "source.docx"
    write_smart_tag(source)
    original = source.read_bytes()
    output = tmp_path / "result.json"
    output.write_bytes(b"Previous result")
    report = ConversionExecutor().execute(ConversionRequest(
        source, output, DocFormat.DOCX, DocFormat.MODEL, quality_policy=QualityPolicy(0)))
    assert not report.success and output.read_bytes() == b"Previous result"
    assert source.read_bytes() == original
    diagnostic = report.metrics["step_metrics"]["docx.model"]["import_diagnostics"][0]
    assert diagnostic["code"] == "docx.smart-tags" and diagnostic["reason"] == "native-opaque"
    assert diagnostic["location"] and diagnostic["measurement"]["state"] == "opaque"
    assert report.metrics["step_metrics"]["docx.model"]["assessment_complete"] is False


@pytest.mark.skipif(sys.platform != "win32" or os.getenv("TEXTALCHEMY_WORD_ACCEPTANCE") != "1",
                    reason="Requires installed Word and explicit local acceptance run")
@pytest.mark.parametrize("case", ["percentage-table", "smart-tag"])
def test_native_word_preserves_selected_text_and_geometry_in_both_cycles(tmp_path, case):
    paths = two_cycles(tmp_path, case, max_loss_issues=1 if case == "smart-tag" else 0)
    directory = ROOT / ".textalchemy/acceptance/minimal-docx-tests" / uuid.uuid4().hex
    directory.mkdir(parents=True)
    observations = []
    for label, path in zip(("source", "cycle-1", "cycle-2"), paths, strict=True):
        before = digest(path)
        record = target_program_check(path)
        record["input_unchanged"] = before == digest(path) == record["artifact_sha256"]
        assert record["input_unchanged"]
        observations.append(record["inventory"])
        (directory / f"{label}.json").write_text(json.dumps(record, indent=2), encoding="utf-8")
    baseline, first, second = observations
    if case == "percentage-table":
        assert baseline["pages"] == 1
        assert first["pages"] == second["pages"] == baseline["pages"]
        original = baseline["table_geometry"][0]
        assert original["available"] and original["scope"] == "first-row-native-cell-widths"
        assert original["widths_points"][1] > 7 * original["widths_points"][0]
        for result in (first, second):
            geometry = result["table_geometry"][0]
            assert geometry["available"]
            assert geometry["widths_points"] == pytest.approx(original["widths_points"], abs=.1)
            assert result["normalized_text_sha256"] == baseline["normalized_text_sha256"]
    else:
        assert baseline["whitespace_tokens"] == 3
        assert first["whitespace_tokens"] == second["whitespace_tokens"] == baseline["whitespace_tokens"]
        assert baseline["normalized_text_sha256"] == first["normalized_text_sha256"] == second["normalized_text_sha256"]
