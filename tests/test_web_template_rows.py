"""Повторяемая строка проверяется через настоящий DOCX и схему копии."""
import io
import json

import pytest
from docx import Document

from tests import test_web_template_conditions as fixtures
from textalchemy.web.services.template_variables import custom_path

row_client = fixtures.condition_client


def make_table(path):
    doc = Document()
    doc.add_paragraph('Outside {{ title }}')
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Unit'
    paragraph = table.cell(1, 0).paragraphs[0]
    paragraph.add_run('{{ti').bold = True
    paragraph.add_run('tle}}')
    table.cell(1, 1).text = 'kg'
    table.cell(2, 0).text = 'Total'
    table.cell(2, 1).text = 'End'
    doc.save(path)


def save(client, name='example', **overrides):
    inspection = client.get(f'/api/generate/templates/{name}/table-loops').json()
    block = inspection['blocks'][0]
    return client.post(f'/api/generate/templates/{name}/table-loops', json={
        'block': block['id'], 'variable': block['variables'][0], 'field': 'items',
        'revision': inspection['revision'], **overrides})


def test_row_copy_generation_and_reediting(row_client):
    client, path = row_client
    make_table(path)
    original = path.read_bytes()
    response = save(client)
    assert response.status_code == 200, response.text
    name = response.json()['name']
    assert path.read_bytes() == original
    for values in ([], ['A'], ['A', 'B', 'A']):
        result = client.post('/api/generate', data={'template': name, 'params': json.dumps({'title': 'Global', 'items': values})})
        assert result.headers['content-type'].startswith('application/vnd.openxmlformats'), result.text
        doc = Document(io.BytesIO(result.content))
        assert doc.paragraphs[0].text == 'Outside Global'
        assert [[c.text for c in row.cells] for row in doc.tables[0].rows] == [
            ['Name', 'Unit'], *[[value, 'kg'] for value in values], ['Total', 'End']]
        assert all(row.cells[0].paragraphs[0].runs[0].bold for row in doc.tables[0].rows[1:-1])
    changed = save(client, name, field='entries').json()['name']
    rows = Document(custom_path(changed)).tables[0].rows
    assert sum('{% for' in row.cells[0].text for row in rows) == 1
    result = client.post('/api/generate', data={'template': changed, 'params': json.dumps({
        'title': 'Global', 'items': ['Old'], 'entries': ['New']})})
    assert Document(io.BytesIO(result.content)).tables[0].cell(1, 0).text == 'New'


@pytest.mark.parametrize('change', [{'field': 'title'}, {'field': 'self'}, {'block': '0:999'},
                                   {'variable': 'missing'}, {'revision': 'stale'}])
def test_invalid_row_preserves_source(row_client, change):
    client, path = row_client
    make_table(path)
    original = path.read_bytes()
    assert save(client, **change).status_code == 422
    assert path.read_bytes() == original
    assert len(client.get('/api/generate/templates').json()) == 1


def test_merged_row_not_offered(row_client):
    client, path = row_client
    make_table(path)
    doc = Document(path)
    doc.tables[0].cell(1, 0).merge(doc.tables[0].cell(1, 1))
    doc.save(path)
    assert client.get('/api/generate/templates/example/table-loops').json()['blocks'] == []
