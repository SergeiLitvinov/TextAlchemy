"""Тесты Web API генерации: схема шаблона, preview, валидация данных."""

import base64

from textalchemy.generate.model_template import TemplateImage
from textalchemy.web.routes.generate import _decode_data_uri, _prepare_web_params, template_schema_for


def test_template_schema_for_derived_from_docx(tmp_path):
    from docx import Document

    template = tmp_path / "letter.docx"
    doc = Document()
    doc.add_paragraph("Dear {{name}}, your code is {{code}}.")
    doc.save(str(template))

    schema, source = template_schema_for("letter", templates_dir=tmp_path)

    assert source == "derived"
    assert {field.name for field in schema.fields} == {"name", "code"}
    assert all(field.type.value == "string" for field in schema.fields)


def test_template_schema_for_sidecar_wins(tmp_path):
    from docx import Document

    template = tmp_path / "letter.docx"
    doc = Document()
    doc.add_paragraph("Hello {{name}}.")
    doc.save(str(template))
    (tmp_path / "letter.schema.json").write_text(
        '{"fields": {"name": {"type": "string", "required": true}}, "allow_extra": false}',
        encoding="utf-8",
    )

    schema, source = template_schema_for("letter", templates_dir=tmp_path)

    assert source == "sidecar"
    assert [field.name for field in schema.fields] == ["name"]
    assert schema.allow_extra is False


def test_prepare_web_params_decodes_data_uri_images(tmp_path):
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+M8AAAMBAQDJ/pLvAAAAAElFTkSuQmCC"
    )
    from textalchemy.generate.template_schema import TemplateField, TemplateSchema, TemplateValueType

    schema = TemplateSchema(fields=[TemplateField("logo", TemplateValueType.IMAGE)])

    prepared = _prepare_web_params(
        schema,
        {"logo": "data:image/png;base64," + base64.b64encode(png).decode()},
    )

    assert isinstance(prepared["logo"], TemplateImage)
    assert prepared["logo"].data == png
    assert prepared["logo"].media_type == "image/png"


def test_decode_data_uri_media_type():
    image = _decode_data_uri("data:image/jpeg;base64,AAAA")
    assert image.data == b"\x00\x00\x00"
    assert image.media_type == "image/jpeg"
