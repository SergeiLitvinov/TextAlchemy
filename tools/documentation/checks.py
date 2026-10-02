"""Check maintained examples and the completeness of historical task migration."""

import contextlib
import io
import json
import re
import shlex
import tempfile

import yaml

from textalchemy.__main__ import _setup_parser, main
from tools.documentation.generated import ROOT


def check_commands(text):
    parser = _setup_parser()
    count = 0
    for fence in re.findall(r"```powershell\n(.*?)```", text, re.S):
        for line in fence.splitlines():
            if line.startswith("uv run textalchemy "):
                try:
                    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                        parser.parse_args(shlex.split(line)[3:])
                except SystemExit as error:
                    if error.code:
                        raise ValueError(f"Неверная команда в руководстве: {line}") from error
                count += 1
    return count


def check_mapping():
    source = (ROOT / "docs/history/todo-2026-09-13.md").read_text(encoding="utf-8")
    expected = [index for index, line in enumerate(source.splitlines(), 1) if re.match(r"\s*- \[[ xX]\]", line)]
    table = (ROOT / "docs/history/todo-milestone-map.md").read_text(encoding="utf-8")
    mapped = [int(number) for number in re.findall(r"^\| L(\d+) \|", table, re.M)]
    if sorted(mapped) != expected:
        raise ValueError("Карта старого TODO содержит пропуски или дубликаты")
    return len(mapped)


def check_active_plan(text: str) -> int:
    """В активном плане остаются только незавершённые задачи с уникальными ID своей вехи."""
    milestone = None
    identifiers = set()
    for line in text.splitlines():
        heading = re.match(r"## M(\d+)\.", line)
        if heading:
            milestone = heading[1]
        elif line.startswith("## "):
            milestone = None
        checkbox = re.match(r"\s*- \[([ xX])\]", line)
        if not checkbox:
            continue
        if checkbox[1] != " ":
            raise ValueError("Закрытые задачи должны быть удалены из активного TODO")
        task = re.match(r"\s*- \[ \] \*\*(M(\d+)\.\d+)\b", line)
        if not task or task[2] != milestone:
            raise ValueError("Задача TODO должна иметь ID своей вехи")
        if task[1] in identifiers:
            raise ValueError(f"Повторяющийся ID TODO: {task[1]}")
        identifiers.add(task[1])
    return len(identifiers)


def run(arguments):
    output = io.StringIO()
    with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
        code = main([str(item) for item in arguments])
    if code:
        raise ValueError(f"Ошибка примера {arguments[0]}: {output.getvalue()}")
    return output.getvalue()


def check_examples(guide, directory):
    import fitz
    from docx import Document

    from textalchemy.pipeline.runner import run_pipeline

    bibliography = json.loads(run(["run", ROOT / "docs/examples/bibliography-to-json.yaml", "--json"]))
    entries = json.loads(bibliography["final"])
    if len(entries) != 1 or entries[0]["title"] != "Исследование электрических сетей" or entries[0]["year"] != 2020:
        raise ValueError("Библиографический конвейер изменил данные примера")

    spec = yaml.safe_load((ROOT / "docs/examples/text-to-docx.yaml").read_text(encoding="utf-8"))
    for step in spec["steps"]:
        params = step.get("params", {})
        if "path" in params:
            params["path"] = str(ROOT / params["path"])
        if "output_path" in params:
            params["output_path"] = str(directory / "pipeline.docx")
    result = run_pipeline(spec)
    if not result.ok:
        raise ValueError(f"Ошибка примера конвейера: {result.error}")
    original = (ROOT / "docs/examples/input.txt").read_text(encoding="utf-8").split()
    actual = " ".join(paragraph.text for paragraph in Document(directory / "pipeline.docx").paragraphs).split()
    if original != actual:
        raise ValueError("Конвейер изменил текст примера")
    template_source = re.search(r"```jinja2\n(.*?)```", guide, re.S)
    if not template_source:
        raise ValueError("В руководстве отсутствует проверяемый пример шаблона")
    template = Document()
    for line in template_source[1].strip().splitlines():
        template.add_paragraph(line)
    template_path = directory / "template.docx"
    template.save(template_path)
    data = ROOT / "docs/examples/template-data.json"
    schema = ROOT / "docs/examples/template-schema.json"
    run(["template-check", template_path, "--schema", schema, "--data", data, "--json"])
    values = json.loads(data.read_text(encoding="utf-8"))
    expected = [values["title"], values["body"], *[f"{item['name']}: {item['value']}" for item in values["items"]]]
    for extension in ["docx", "html", "pdf"]:
        output = directory / f"result.{extension}"
        run(["generate", template_path, output, "--schema", schema, "--data", data, "--json"])
        if extension == "docx":
            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        elif extension == "pdf":
            with fitz.open(output) as pdf:
                text = "\n".join(page.get_text() for page in pdf)
        else:
            from bs4 import BeautifulSoup

            text = BeautifulSoup(output.read_text(encoding="utf-8"), "html.parser").get_text()
        if not all(value in text for value in expected):
            raise ValueError(f"Пример {extension} не содержит ожидаемого текста")


def check():
    guide = (ROOT / "docs/guide/index.md").read_text(encoding="utf-8")
    count = sum(check_commands(path.read_text(encoding="utf-8")) for path in (ROOT / "docs/guide").glob("*.md"))
    mapped = check_mapping()
    active = check_active_plan((ROOT / "TODO.md").read_text(encoding="utf-8"))
    parent = ROOT / ".textalchemy/docs-check"
    parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=parent) as temporary:
        from pathlib import Path

        directory = Path(temporary).resolve()
        if not directory.is_relative_to(parent.resolve()):
            raise ValueError("Unexpected examples workspace")
        check_examples(guide, directory)
    print(
        f"Checked {count} command examples, {mapped} migrated tasks, {active} active tasks, "
        "bibliography/text pipelines and DOCX/HTML/PDF generation",
        flush=True,
    )
