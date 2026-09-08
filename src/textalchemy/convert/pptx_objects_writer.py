"""Native PPTX tables and category charts."""

from textalchemy.convert.pptx_cell_writer import configure_cell
from textalchemy.convert.pptx_text_writer import write_text
from textalchemy.core.diagnostics import IssueSeverity
from textalchemy.core.document_model import Paragraph


def write_table(slide, block, geometry, report, location):
    from pptx.util import Pt

    rows = len(block.rows)
    columns = max((len(row.cells) for row in block.rows), default=0)
    if not rows or not columns:
        report.add(IssueSeverity.LOSS, "tables", "Пустая таблица пропущена.", location)
        return
    table = slide.shapes.add_table(rows, columns, *geometry).table
    widths = block.properties.get("column_widths_pt", [])
    if len(widths) == columns and all(width > 0 for width in widths):
        for column, width in zip(table.columns, widths):
            column.width = Pt(width)
    for r, row in enumerate(block.rows):
        if row.properties.get("height_pt", 0) > 0:
            table.rows[r].height = Pt(row.properties["height_pt"])
        for c, cell in enumerate(row.cells):
            target = table.cell(r, c)
            cell_location = f"{location}.rows[{r}].cells[{c}]"
            if target.is_spanned:
                configure_cell(target, cell.properties, report, cell_location)
                continue
            if cell.row_span > 1 or cell.column_span > 1:
                end_r, end_c = r + cell.row_span - 1, c + cell.column_span - 1
                if end_r < rows and end_c < columns:
                    target.merge(table.cell(end_r, end_c))
                else:
                    report.add(IssueSeverity.LOSS, "tables", "Некорректное объединение ячеек пропущено.", location)
            paragraph_count = 0
            for child in cell.blocks:
                if isinstance(child, Paragraph):
                    write_text(target.text_frame, child, report, cell_location, append=paragraph_count > 0)
                    paragraph_count += 1
                else:
                    report.add(IssueSeverity.LOSS, "tables", "Нетекстовый блок внутри ячейки не перенесён.", location)
            configure_cell(target, cell.properties, report, cell_location)
    if block.properties:
        report.add(IssueSeverity.LOSS, "styles", "Тема таблицы не восстановлена; явное оформление ячеек сохранено.", location)


def write_chart(slide, data, geometry, report, location):
    from textalchemy.convert.pptx_chart_writer import write_chart as export_chart

    return export_chart(slide, data, geometry, report, location)
