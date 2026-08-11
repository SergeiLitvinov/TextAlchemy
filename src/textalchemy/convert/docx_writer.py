"""Экспорт богатой промежуточной модели в редактируемый DOCX."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from textalchemy.convert.docx_drawing_writer import write_image
from textalchemy.convert.docx_section_writer import configure_section, section_start_type
from textalchemy.convert.docx_style_writer import write_styles
from textalchemy.convert.docx_table_writer import write_table
from textalchemy.convert.docx_text_writer import add_paragraph, write_formula, write_paragraph_content
from textalchemy.core.diagnostics import ConversionReport, IssueSeverity
from textalchemy.core.document_model import (
    Block,
    DocumentModel,
    Formula,
    Image,
    Paragraph,
    Section,
    Table,
)
from textalchemy.ooxml.package import restore_package_graph


def write_docx_model(document: DocumentModel, output_path: str | Path) -> ConversionReport:
    """Записать ``DocumentModel`` в DOCX и вернуть отчёт о потерях."""

    from docx import Document
    from docx.enum.section import WD_SECTION

    output = Path(output_path)
    report = ConversionReport(output)
    target = Document()
    _write_metadata(target, document.metadata)
    write_styles(target, document.styles, report)
    target.settings.odd_and_even_pages_header_footer = any(
        bool(section.properties.odd_and_even_pages_header_footer)
        or bool(section.even_page_headers)
        or bool(section.even_page_footers)
        for section in document.sections
    )

    sections = document.sections or [Section()]
    counters = {"paragraphs": 0, "tables": 0, "images": 0, "formulas": 0}
    for section_index, source_section in enumerate(sections):
        if section_index == 0:
            target_section = target.sections[0]
        else:
            target_section = target.add_section(section_start_type(source_section, WD_SECTION))
        configure_section(target_section, source_section, document, report, section_index, counters, _write_blocks)
        _write_blocks(target, source_section.blocks, document, report, f"sections[{section_index}]", counters)

    _restore_package_graph(target, document, report)
    try:
        from textalchemy.core.artifacts import ArtifactWorkspace
        from textalchemy.core.io import atomic_copy

        with ArtifactWorkspace(prefix="textalchemy_docx_writer_") as workspace:
            staged = workspace.artifact_path("output.docx")
            target.save(staged)
            workspace.validate_artifact(staged)
            atomic_copy(staged, output)
    except Exception as error:  # noqa: BLE001 - diagnostic boundary must return a report
        report.add(IssueSeverity.ERROR, "docx-write", str(error))
        return report
    report.metrics.update(counters)
    report.metrics["sections"] = len(sections)
    report.metrics["resources"] = len(document.resources)
    return report


def _write_metadata(target: Any, metadata: dict[str, Any]) -> None:
    properties = target.core_properties
    for name in ("title", "subject", "author", "keywords", "comments"):
        value = metadata.get(name)
        if value is not None:
            setattr(properties, name, str(value))


def _restore_package_graph(target: Any, document: DocumentModel, report: ConversionReport) -> None:
    graph = document.package
    if graph is None or graph.format != "ooxml":
        return
    try:
        restore_package_graph(target.part, graph)
    except Exception as error:  # noqa: BLE001 - foreign package parts are a diagnostic boundary
        report.add(IssueSeverity.LOSS, "package-graph", f"OOXML package graph could not be restored: {error}")


def _write_blocks(
    container: Any,
    blocks: list[Block],
    document: DocumentModel,
    report: ConversionReport,
    location: str,
    counters: dict[str, int],
) -> None:
    for index, block in enumerate(blocks):
        block_location = f"{location}.blocks[{index}]"
        if isinstance(block, Paragraph):
            paragraph = add_paragraph(container, block, document, report, block_location)
            write_paragraph_content(paragraph, block, document, report, block_location, counters)
            counters["paragraphs"] += 1
        elif isinstance(block, Table):
            write_table(container, block, document, report, block_location, counters, _write_blocks)
            counters["tables"] += 1
        elif isinstance(block, Formula):
            paragraph = container.add_paragraph()
            write_formula(paragraph, block, report, block_location)
            counters["formulas"] += 1
        elif isinstance(block, Image):
            paragraph = container.add_paragraph()
            write_image(paragraph, block, document, report, block_location)
            counters["images"] += 1


__all__ = ["write_docx_model"]
