"""Native DrawingML text and Office Math, without importing PPTX at package import."""

from textalchemy.convert.pptx_paragraph_writer import configure_frame, configure_paragraph
from textalchemy.core.color import ColorValue
from textalchemy.core.diagnostics import IssueSeverity
from textalchemy.core.document_model import Formula, FormulaFormat, TextRun


def set_color(color_format, value, report, location):
    from pptx.dml.color import RGBColor

    if value is None:
        return
    try:
        if isinstance(value, dict):
            value = ColorValue.from_dict(value)
        color = value if isinstance(value, ColorValue) else ColorValue.from_hex(value)
        color_format.rgb = RGBColor.from_string(color.to_hex().lstrip("#"))
        if color.alpha != 1 or color.icc_profile or color.blend_mode != "normal":
            report.add(IssueSeverity.LOSS, "color", "Цвет приведён к непрозрачному RGB.", location)
    except ValueError:
        report.add(IssueSeverity.LOSS, "color", "Цвет не удалось перенести.", location)


def write_text(frame, block, report, location, *, append=False):
    from pptx.util import Pt

    if append:
        paragraph = frame.add_paragraph()
    else:
        frame.clear()
        configure_frame(frame, block, report, location)
        paragraph = frame.paragraphs[0]
    paragraph_index = 0
    configure_paragraph(paragraph, block, paragraph_index)
    for item in block.content:
        if isinstance(item, Formula):
            write_formula(paragraph, item, report, location)
        elif isinstance(item, TextRun):
            for index, text in enumerate(item.text.split("\n")):
                if index:
                    if item.properties.get("pptx_break") == "line":
                        paragraph.add_line_break()
                    else:
                        paragraph = frame.add_paragraph()
                        paragraph_index += 1
                        configure_paragraph(paragraph, block, paragraph_index)
                run = paragraph.add_run()
                run.text = text
                style = item.style
                run.font.name = style.font_family
                run.font.size = Pt(style.font_size.pt) if style.font_size else None
                run.font.bold, run.font.italic = style.bold, style.italic
                run.font.underline = style.underline
                set_color(run.font.color, style.color, report, location)
                if style.superscript or style.subscript:
                    run._r.get_or_add_rPr().set("baseline", "30000" if style.superscript else "-25000")
                if item.link:
                    run.hyperlink.address = item.link
                if style.background:
                    report.add(IssueSeverity.LOSS, "styles", "Фон текста не перенесён.", location)
        else:
            report.add(IssueSeverity.LOSS, "raster_images", "Встроенное в строку изображение не перенесено.", location)


def write_formula(paragraph, formula, report, location):
    from lxml import etree

    if formula.format is FormulaFormat.OMML:
        try:
            root = etree.fromstring(formula.value.encode(), etree.XMLParser(resolve_entities=False, no_network=True))
            ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
            math = root if root.tag == f"{{{ns}}}oMath" else root.find(f".//{{{ns}}}oMath")
            if math is not None:
                wrapper = etree.Element("{http://schemas.microsoft.com/office/drawing/2010/main}m", nsmap={
                    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
                })
                wrapper.append(math)
                paragraph._p.append(wrapper)
                return
        except etree.XMLSyntaxError:
            pass
    paragraph.add_run().text = formula.fallback_text or formula.value
    report.add(IssueSeverity.LOSS, "formulas", "Формула сохранена как текст; необходим корректный OMML.", location)
