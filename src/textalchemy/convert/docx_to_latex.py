"""Конвертер DOCX → LaTeX.

Читает .docx через python-docx, восстанавливает структуру (заголовки,
абзацы, списки, таблицы) в объект ``Text`` и отдаёт готовый ``.tex``
через существующий рендерер ``render.latex``. Текст сохраняется как
текст (никаких растровых вставок).
"""
from __future__ import annotations

from pathlib import Path

from textalchemy.convert.base import BaseConverter, ConversionResult
from textalchemy.core.types import Block, BlockType, Table, Text
from textalchemy.pipeline.render import render_latex


class DocxToLatexConverter(BaseConverter):
    @property
    def name(self) -> str:
        return "docx2latex"

    def convert(self, input_path: str | Path, output_path: str | Path) -> ConversionResult:
        prepared = self._prepare(input_path, output_path)
        if prepared is None:
            return ConversionResult(Path(input_path), Path(output_path), False, "Input file not found")
        input_path, output_path = prepared
        try:
            from docx import Document as DocxDocument
        except ImportError:
            return ConversionResult(
                input_path, output_path, False,
                "python-docx не установлен (добавьте в зависимости для DOCX→LaTeX)",
            )

        try:
            doc = DocxDocument(str(input_path))
            text = Text(source_format=__import__("textalchemy.core.types", fromlist=["DocFormat"]).DocFormat.DOCX)
            text.engine = "docx2latex"

            for para in doc.paragraphs:
                style = (para.style.name or "") if para.style else ""
                content = (para.text or "").strip()
                if not content:
                    continue

                if style.startswith("Heading") or style.startswith("Title"):
                    try:
                        level = int(style.replace("Heading", "").strip() or "1")
                    except ValueError:
                        level = 1
                    if style.startswith("Title"):
                        level = 1
                    text.blocks.append(Block(type=BlockType.HEADING, text=content, level=level))
                elif style.startswith("List") or style.startswith("List Bullet") or style.startswith("List Number"):
                    text.blocks.append(Block(type=BlockType.LIST_ITEM, text=content))
                else:
                    text.blocks.append(Block(type=BlockType.PARAGRAPH, text=content))

            for table in doc.tables:
                rows: list[list[str]] = []
                for row in table.rows:
                    rows.append([cell.text.strip() for cell in row.cells])
                if rows:
                    text.tables.append(Table(rows=rows))

            tex = render_latex(text=text, title=input_path.stem, author="")
            Path(output_path).write_text(tex, encoding="utf-8")
        except Exception as e:  # noqa: BLE001
            return ConversionResult(input_path, output_path, False, str(e))

        return ConversionResult(input_path, output_path, output_path.exists())
