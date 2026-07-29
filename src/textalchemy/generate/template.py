from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from textalchemy.core.exceptions import GenerateError

TEMPLATES_DIR = Path(__file__).parent / "templates"


def _replace_in_paragraph(paragraph, items):
    """Подставить значения плейсхолдеров {{key}} в параграфе.

    Текст собирается из всех runs (плейсхолдеры в docx часто разбиты между
    несколькими runs), заменяется, и результат записывается в первый run с
    сохранением его стиля; остальные runs очищаются. Это сохраняет базовое
    форматирование параграфа (в отличие от перезаписи ``para.text``).
    """
    if not paragraph.runs:
        if not items:
            return
        text = paragraph.text
        replaced = text
        for key, value in items:
            replaced = replaced.replace("{{" + key + "}}", str(value))
        if replaced != text:
            paragraph.add_run(replaced)
        return

    text = "".join(run.text for run in paragraph.runs)
    replaced = text
    for key, value in items:
        replaced = replaced.replace("{{" + key + "}}", str(value))
    if replaced == text:
        return
    first_run = paragraph.runs[0]
    first_run.text = replaced
    for run in paragraph.runs[1:]:
        run.text = ""


@dataclass
class DocumentTemplate:
    name: str
    description: str
    template_type: str = "docx"


class TemplateEngine:
    def __init__(self, templates_dir: str | Path | None = None):
        self.templates_dir = Path(templates_dir) if templates_dir else TEMPLATES_DIR

    def list_templates(self) -> List[DocumentTemplate]:
        if not self.templates_dir.exists():
            return []
        templates: List[DocumentTemplate] = []
        for f in sorted(self.templates_dir.iterdir()):
            if f.suffix.lower() in (".docx", ".doc", ".dotx"):
                templates.append(DocumentTemplate(
                    name=f.stem,
                    description=f"Template: {f.name}",
                    template_type=f.suffix[1:],
                ))
        return templates

    def resolve_template(self, template_name: str) -> Path:
        """Найти шаблон по пути или имени в каталоге шаблонов."""

        direct = Path(template_name)
        candidates = [direct] if direct.is_file() else []
        candidates.extend(
            [
                self.templates_dir / template_name,
                self.templates_dir / f"{template_name}.docx",
                self.templates_dir / f"{template_name}.doc",
            ]
        )
        for candidate in candidates:
            if candidate.is_file():
                return candidate
        available = ", ".join(t.name for t in self.list_templates())
        msg = f"Template '{template_name}' not found. Available: {available or 'none'}"
        raise GenerateError(msg)

    def generate(
        self,
        template_name: str,
        output_path: str | Path,
        params: Optional[Dict[str, Any]] = None,
    ) -> Path:
        output_path = Path(output_path)
        template_path = self.resolve_template(template_name)

        output_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            from docx import Document
            doc = Document(str(template_path))
            params = params or {}

            # Заменяем по убыванию длины ключа, чтобы более длинные
            # плейсхолдеры (возможно, содержащие префиксы других) не
            # конфликтовали при подстановке.
            items = sorted(params.items(), key=lambda kv: len(kv[0]), reverse=True)

            for para in doc.paragraphs:
                _replace_in_paragraph(para, items)

            # Заменяем плейсхолдеры и внутри таблиц шаблона.
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            _replace_in_paragraph(para, items)

            doc.save(str(output_path))
        except Exception as e:
            raise GenerateError(f"Failed to generate document: {e}") from e

        return output_path


def generate_document(
    template_name: str,
    output_path: str | Path,
    params: Optional[Dict[str, Any]] = None,
    templates_dir: Optional[str | Path] = None,
) -> Path:
    engine = TemplateEngine(templates_dir)
    return engine.generate(template_name, output_path, params)


def list_templates(templates_dir: Optional[str | Path] = None) -> List[DocumentTemplate]:
    engine = TemplateEngine(templates_dir)
    return engine.list_templates()
