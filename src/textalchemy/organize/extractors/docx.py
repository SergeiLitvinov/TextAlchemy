def read_docx(file_path: str) -> str:
    try:
        from docx import Document

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs]
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    table_texts.append(cell.text)
        return "\n".join(paragraphs + table_texts)
    except Exception:
        return ""
