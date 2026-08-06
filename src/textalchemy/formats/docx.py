"""DOCX → Text.

Единый ридер DOCX; раньше было две копии (extract/text.py и
organize/extractors/docx.py).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Union

from textalchemy.core.document_model import (
    Block as RichBlock,
)
from textalchemy.core.document_model import (
    ConversionMode,
    DocumentModel,
    PackageGraph,
)
from textalchemy.core.io import check_archive_safety
from textalchemy.core.types import Block, BlockType, DocFormat, Table, Text
from textalchemy.formats.docx_section import read_section
from textalchemy.formats.docx_style import (
    read_document_defaults,
    read_document_styles,
)
from textalchemy.formats.docx_table import read_table
from textalchemy.formats.docx_text import read_paragraph
from textalchemy.ooxml.package import (
    RELATIONSHIP_TYPE,
    load_package_graph,
)


def read_docx(path: Union[str, Path], *, include_tables: bool = True) -> Text:
    from docx import Document  # type: ignore[import-not-found]

    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(p)
    check_archive_safety(p)
    doc = Document(str(p))
    blocks: list[Block] = []
    tables: list[Table] = []
    plain: list[str] = []

    for para in doc.paragraphs:
        t = para.text
        blocks.append(Block(type=BlockType.PARAGRAPH, text=t))
        if t:
            plain.append(t)

    if include_tables:
        for tbl in doc.tables:
            rows = [[cell.text for cell in row.cells] for row in tbl.rows]
            tables.append(Table(rows=rows))
            for row in rows:
                plain.append(" | ".join(row))

    return Text(
        blocks=blocks,
        tables=tables,
        plain="\n".join(plain),
        source_format=DocFormat.DOCX,
        engine="python-docx",
    )


def read_docx_model(
    path: Union[str, Path],
    *,
    mode: ConversionMode = ConversionMode.BALANCED,
) -> DocumentModel:
    """Импортировать DOCX в богатую модель с оформлением и ресурсами."""

    from docx import Document as OpenDocument  # type: ignore[import-not-found]
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    check_archive_safety(source)
    document = OpenDocument(str(source))
    model = DocumentModel(
        mode=mode,
        source_format=DocFormat.DOCX.value,
        metadata=_document_metadata(document, source),
        styles=read_document_styles(document),
    )
    model.package = _load_docx_package_graph(document)
    document_defaults = read_document_defaults(document, model)
    if document_defaults is not None:
        model.styles["__doc_defaults__"] = document_defaults

    section_blocks: list[list[RichBlock]] = [[]]
    for item in document.iter_inner_content():
        if isinstance(item, DocxParagraph):
            section_blocks[-1].append(read_paragraph(item, model))
            if item._p.pPr is not None and item._p.pPr.sectPr is not None:
                section_blocks.append([])
        elif isinstance(item, DocxTable):
            section_blocks[-1].append(read_table(item, model, _container_blocks))

    if len(section_blocks) > 1 and not section_blocks[-1]:
        section_blocks.pop()
    for index, blocks in enumerate(section_blocks):
        source_section = document.sections[min(index, len(document.sections) - 1)]
        model.sections.append(
            read_section(
                source_section,
                blocks,
                model,
                index,
                _container_blocks,
                odd_and_even_pages=bool(document.settings.odd_and_even_pages_header_footer),
            )
        )
    return model


def _load_docx_package_graph(document: Any) -> PackageGraph | None:
    supported = set(RELATIONSHIP_TYPE.values())
    recursive = {RELATIONSHIP_TYPE["footnotes"], RELATIONSHIP_TYPE["endnotes"]}
    return load_package_graph(
        document.part,
        format_name="ooxml",
        supported_relationships=supported,
        recursive_relationships=recursive,
        include=lambda relationship: (
            relationship.reltype != RELATIONSHIP_TYPE["numbering"] or _document_uses_numbering(document)
        ),
    )


def _document_uses_numbering(document: Any) -> bool:
    from docx.oxml.ns import qn

    if document.element.xpath(".//w:pPr/w:numPr"):
        return True
    used_style_ids = {
        node.get(qn("w:val"))
        for node in document.element.xpath(".//w:pPr/w:pStyle")
        if node.get(qn("w:val"))
    }
    return any(
        style.style_id in used_style_ids and bool(style.element.xpath("./w:pPr/w:numPr"))
        for style in document.styles
    )


def _document_metadata(document: Any, source: Path) -> dict[str, Any]:
    props = document.core_properties
    return {
        "title": props.title or "",
        "subject": props.subject or "",
        "author": props.author or "",
        "keywords": props.keywords or "",
        "comments": props.comments or "",
        "created": props.created.isoformat() if props.created else None,
        "modified": props.modified.isoformat() if props.modified else None,
        "source_name": source.name,
    }


def _container_blocks(container: Any, model: DocumentModel) -> list[RichBlock]:
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    result: list[RichBlock] = []
    for item in container.iter_inner_content():
        if isinstance(item, DocxParagraph):
            result.append(read_paragraph(item, model))
        elif isinstance(item, DocxTable):
            result.append(read_table(item, model, _container_blocks))
    return result


__all__ = ["read_docx", "read_docx_model"]
