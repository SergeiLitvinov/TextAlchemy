import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from textalchemy.core.exceptions import ExtractError


def clean_text(text: str) -> str:
    _bs_placeholder = "\x00BS\x00"
    text = text.replace("\\", _bs_placeholder)
    replacements = {
        "{": "\\{",
        "}": "\\}",
        "$": "\\$",
        "&": "\\&",
        "#": "\\#",
        "^": "\\^{}",
        "_": "\\_{}",
        "~": "\\textasciitilde{}",
        "%": "\\%",
        "[": "\\[",
        "]": "\\]",
    }
    for char, repl in replacements.items():
        text = text.replace(char, repl)
    text = text.replace(_bs_placeholder, "\\textbackslash{}")
    text = text.replace("…", "\\dots{}")
    return text


def _format_run_text(run) -> str:
    text = clean_text(run.text)
    if run.bold:
        text = f"\\textbf{{{text}}}"
    if run.italic:
        text = f"\\textit{{{text}}}"
    return text


def _format_paragraph_text(para) -> str:
    parts = [_format_run_text(run) for run in para.runs if run.text.strip()]
    if not parts:
        return clean_text(para.text)
    return "".join(parts)


def _get_paragraph_style(paragraph):
    style_name = paragraph.style.name.lower() if paragraph.style and paragraph.style.name else ""
    if "heading 1" in style_name or "заголовок 1" in style_name:
        return "section"
    if "heading 2" in style_name or "заголовок 2" in style_name:
        return "subsection"
    if "heading 3" in style_name or "заголовок 3" in style_name:
        return "subsubsection"
    if "title" in style_name or "название" in style_name:
        return "title"
    return None


_PREAMBLE = """\\documentclass[12pt,a4paper]{article}
\\usepackage[T2A]{fontenc}
\\usepackage[utf8]{inputenc}
\\usepackage[russian]{babel}
\\usepackage{amsmath,amssymb}
\\usepackage{graphicx}
\\usepackage{geometry}
\\geometry{left=3cm,right=1.5cm,top=2cm,bottom=2cm}
\\usepackage{setspace}
\\onehalfspacing
"""


def docx_to_latex(input_path: str | Path, output_path: str | Path | None = None, doc_type: str = "manuscript") -> str:
    input_path = Path(input_path)
    if not input_path.exists():
        raise ExtractError(f"File not found: {input_path}")
    try:
        doc = Document(str(input_path))
    except Exception as e:
        raise ExtractError(f"Failed to read document: {e}") from e

    lines = [_PREAMBLE, "\\begin{document}"]
    if doc_type == "abstract":
        lines.append("\\thispagestyle{empty}")
    lines.append("")

    for para in doc.paragraphs:
        if not para.runs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            clean = clean_text(text)
        else:
            formatted = _format_paragraph_text(para)
            clean = formatted.strip()
            if not clean:
                lines.append("")
                continue

        style = _get_paragraph_style(para)
        if style == "title":
            continue
        elif style == "section":
            lines.append(f"\\section{{{clean}}}")
        elif style == "subsection":
            lines.append(f"\\subsection{{{clean}}}")
        elif style == "subsubsection":
            lines.append(f"\\subsubsection{{{clean}}}")
        else:
            align = para.alignment
            if align == WD_ALIGN_PARAGRAPH.CENTER:
                lines.append(f"\\begin{{center}}{clean}\\end{{center}}")
            elif align == WD_ALIGN_PARAGRAPH.RIGHT:
                lines.append(f"\\begin{{flushright}}{clean}\\end{{flushright}}")
            else:
                lines.append(clean)
        lines.append("")

    for table in doc.tables:
        if not table.rows:
            continue
        col_count = len(table.columns)
        col_spec = "|" + "|".join(["c"] * col_count) + "|"
        lines.append(f"\\begin{{tabular}}{{{col_spec}}}")
        lines.append("\\hline")
        for i, row in enumerate(table.rows):
            cells = [clean_text(cell.text.strip()) for cell in row.cells]
            lines.append(" & ".join(cells) + " \\\\")
            if i < len(table.rows) - 1:
                lines.append("\\hline")
        lines.append("\\hline")
        lines.append("\\end{tabular}")
        lines.append("")

    lines.append("\\end{document}")
    result = "\n".join(lines)

    if output_path:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        Path(output_path).write_text(result, encoding="utf-8")
    return result


def docx_to_latex_pandoc(input_path: str | Path, output_path: str | Path, doc_type: str = "manuscript") -> str:
    input_path = Path(input_path)
    output_path = Path(output_path)
    if not input_path.exists():
        raise ExtractError(f"File not found: {input_path}")

    lua_filter = Path(__file__).parent / "lua-filters" / "sanitize.lua"
    lua_filter_str = str(lua_filter) if lua_filter.exists() else ""

    cmd = [
        "pandoc",
        str(input_path),
        "-o",
        str(output_path),
        "--from",
        "docx",
        "--to",
        "latex",
        "--standalone",
        "--top-level-division=chapter",
    ]
    if lua_filter_str:
        cmd.extend(["--lua-filter", lua_filter_str])

    try:
        subprocess.run(cmd, capture_output=True, text=True, check=True, timeout=120)
    except FileNotFoundError:
        raise ExtractError("pandoc not found. Install pandoc or use 'latex' (python-docx) format.")
    except subprocess.TimeoutExpired:
        raise ExtractError("pandoc timed out")
    except subprocess.CalledProcessError as e:
        raise ExtractError(f"pandoc error: {e.stderr[:500]}")

    tex = output_path.read_text(encoding="utf-8")
    title_text = "Document"
    author_text = "Author"
    russian_support = (
        "\\usepackage[T2A]{fontenc}\n"
        "\\usepackage[utf8]{inputenc}\n"
        "\\usepackage[russian]{babel}\n"
        "\\usepackage{amsmath,amsfonts,amssymb}\n"
        "\\usepackage{graphicx}\n"
        "\\usepackage{geometry}\n"
        "\\geometry{a4paper, margin=2cm}\n"
    )
    tex = tex.replace("\\begin{document}", russian_support + "\n\\begin{document}")
    tex = tex.replace(
        "\\maketitle",
        f"\\title{{{title_text}}}\n\\author{{{author_text}}}\n\\date{{\\today}}\n\\maketitle",
    )
    output_path.write_text(tex, encoding="utf-8")
    return tex
