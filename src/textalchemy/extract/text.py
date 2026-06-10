from pathlib import Path

from docx import Document

from textalchemy.core.exceptions import ExtractError


def extract_text(input_path: str | Path, output_path: str | Path | None = None) -> str:
    input_path = Path(input_path)
    if not input_path.exists():
        raise ExtractError(f"File not found: {input_path}")
    if input_path.suffix.lower() not in (".docx",):
        raise ExtractError(f"Unsupported format: {input_path.suffix}")
    try:
        doc = Document(str(input_path))
        text = "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        raise ExtractError(f"Failed to extract text: {e}") from e
    if output_path:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        Path(output_path).write_text(text, encoding="utf-8")
    return text


def extract_text_with_tables(input_path: str | Path, output_path: str | Path | None = None) -> str:
    input_path = Path(input_path)
    if not input_path.exists():
        raise ExtractError(f"File not found: {input_path}")
    try:
        doc = Document(str(input_path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    except Exception as e:
        raise ExtractError(f"Failed to extract text: {e}") from e
    if output_path:
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        Path(output_path).write_text(text, encoding="utf-8")
    return text
